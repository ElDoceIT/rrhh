import io
import mimetypes
import os
import secrets
from contextlib import asynccontextmanager
from dataclasses import replace
from datetime import date, datetime, time, timedelta
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path
from typing import Annotated, NamedTuple
from urllib.parse import quote, unquote

from fastapi import Depends, FastAPI, Form, HTTPException, Query, Request
from fastapi.responses import HTMLResponse, RedirectResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from starlette.middleware.sessions import SessionMiddleware
from sqlalchemy import and_, func, select, text
from sqlalchemy.exc import IntegrityError, SQLAlchemyError
from sqlalchemy.orm import Session, selectinload
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from app.database.connection import SessionLocal, engine, get_db
from app.database.models import ConceptoExcepcional, Feriado, HoraExtra, ReglaHora, TipoContratacion, Usuario, UsuarioConceptoExcepcional, UsuarioRol
from app.services.ad_auth import ActiveDirectoryAuthError, authenticate_ad_user, list_ad_group_users
from app.services.access_catalog import CONVENIOS_DISPONIBLES, PERFILES_DISPONIBLES, ROLES_DISPONIBLES
from app.services.calculo_horas import (
    CalculoHorasError,
    calcular_horas_totales,
    calcular_resultado_nocturnas_en_jornada,
    calcular_resultado_concepto_excepcional,
    calcular_resultado_dia_trabajado,
    calcular_resultado_domingo,
    calcular_resultado_horas_extra,
    calcular_resultado_otra_carga,
    calcular_resultado_reintegro,
    clasificar_tipo_dia,
    dividir_carga_en_fechas,
    TIPO_NOCTURNAS_EN_JORNADA,
    TIPOS_NOCTURNAS_EN_JORNADA,
)
from app.services.local_auth import (
    authenticate_local_user,
    find_user,
    hash_password,
    initialize_seed_admin,
    is_local_username,
    local_usernames,
    normalize_username,
    session_user,
)


@asynccontextmanager
async def lifespan(_: FastAPI):
    initialize_seed_admin(SessionLocal)
    yield


mimetypes.add_type("text/css", ".css")
mimetypes.add_type("application/javascript", ".js")

app = FastAPI(title="Horas extras", lifespan=lifespan)
app.mount("/static", StaticFiles(directory="app/static"), name="static")
templates = Jinja2Templates(directory="app/templates")

TUTORIAL_HTML_PATH = Path(__file__).resolve().parent / "resources" / "tutorial" / "index.html"
TUTORIAL_KEYS = ("CISPREN", "SAL", "SAT", "FC", "MONOTRIBUTISTA", "JEFATURA", "RRHH")
TUTORIAL_LABELS = {
    "ADMIN": "Abrir tutorial completo",
    "RRHH": "Abrir guía de RRHH",
    "JEFATURA": "Abrir guía para Jefatura",
    "CISPREN": "Abrir guía de carga CISPREN",
    "SAL": "Abrir guía de carga SAL",
    "SAT": "Abrir guía de carga SAT",
    "FC": "Abrir guía de carga FC",
    "MONOTRIBUTISTA": "Abrir guía de carga MONOTRIBUTISTA",
}
PUBLIC_PATHS = {"/login", "/health", "/health/db", "/health/database"}


@app.middleware("http")
async def require_login(request: Request, call_next):
    path = request.url.path
    if path.startswith("/static/") or path in PUBLIC_PATHS:
        return await call_next(request)
    session_user_data = request.session.get("user")
    if session_user_data is not None:
        active_role = str(session_user_data.get("role") or "").upper()
        users_path = path.startswith("/configuracion/usuarios")
        holidays_path = path.startswith("/configuracion/feriados")
        exceptional_path = path.startswith("/configuracion/conceptos-excepcionales")
        preloads_path = path.startswith("/configuracion/precargas-autorizadas")
        active_profile = str(session_user_data.get("profile") or "").upper()
        if path.startswith("/configuracion") and not (
            active_role == "ADMIN"
            or ((users_path or holidays_path) and active_role == "RRHH")
            or (preloads_path and active_role == "RRHH")
            or (exceptional_path and (active_role == "RRHH" or active_profile == "JEFE"))
        ):
            return HTMLResponse("Acceso denegado: se requiere el rol ADMIN.", status_code=403)
        return await call_next(request)
    return RedirectResponse(f"/login?next={quote(path, safe='/')}", status_code=303)


app.add_middleware(
    SessionMiddleware,
    secret_key=os.getenv("SESSION_SECRET_KEY", "rrhh-dev-session-key-change-me"),
    same_site="lax",
    https_only=os.getenv("SESSION_COOKIE_SECURE", "false").lower() in {"1", "true", "yes", "on"},
)


def parse_optional_time(value: str | None) -> time | None:
    if not value:
        return None
    try:
        return time.fromisoformat(value)
    except ValueError as error:
        raise HTTPException(status_code=422, detail="El horario ingresado no es válido.") from error


def rule_form_data(
    convenio: str,
    tipo_dia: str,
    tipo_hora: str,
    hora_nocturna_desde: str | None,
    hora_nocturna_hasta: str | None,
    permite_reintegro: str | None,
    observaciones: str | None,
    modalidad_carga: str | None = None,
    cantidad_predeterminada: Decimal | None = None,
) -> dict:
    convenio = convenio.strip().upper()
    tipo_dia = tipo_dia.strip().upper().replace("Á", "A")
    tipo_hora = tipo_hora.strip()
    if not convenio or not tipo_dia or not tipo_hora:
        raise HTTPException(status_code=422, detail="Completá los campos obligatorios.")
    if convenio not in CONVENIOS_DISPONIBLES:
        raise HTTPException(status_code=422, detail="El convenio seleccionado no es válido.")
    if tipo_dia not in {"HABIL", "FERIADO", "FRANCO", "TODOS"}:
        raise HTTPException(status_code=422, detail="El tipo de día seleccionado no es válido.")
    modalidad_normalizada = (modalidad_carga or "").strip().upper()
    if modalidad_normalizada == "ACTUAL":
        modalidad_normalizada = ""
    if modalidad_normalizada not in {"", "CARGA_POR_DIA", "CARGA_MANUAL"}:
        raise HTTPException(status_code=422, detail="La modalidad de carga no es válida.")
    if modalidad_normalizada and tipo_dia != "TODOS":
        raise HTTPException(
            status_code=422,
            detail="Las modalidades por día o manual sólo corresponden a reglas de tipo TODOS.",
        )
    cantidad_configurada = None
    if modalidad_normalizada == "CARGA_POR_DIA":
        if (
            cantidad_predeterminada is None
            or cantidad_predeterminada <= 0
            or cantidad_predeterminada % Decimal("0.5") != 0
        ):
            raise HTTPException(
                status_code=422,
                detail="Ingresá una cantidad predeterminada válida, de a 0,5.",
            )
        cantidad_configurada = cantidad_predeterminada.quantize(Decimal("0.01"))
    return {
        "convenio": convenio,
        "tipo_dia": tipo_dia,
        "tipo_hora": tipo_hora,
        "hora_nocturna_desde": parse_optional_time(hora_nocturna_desde),
        "hora_nocturna_hasta": parse_optional_time(hora_nocturna_hasta),
        "permite_reintegro": permite_reintegro == "on",
        "observaciones": observaciones.strip() if observaciones else None,
        "modalidad_carga": modalidad_normalizada or None,
        "cantidad_predeterminada": cantidad_configurada,
    }


def clean_optional(value: str | None) -> str | None:
    value = value.strip() if value else ""
    return value or None


def safe_next_url(value: str | None) -> str:
    if not value or not value.startswith("/") or value.startswith("//"):
        return "/"
    return value


def obtener_asignacion_activa(request: Request, db: Session) -> UsuarioRol:
    session_data = request.session.get("user") or {}
    assignment = db.scalar(
        select(UsuarioRol).where(
            UsuarioRol.id_rol == session_data.get("active_assignment_id"),
            UsuarioRol.usuario_id == session_data.get("id"),
        )
    )
    if assignment is None:
        raise HTTPException(
            status_code=403,
            detail="Seleccioná un rol y perfil activo.",
        )
    return assignment


def tutorial_key_for_assignment(assignment: UsuarioRol, usuario: Usuario) -> str:
    role = str(assignment.rol or "").strip().upper()
    profile = str(assignment.perfil or "").strip().upper()
    if role == "ADMIN":
        return "ADMIN"
    if role == "RRHH":
        return "RRHH"
    if profile == "JEFE":
        return "JEFATURA"
    convenio = str(usuario.convenio or "").strip().upper()
    if convenio not in CONVENIOS_DISPONIBLES:
        raise HTTPException(
            status_code=422,
            detail="El usuario activo no tiene un convenio válido para mostrar el tutorial.",
        )
    return convenio


def tutorial_html_for_key(tutorial_key: str) -> str:
    if tutorial_key != "ADMIN" and tutorial_key not in TUTORIAL_KEYS:
        raise HTTPException(status_code=404, detail="El tutorial solicitado no existe.")
    try:
        html = TUTORIAL_HTML_PATH.read_text(encoding="utf-8")
    except OSError as error:
        raise HTTPException(status_code=500, detail="No se pudo abrir el tutorial.") from error

    html = html.replace(
        '<span class="demo-chip">Instructivo · No guarda datos</span>',
        '<a class="demo-chip" href="/" style="text-decoration:none">← Volver al sistema</a>',
        1,
    )
    if tutorial_key == "ADMIN":
        return html

    entry_prefixes = tuple(f"{key}:{{steps:[" for key in TUTORIAL_KEYS)
    filtered_lines = []
    for line in html.splitlines(keepends=True):
        stripped = line.lstrip()
        if stripped.startswith(entry_prefixes) and not stripped.startswith(
            f"{tutorial_key}:{{steps:["
        ):
            continue
        filtered_lines.append(line)
    html = "".join(filtered_lines)
    html = html.replace(
        '<section class="card">',
        '<section class="card" style="display:none" aria-hidden="true">',
        1,
    )
    html = html.replace(
        "<h1>Aprendé a realizar una carga</h1><p>Elegí tu convenio y avanzá a tu ritmo. Podés volver atrás cuando quieras.</p>",
        f"<h1>{TUTORIAL_LABELS[tutorial_key]}</h1><p>Recorrido personalizado según tu rol, perfil y convenio activos.</p>",
        1,
    )
    html = html.replace("  </script>", f'    start("{tutorial_key}");\n  </script>', 1)
    return html


def convenios_visibles_para_reglas(
    assignment: UsuarioRol,
    usuario: Usuario,
    db: Session,
) -> list[str]:
    role = str(assignment.rol or "").strip().upper()
    profile = str(assignment.perfil or "").strip().upper()
    if role in {"ADMIN", "RRHH"}:
        return list(CONVENIOS_DISPONIBLES)
    if profile == "JEFE":
        convenios = db.scalars(
            select(Usuario.convenio)
            .join(Usuario.roles)
            .where(
                Usuario.status.is_(True),
                Usuario.convenio.is_not(None),
                func.upper(UsuarioRol.rol) == role,
                func.upper(UsuarioRol.perfil) == "USUARIO",
            )
            .distinct()
            .order_by(Usuario.convenio)
        ).all()
        return [
            convenio for convenio in CONVENIOS_DISPONIBLES
            if convenio in {str(item).strip().upper() for item in convenios}
        ]
    convenio = str(usuario.convenio or "").strip().upper()
    if convenio not in CONVENIOS_DISPONIBLES:
        raise HTTPException(
            status_code=422,
            detail="El usuario activo no tiene un convenio válido para consultar sus reglas.",
        )
    return [convenio]


def descripcion_concepto_regla(tipo_hora: str) -> dict[str, str]:
    original = str(tipo_hora or "").strip()
    normalizado = original.upper().replace("_", " ")
    if normalizado in {"50", "50%"}:
        return {
            "nombre": "Horas extras al 50 %",
            "categoria": "Horas extras",
            "descripcion": "Tiempo trabajado fuera de la jornada habitual que se liquida con recargo del 50 %.",
            "icono": "bi-clock-history",
        }
    if normalizado in {"100", "100%"}:
        return {
            "nombre": "Horas extras al 100 %",
            "categoria": "Horas extras",
            "descripcion": "Tiempo trabajado fuera de la jornada habitual que se liquida con recargo del 100 %.",
            "icono": "bi-clock-history",
        }
    if normalizado == "DIA TRABAJADO":
        return {
            "nombre": "Día trabajado completo",
            "categoria": "Día trabajado",
            "descripcion": "Representa un franco o feriado trabajado completo. La cantidad registrada es 1 día.",
            "icono": "bi-calendar-check",
        }
    explicaciones = {
        "COMIDA": ("Comida", "Concepto automático", "En SAT se genera una comida cada 3 horas extra acumuladas en la misma jornada.", "bi-cup-hot"),
        "MERIENDA": ("Merienda", "Concepto automático", "En SAT se genera una merienda cada 2 horas extra acumuladas en la misma jornada.", "bi-cup-straw"),
        "DOMINGO": ("Domingo trabajado", "Día trabajado", "En SAT representa el trabajo realizado un domingo y se registra con cantidad 1.", "bi-calendar-week"),
        "HS ARTICULO": ("Horas artículo", "Descanso entre jornadas", "Compensa las horas faltantes para completar el descanso obligatorio entre jornadas.", "bi-moon-stars"),
    }
    if normalizado in explicaciones:
        nombre, categoria, descripcion, icono = explicaciones[normalizado]
        return {
            "nombre": nombre,
            "categoria": categoria,
            "descripcion": descripcion,
            "icono": icono,
        }
    return {
        "nombre": original,
        "categoria": "Concepto de convenio",
        "descripcion": "Concepto general habilitado por las reglas del convenio. La forma de carga y la cantidad dependen de su configuración vigente.",
        "icono": "bi-plus-circle",
    }


def descripcion_modalidad_regla(
    modalidad_carga: str | None,
    cantidad_predeterminada: Decimal | None,
) -> str | None:
    modalidad = str(modalidad_carga or "").strip().upper()
    descripciones = {
        "CARGA_MANUAL": "La cantidad se informa manualmente al realizar la carga.",
        "CARGA_POR_DIA": "Se registra por cada día seleccionado.",
    }
    descripcion = descripciones.get(modalidad)
    if cantidad_predeterminada is not None:
        cantidad = format(cantidad_predeterminada, "f").rstrip("0").rstrip(".")
        detalle = f"Cantidad predeterminada: {cantidad}."
        return f"{descripcion} {detalle}" if descripcion else detalle
    return descripcion


def descripcion_tipo_dia(convenio: str, tipo_dia: str) -> tuple[str, str, str]:
    convenio = convenio.upper()
    tipo_dia = tipo_dia.upper()
    if tipo_dia == "HABIL":
        return (
            "Día hábil", "bi-briefcase",
            "Las horas informadas se clasifican con el porcentaje configurado para el convenio.",
        )
    if tipo_dia in {"FRANCO", "FERIADO"}:
        nombre = "Franco trabajado" if tipo_dia == "FRANCO" else "Feriado trabajado"
        if convenio in {"FC", "MONOTRIBUTISTA"}:
            detalle = "Genera 1 día trabajado y las horas extra se agregan por separado al 100 %."
        else:
            detalle = "El sistema resuelve la carga según la jornada informada y las reglas del convenio."
        if convenio == "CISPREN":
            detalle += " Puede solicitar medio reintegro, reintegro completo o ninguno."
        elif convenio in {"SAL", "SAT"}:
            detalle += " Puede solicitar reintegro completo cuando la regla y el feriado lo permiten."
        else:
            detalle += " No permite solicitar reintegro."
        return nombre, "bi-calendar-check", detalle
    if tipo_dia == "DOMINGO":
        return (
            "Domingo", "bi-calendar-week",
            "Para SAT registra 1 domingo trabajado, sin reintegro. Puede calcular nocturnidad si se informa horario.",
        )
    return (
        "Conceptos para cualquier día", "bi-grid",
        "Conceptos adicionales o automáticos que no dependen de que el día sea hábil, franco o feriado.",
    )


def usuarios_habilitados_para_carga(
    request: Request,
    db: Session,
    destino: str | None = None,
) -> tuple[UsuarioRol, list[Usuario]]:
    assignment = obtener_asignacion_activa(request, db)
    if destino == "rrhh":
        if assignment.rol.upper() not in {"ADMIN", "RRHH"}:
            raise HTTPException(
                status_code=403,
                detail="Sólo RRHH o ADMIN pueden realizar cargas administrativas.",
            )
        usuarios = list(db.scalars(
            select(Usuario)
            .options(selectinload(Usuario.tipo_contratacion))
            .where(
                Usuario.status.is_(True),
                Usuario.convenio.is_not(None),
            )
            .order_by(Usuario.apellido, Usuario.nombre)
        ))
        return assignment, usuarios
    if assignment.perfil.upper() != "USUARIO":
        raise HTTPException(
            status_code=403,
            detail="Sólo los usuarios pueden cargar sus propias horas.",
        )

    current_user = db.get(Usuario, assignment.usuario_id)
    if current_user is None or not current_user.status:
        raise HTTPException(status_code=403, detail="El usuario activo no está habilitado.")
    if not usuario_tiene_funciones_carga(current_user, db):
        raise HTTPException(
            status_code=403,
            detail="No tenés funciones de carga habilitadas.",
        )

    return assignment, [current_user]


CODIGO_PERMISO_CARGAS_FC = "PERMISO_CARGAS_FC"
CODIGO_SWITCHER_NOCTURNIDAD = "SWITCHER_NOCTURNIDAD"
TIPO_SWITCHER_NOCTURNIDAD = "DERIVADO_NOCTURNIDAD"
TIPO_PRECARGA_AUTORIZADA = "PRECARGA_AUTORIZADA"
TIPOS_CONCEPTO_EXCEPCIONAL = (
    "PERMISO",
    "CARGA_POR_DIA",
    "CARGA_MANUAL",
    "PRECARGA_PERIODO",
    TIPO_SWITCHER_NOCTURNIDAD,
)
TIPOS_CONCEPTO_EXCEPCIONAL_CONFIGURABLES = tuple(
    tipo for tipo in TIPOS_CONCEPTO_EXCEPCIONAL
    if tipo not in {"PRECARGA_PERIODO", TIPO_SWITCHER_NOCTURNIDAD}
)


def configurar_concepto_excepcional(
    concepto: ConceptoExcepcional,
    nombre: str,
    codigo: str | None,
    tipo: str,
    descripcion: str | None,
    cantidad_predeterminada: Decimal | None,
    tipo_hora_destino: str | None,
    requiere_observacion: bool,
    requiere_autorizacion: bool,
) -> None:
    nombre_limpio = nombre.strip()
    tipo_normalizado = tipo.strip().upper()
    if not nombre_limpio:
        raise CalculoHorasError("El nombre es obligatorio.")
    if tipo_normalizado not in TIPOS_CONCEPTO_EXCEPCIONAL:
        raise CalculoHorasError("La modalidad del concepto no es válida.")
    codigo_limpio = clean_optional(codigo)
    if codigo_limpio:
        codigo_limpio = codigo_limpio.upper().replace(" ", "_")

    cantidad_configurada = None
    destino_configurado = None
    if tipo_normalizado == "CARGA_POR_DIA":
        if (
            cantidad_predeterminada is None
            or cantidad_predeterminada <= 0
            or cantidad_predeterminada % Decimal("0.5") != 0
        ):
            raise CalculoHorasError("Ingresá una cantidad predeterminada válida, de a 0,5.")
        cantidad_configurada = cantidad_predeterminada.quantize(Decimal("0.01"))
    elif tipo_normalizado == "PRECARGA_PERIODO":
        destino_configurado = clean_optional(tipo_hora_destino)
        if destino_configurado is None:
            raise CalculoHorasError("Seleccioná el tipo de hora que liquidará la precarga.")
        destino_configurado = destino_configurado.upper()

    concepto.nombre = nombre_limpio
    concepto.codigo = codigo_limpio
    concepto.tipo = tipo_normalizado
    concepto.descripcion = clean_optional(descripcion)
    concepto.cantidad_predeterminada = cantidad_configurada
    concepto.tipo_hora_destino = destino_configurado
    concepto.requiere_observacion = requiere_observacion
    concepto.requiere_autorizacion = (
        requiere_autorizacion if tipo_normalizado != "PERMISO" else False
    )


def validar_destino_precarga(concepto: ConceptoExcepcional, db: Session) -> None:
    if concepto.tipo != "PRECARGA_PERIODO":
        return
    destino_existe = db.scalar(
        select(ReglaHora.id).where(
            func.upper(ReglaHora.tipo_hora) == concepto.tipo_hora_destino,
        ).limit(1)
    )
    if destino_existe is None:
        raise CalculoHorasError(
            "El tipo de hora a liquidar no está configurado en las reglas de horas."
        )


def periodo_precarga_usuario(
    usuario: Usuario, fecha_referencia: date,
) -> tuple[date, date]:
    contratacion = (
        usuario.tipo_contratacion.contratacion.strip().lower()
        if usuario.tipo_contratacion is not None else ""
    )
    periodos = periodos_rapidos_rrhh(fecha_referencia)
    if contratacion in {"nómina", "nomina"}:
        return periodos["nomina"]
    return periodos["monotributo"]


def periodos_vigencia_precarga(
    usuario: Usuario, fecha_desde: date, fecha_hasta: date,
) -> list[tuple[date, date]]:
    periodos: list[tuple[date, date]] = []
    cursor = fecha_desde
    while cursor <= fecha_hasta:
        periodo_desde, periodo_hasta = periodo_precarga_usuario(usuario, cursor)
        if not periodos or periodos[-1][0] != periodo_desde:
            periodos.append((periodo_desde, periodo_hasta))
        cursor = periodo_hasta + timedelta(days=1)
    return periodos


def opciones_periodos_precarga(
    fecha_referencia: date | None = None,
    meses_anteriores: int = 0,
    meses_futuros: int = 11,
) -> dict[str, list[dict[str, str]]]:
    referencia = fecha_referencia or date.today()
    inicios_actuales = {
        "nomina": periodos_rapidos_rrhh(referencia)["nomina"][0],
        "calendario": referencia.replace(day=1),
    }
    opciones: dict[str, list[dict[str, str]]] = {
        "nomina": [], "calendario": [],
    }
    for tipo, inicio_actual in inicios_actuales.items():
        for desplazamiento in range(-meses_anteriores, meses_futuros + 1):
            indice_mes = (
                inicio_actual.year * 12 + inicio_actual.month - 1 + desplazamiento
            )
            anio, indice = divmod(indice_mes, 12)
            inicio = date(anio, indice + 1, 16 if tipo == "nomina" else 1)
            _, fin = (
                periodos_rapidos_rrhh(inicio)["nomina"]
                if tipo == "nomina"
                else periodos_rapidos_rrhh(inicio)["monotributo"]
            )
            opciones[tipo].append({
                "valor": inicio.isoformat(),
                "etiqueta": (
                    f"{inicio.strftime('%d/%m/%Y')} al {fin.strftime('%d/%m/%Y')}"
                ),
            })
    return opciones


def generar_precargas_autorizadas(
    asignacion: UsuarioConceptoExcepcional,
    concepto: ConceptoExcepcional,
    usuario: Usuario,
    aprobado_por: int | None,
    db: Session,
) -> int:
    if (
        concepto.tipo != "PRECARGA_PERIODO"
        or asignacion.fecha_hasta is None
        or asignacion.cantidad_asignada is None
        or not concepto.tipo_hora_destino
    ):
        return 0
    existentes = set(db.scalars(
        select(HoraExtra.fecha).where(
            HoraExtra.usuario_id == usuario.id,
            HoraExtra.concepto_excepcional_id == concepto.id_concepto_excepcional,
            HoraExtra.tipo_registro == TIPO_PRECARGA_AUTORIZADA,
        )
    ))
    ahora = datetime.now()
    creadas = 0
    for periodo_desde, _ in periodos_vigencia_precarga(
        usuario, asignacion.fecha_desde, asignacion.fecha_hasta,
    ):
        if periodo_desde in existentes:
            continue
        db.add(HoraExtra(
            usuario_id=usuario.id,
            fecha=periodo_desde,
            hora_inicio=None,
            hora_fin=None,
            horas_totales=None,
            cantidad=asignacion.cantidad_asignada,
            concepto_excepcional_id=concepto.id_concepto_excepcional,
            concepto_excepcional_nombre=concepto.nombre,
            tipo_dia=clasificar_tipo_dia(periodo_desde, db),
            tipo_hora=concepto.tipo_hora_destino,
            horas_nocturnas=None,
            tipo_registro=TIPO_PRECARGA_AUTORIZADA,
            solicita_reintegro=False,
            observaciones=(
                f"Precarga autorizada por período. {asignacion.observacion or ''}"
            ).strip(),
            estado="APROBADA",
            aprobado_por=aprobado_por,
            fecha_resolucion=ahora,
        ))
        existentes.add(periodo_desde)
        creadas += 1
    return creadas


def generar_pack_precargas_autorizadas(
    usuario: Usuario,
    tipo_hora: str,
    cantidad: Decimal,
    fecha_desde: date,
    fecha_hasta: date,
    observacion: str,
    aprobado_por: int,
    db: Session,
) -> int:
    tipo_hora_normalizado = tipo_hora.strip().upper()
    existentes = set(db.scalars(
        select(HoraExtra.fecha).where(
            HoraExtra.usuario_id == usuario.id,
            HoraExtra.tipo_registro == TIPO_PRECARGA_AUTORIZADA,
            func.upper(HoraExtra.tipo_hora) == tipo_hora_normalizado,
        )
    ))
    ahora = datetime.now()
    creadas = 0
    for periodo_desde, _ in periodos_vigencia_precarga(
        usuario, fecha_desde, fecha_hasta,
    ):
        if periodo_desde in existentes:
            continue
        db.add(HoraExtra(
            usuario_id=usuario.id,
            fecha=periodo_desde,
            hora_inicio=None,
            hora_fin=None,
            horas_totales=None,
            cantidad=cantidad.quantize(Decimal("0.01")),
            concepto_excepcional_id=None,
            concepto_excepcional_nombre=None,
            tipo_dia=clasificar_tipo_dia(periodo_desde, db),
            tipo_hora=tipo_hora_normalizado,
            horas_nocturnas=None,
            tipo_registro=TIPO_PRECARGA_AUTORIZADA,
            solicita_reintegro=False,
            observaciones=f"Precarga autorizada por período. {observacion}".strip(),
            estado="APROBADA",
            aprobado_por=aprobado_por,
            fecha_resolucion=ahora,
        ))
        existentes.add(periodo_desde)
        creadas += 1
    return creadas


def usuario_tiene_funciones_carga(usuario: Usuario, db: Session) -> bool:
    convenio = str(usuario.convenio or "").strip().upper()
    if convenio != "FC":
        return True
    hoy = date.today()
    return db.scalar(
        select(UsuarioConceptoExcepcional.id_asignacion)
        .join(ConceptoExcepcional)
        .where(
            UsuarioConceptoExcepcional.usuario_id == usuario.id,
            UsuarioConceptoExcepcional.activo.is_(True),
            UsuarioConceptoExcepcional.fecha_desde <= hoy,
            (
                UsuarioConceptoExcepcional.fecha_hasta.is_(None)
                | (UsuarioConceptoExcepcional.fecha_hasta >= hoy)
            ),
            ConceptoExcepcional.codigo == CODIGO_PERMISO_CARGAS_FC,
            ConceptoExcepcional.tipo == "PERMISO",
            ConceptoExcepcional.activo.is_(True),
        )
        .limit(1)
    ) is not None


def ids_habilitados_para_confirmar(
    request: Request,
    db: Session,
    destino: str = "propio",
) -> set[int]:
    _, propios = usuarios_habilitados_para_carga(request, db, destino)
    return {item.id for item in propios}


def tipos_otras_cargas(
    db: Session, convenio: str | None = None,
) -> list[tuple[str, str, str, str | None, Decimal | None]]:
    consulta = (
        select(
            ReglaHora.convenio,
            ReglaHora.tipo_hora,
            ReglaHora.observaciones,
            ReglaHora.modalidad_carga,
            ReglaHora.cantidad_predeterminada,
        )
        .where(
            func.upper(ReglaHora.tipo_dia) == "TODOS",
            func.upper(ReglaHora.tipo_hora).notin_(("COMIDA", "MERIENDA", "DOMINGO")),
        )
        .distinct()
        .order_by(ReglaHora.convenio, ReglaHora.tipo_hora)
    )
    if convenio:
        consulta = consulta.where(
            func.upper(ReglaHora.convenio) == convenio.strip().upper(),
        )
    return list(db.execute(consulta).tuples())


def regla_otra_carga(
    db: Session, convenio: str, tipo_hora: str,
) -> ReglaHora | None:
    return db.scalar(
        select(ReglaHora).where(
            func.upper(ReglaHora.convenio) == convenio.strip().upper(),
            func.upper(ReglaHora.tipo_dia) == "TODOS",
            func.upper(ReglaHora.tipo_hora) == tipo_hora.strip().upper(),
        ).order_by(ReglaHora.id).limit(1)
    )


def convenio_permite_hs_articulo(db: Session, convenio: str | None) -> bool:
    if not convenio:
        return False
    return db.scalar(
        select(ReglaHora.id).where(
            func.upper(ReglaHora.convenio) == convenio.strip().upper(),
            func.upper(ReglaHora.tipo_dia) == "TODOS",
            func.upper(ReglaHora.tipo_hora) == "HS ARTICULO",
        ).limit(1)
    ) is not None


def convenio_permite_reintegro(db: Session, convenio: str | None) -> bool:
    if not convenio:
        return False
    return db.scalar(
        select(ReglaHora.id).where(
            func.upper(ReglaHora.convenio) == convenio.strip().upper(),
            ReglaHora.permite_reintegro.is_(True),
        ).limit(1)
    ) is not None


def conceptos_excepcionales_habilitados(db: Session, usuario_id: int) -> list[ConceptoExcepcional]:
    return list(db.scalars(
        select(ConceptoExcepcional)
        .join(UsuarioConceptoExcepcional)
        .where(
            UsuarioConceptoExcepcional.usuario_id == usuario_id,
            UsuarioConceptoExcepcional.activo.is_(True),
            UsuarioConceptoExcepcional.fecha_desde <= date.today(),
            (
                UsuarioConceptoExcepcional.fecha_hasta.is_(None)
                | (UsuarioConceptoExcepcional.fecha_hasta >= date.today())
            ),
            ConceptoExcepcional.activo.is_(True),
            ConceptoExcepcional.tipo.in_(("CARGA_MANUAL", "CARGA_POR_DIA")),
        )
        .distinct()
        .order_by(ConceptoExcepcional.nombre)
    ))


def franja_nocturna_convenio(db: Session, convenio: str | None) -> tuple[time, time] | None:
    if not convenio:
        return None
    franjas = list(db.execute(
        select(ReglaHora.hora_nocturna_desde, ReglaHora.hora_nocturna_hasta)
        .where(
            func.upper(ReglaHora.convenio) == convenio.strip().upper(),
            ReglaHora.hora_nocturna_desde.is_not(None),
            ReglaHora.hora_nocturna_hasta.is_not(None),
        )
        .distinct()
        .limit(2)
    ).tuples())
    return franjas[0] if len(franjas) == 1 else None


def intervalo_jornada(fecha: date, inicio: time, fin: time) -> tuple[datetime, datetime]:
    desde = datetime.combine(fecha, inicio)
    hasta = datetime.combine(fecha, fin)
    if hasta <= desde:
        hasta += timedelta(days=1)
    return desde, hasta


def validar_nocturnidad_sin_superposicion(
    usuario_id: int,
    fecha: date,
    hora_inicio: time,
    hora_fin: time,
    db: Session,
    resultados_borrador: list | None = None,
    excluir_hora_id: int | None = None,
) -> None:
    desde, hasta = intervalo_jornada(fecha, hora_inicio, hora_fin)
    candidatos = list(resultados_borrador or [])
    consulta_existentes = select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
            HoraExtra.hora_inicio.is_not(None),
            HoraExtra.hora_fin.is_not(None),
            HoraExtra.fecha >= desde.date() - timedelta(days=1),
            HoraExtra.fecha <= hasta.date(),
        )
    if excluir_hora_id is not None:
        consulta_existentes = consulta_existentes.where(HoraExtra.id != excluir_hora_id)
    existentes = db.scalars(consulta_existentes)
    candidatos.extend(existentes)
    for item in candidatos:
        if item.hora_inicio is None or item.hora_fin is None:
            continue
        item_desde, item_hasta = intervalo_jornada(
            item.fecha, item.hora_inicio, item.hora_fin,
        )
        if desde < item_hasta and item_desde < hasta:
            raise CalculoHorasError(
                "El horario nocturno se superpone con otra carga pendiente, autorizada o incluida en este borrador."
            )


def fechas_feriados_sin_devolucion(db: Session) -> list[str]:
    return [
        fecha.isoformat()
        for fecha in db.scalars(
            select(Feriado.fecha)
            .where(Feriado.devuelve.is_(False))
            .order_by(Feriado.fecha)
        )
    ]


def fechas_feriados(db: Session) -> list[str]:
    return [
        fecha.isoformat()
        for fecha in db.scalars(select(Feriado.fecha).order_by(Feriado.fecha))
    ]


CONCEPTOS_DERIVADOS = {"MERIENDA", "COMIDA"}
CONCEPTOS_AUTOMATICOS = CONCEPTOS_DERIVADOS | {"DOMINGO", CODIGO_SWITCHER_NOCTURNIDAD}


def es_concepto_derivado(hora: HoraExtra) -> bool:
    return (
        hora.tipo_registro == "OTRAS"
        and str(hora.tipo_hora or "").upper() in CONCEPTOS_DERIVADOS
    )


def es_concepto_automatico(hora: HoraExtra) -> bool:
    if (
        hora.tipo_registro == "EXCEPCIONAL"
        and str(hora.tipo_hora or "").upper() == CODIGO_SWITCHER_NOCTURNIDAD
    ):
        return True
    return (
        hora.tipo_registro == "OTRAS"
        and str(hora.tipo_hora or "").upper() in CONCEPTOS_AUTOMATICOS
        and not (
            str(hora.tipo_hora or "").upper() == "DOMINGO"
            and hora.hora_inicio is not None
            and hora.hora_fin is not None
        )
    )


def cantidad_nocturna_registro(hora: HoraExtra) -> Decimal:
    if es_concepto_automatico(hora):
        return Decimal("0.00")
    if (
        hora.tipo_registro == "OTRAS"
        and str(hora.tipo_hora or "").strip().upper()
        in TIPOS_NOCTURNAS_EN_JORNADA
    ):
        return hora.cantidad or Decimal("0.00")
    return hora.horas_nocturnas or Decimal("0.00")


def concepto_switcher_vigente(
    db: Session, usuario_id: int, fecha: date,
) -> ConceptoExcepcional | None:
    return db.scalar(
        select(ConceptoExcepcional)
        .join(UsuarioConceptoExcepcional)
        .where(
            UsuarioConceptoExcepcional.usuario_id == usuario_id,
            UsuarioConceptoExcepcional.activo.is_(True),
            UsuarioConceptoExcepcional.fecha_desde <= fecha,
            (
                UsuarioConceptoExcepcional.fecha_hasta.is_(None)
                | (UsuarioConceptoExcepcional.fecha_hasta >= fecha)
            ),
            ConceptoExcepcional.activo.is_(True),
            func.upper(ConceptoExcepcional.codigo) == CODIGO_SWITCHER_NOCTURNIDAD,
            ConceptoExcepcional.tipo == TIPO_SWITCHER_NOCTURNIDAD,
        )
        .order_by(UsuarioConceptoExcepcional.fecha_desde.desc())
        .limit(1)
    )


def recalcular_switcher_nocturnidad_estado(
    usuario_id: int,
    fecha: date,
    db: Session,
    estado: str,
    aprobado_por: int | None = None,
    fecha_resolucion: datetime | None = None,
) -> None:
    concepto = concepto_switcher_vigente(db, usuario_id, fecha)
    existentes = list(db.scalars(
        select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.fecha == fecha,
            HoraExtra.estado == estado,
            HoraExtra.tipo_registro == "EXCEPCIONAL",
            func.upper(HoraExtra.tipo_hora) == CODIGO_SWITCHER_NOCTURNIDAD,
        ).order_by(HoraExtra.id)
    ))
    fuentes = list(db.scalars(
        select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.fecha == fecha,
            HoraExtra.estado == estado,
        )
    ))
    cantidad = sum(
        (cantidad_nocturna_registro(item) for item in fuentes),
        start=Decimal("0.00"),
    ).quantize(Decimal("0.01"))
    if concepto is None or cantidad <= 0:
        for registro in existentes:
            db.delete(registro)
        return

    tipo_dia_fuente = next(
        (
            item.tipo_dia for item in fuentes
            if cantidad_nocturna_registro(item) > 0
        ),
        clasificar_tipo_dia(fecha, db),
    )

    registro = existentes[0] if existentes else HoraExtra(
        usuario_id=usuario_id,
        fecha=fecha,
        tipo_dia=tipo_dia_fuente,
        tipo_hora=CODIGO_SWITCHER_NOCTURNIDAD,
        tipo_registro="EXCEPCIONAL",
        estado=estado,
        fecha_carga=datetime.now(),
        solicita_reintegro=False,
    )
    if not existentes:
        db.add(registro)
    registro.concepto_excepcional_id = concepto.id_concepto_excepcional
    registro.concepto_excepcional_nombre = concepto.nombre
    registro.tipo_dia = tipo_dia_fuente
    registro.cantidad = cantidad
    registro.hora_inicio = None
    registro.hora_fin = None
    registro.horas_totales = None
    registro.horas_nocturnas = None
    registro.observaciones = f"Cálculo automático sobre {cantidad:.2f} horas nocturnas."
    if estado == "APROBADA":
        registro.aprobado_por = aprobado_por
        registro.fecha_resolucion = fecha_resolucion
    for duplicado in existentes[1:]:
        db.delete(duplicado)


def es_domingo_sat(usuario: Usuario, fecha: date) -> bool:
    return str(usuario.convenio or "").strip().upper() == "SAT" and fecha.weekday() == 6


def existe_domingo_activo_o_en_borrador(
    usuario_id: int,
    fecha: date,
    resultados: list,
    db: Session,
) -> bool:
    if any(
        item.usuario_id == usuario_id
        and item.fecha == fecha
        and item.tipo_registro == "OTRAS"
        and str(item.tipo_hora or "").upper() == "DOMINGO"
        for item in resultados
    ):
        return True
    return db.scalar(
        select(HoraExtra.id).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.fecha == fecha,
            HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
            func.upper(HoraExtra.tipo_hora) == "DOMINGO",
        ).limit(1)
    ) is not None


def fecha_jornada_de_hora_estado(
    hora: HoraExtra,
    db: Session,
    estado: str,
) -> date:
    """Asigna una continuación 00:00 a la jornada del día anterior."""
    if hora.tipo_registro != "HORAS" or hora.hora_inicio != time(0, 0):
        return hora.fecha
    fecha_anterior = hora.fecha - timedelta(days=1)
    tramo_anterior = db.scalar(
        select(HoraExtra.id).where(
            HoraExtra.usuario_id == hora.usuario_id,
            HoraExtra.estado == estado,
            HoraExtra.tipo_registro == "HORAS",
            HoraExtra.fecha == fecha_anterior,
            HoraExtra.hora_fin == time(0, 0),
        ).limit(1)
    )
    return fecha_anterior if tramo_anterior is not None else hora.fecha


def fecha_jornada_de_hora_pendiente(hora: HoraExtra, db: Session) -> date:
    return fecha_jornada_de_hora_estado(hora, db, "PENDIENTE")


def horas_de_jornada_estado(
    usuario_id: int,
    fecha_jornada: date,
    db: Session,
    estado: str,
) -> list[HoraExtra]:
    fecha_siguiente = fecha_jornada + timedelta(days=1)
    candidatas = list(db.scalars(
        select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.estado == estado,
            HoraExtra.tipo_registro == "HORAS",
            HoraExtra.fecha.in_((fecha_jornada, fecha_siguiente)),
        )
    ))
    hay_cierre_en_medianoche = any(
        item.fecha == fecha_jornada and item.hora_fin == time(0, 0)
        for item in candidatas
    )
    hay_cierre_dia_anterior = db.scalar(
        select(HoraExtra.id).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.estado == estado,
            HoraExtra.tipo_registro == "HORAS",
            HoraExtra.fecha == fecha_jornada - timedelta(days=1),
            HoraExtra.hora_fin == time(0, 0),
        ).limit(1)
    ) is not None
    return [
        item for item in candidatas
        if (
            item.fecha == fecha_jornada
            and not (hay_cierre_dia_anterior and item.hora_inicio == time(0, 0))
        ) or (
            item.fecha == fecha_siguiente
            and hay_cierre_en_medianoche
            and item.hora_inicio == time(0, 0)
        )
    ]


def horas_pendientes_de_jornada(
    usuario_id: int,
    fecha_jornada: date,
    db: Session,
) -> list[HoraExtra]:
    return horas_de_jornada_estado(usuario_id, fecha_jornada, db, "PENDIENTE")


def recalcular_conceptos_jornada_pendiente(
    usuario_id: int,
    fecha_jornada: date,
    db: Session,
) -> None:
    recalcular_conceptos_jornada_estado(
        usuario_id, fecha_jornada, db, "PENDIENTE",
    )


def recalcular_conceptos_jornada_estado(
    usuario_id: int,
    fecha_jornada: date,
    db: Session,
    estado: str,
    aprobado_por: int | None = None,
    fecha_resolucion: datetime | None = None,
) -> None:
    horas = horas_de_jornada_estado(usuario_id, fecha_jornada, db, estado)
    total = sum(
        (item.horas_totales or Decimal("0.00") for item in horas),
        start=Decimal("0.00"),
    )
    tipo_dia_jornada = next(
        (item.tipo_dia for item in horas if item.fecha == fecha_jornada),
        horas[0].tipo_dia if horas else "HABIL",
    )
    cantidades = {
        "MERIENDA": int(total // Decimal("2")),
        "COMIDA": int(total // Decimal("3")),
    }
    existentes = list(db.scalars(
        select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.fecha == fecha_jornada,
            HoraExtra.estado == estado,
            HoraExtra.tipo_registro == "OTRAS",
            func.upper(HoraExtra.tipo_hora).in_(CONCEPTOS_DERIVADOS),
        ).order_by(HoraExtra.id)
    ))
    por_tipo: dict[str, list[HoraExtra]] = {"MERIENDA": [], "COMIDA": []}
    for item in existentes:
        por_tipo[str(item.tipo_hora).upper()].append(item)

    usuario = db.get(Usuario, usuario_id)
    if usuario is None:
        raise CalculoHorasError("El usuario de la jornada no existe.")
    reglas_automaticas = list(db.scalars(
        select(ReglaHora).where(
            func.upper(ReglaHora.convenio) == str(usuario.convenio or "").upper(),
            func.upper(ReglaHora.tipo_dia) == "TODOS",
            func.upper(ReglaHora.tipo_hora).in_(CONCEPTOS_DERIVADOS),
        ).order_by(ReglaHora.id)
    ))
    reglas_por_concepto: dict[str, ReglaHora] = {}
    for regla_automatica in reglas_automaticas:
        reglas_por_concepto.setdefault(
            str(regla_automatica.tipo_hora).upper(), regla_automatica,
        )
    for concepto, cantidad in cantidades.items():
        registros = por_tipo[concepto]
        regla = reglas_por_concepto.get(concepto)
        if regla is None or cantidad <= 0:
            for registro in registros:
                db.delete(registro)
            continue
        registro = registros[0] if registros else HoraExtra(
            usuario_id=usuario_id,
            fecha=fecha_jornada,
            tipo_dia="TODOS",
            tipo_hora=regla.tipo_hora,
            tipo_registro="OTRAS",
            estado=estado,
            aprobado_por=aprobado_por if estado == "APROBADA" else None,
            fecha_resolucion=fecha_resolucion if estado == "APROBADA" else None,
            fecha_carga=datetime.now(),
            solicita_reintegro=False,
        )
        if not registros:
            db.add(registro)
        registro.cantidad = Decimal(cantidad).quantize(Decimal("0.01"))
        registro.tipo_dia = tipo_dia_jornada
        registro.hora_inicio = None
        registro.hora_fin = None
        registro.horas_totales = None
        registro.horas_nocturnas = None
        registro.observaciones = f"Cálculo automático sobre {total:.2f} horas extras."
        if estado == "APROBADA":
            registro.aprobado_por = aprobado_por
            registro.fecha_resolucion = fecha_resolucion
        for duplicado in registros[1:]:
            db.delete(duplicado)


def conceptos_automaticos_borrador(
    resultados: list,
    db: Session,
) -> tuple[list[dict], list[str | None], list[str | None]]:
    """Proyecta conceptos derivados sin incorporarlos a las cargas confirmables."""
    cierres_medianoche = {
        (item.usuario_id, item.fecha)
        for item in resultados
        if item.tipo_registro == "HORAS" and item.hora_fin == time(0, 0)
    }
    claves: list[str | None] = []
    claves_nocturnidad: list[str | None] = []
    grupos: dict[tuple[int, date], dict] = {}
    grupos_nocturnidad: dict[tuple[int, date], dict] = {}
    for item in resultados:
        cantidad_nocturna = cantidad_nocturna_registro(item)
        clave_nocturna = None
        if cantidad_nocturna > 0:
            clave_nocturna = f"nocturna:{item.usuario_id}:{item.fecha.isoformat()}"
            grupo_nocturno = grupos_nocturnidad.setdefault(
                (item.usuario_id, item.fecha),
                {"muestra": item, "cantidad_borrador": Decimal("0.00"), "clave": clave_nocturna},
            )
            grupo_nocturno["cantidad_borrador"] += cantidad_nocturna
        claves_nocturnidad.append(clave_nocturna)
        if item.tipo_registro != "HORAS":
            claves.append(None)
            continue
        fecha_jornada = item.fecha
        if (
            item.hora_inicio == time(0, 0)
            and (item.usuario_id, item.fecha - timedelta(days=1))
            in cierres_medianoche
        ):
            fecha_jornada -= timedelta(days=1)
        clave_html = f"{item.usuario_id}:{fecha_jornada.isoformat()}"
        claves.append(clave_html)
        grupo = grupos.setdefault((item.usuario_id, fecha_jornada), {
            "muestra": item,
            "horas_borrador": Decimal("0.00"),
            "clave": clave_html,
        })
        grupo["horas_borrador"] += item.horas_totales or Decimal("0.00")

    convenios = {
        str(grupo["muestra"].convenio or "").strip().upper()
        for grupo in grupos.values()
    }
    reglas = list(db.scalars(
        select(ReglaHora).where(
            func.upper(ReglaHora.convenio).in_(convenios),
            func.upper(ReglaHora.tipo_dia) == "TODOS",
            func.upper(ReglaHora.tipo_hora).in_(CONCEPTOS_DERIVADOS),
        )
    ))
    habilitados = {
        (str(regla.convenio).strip().upper(), str(regla.tipo_hora).strip().upper())
        for regla in reglas
    }
    conceptos = []
    for (usuario_id, fecha_jornada), grupo in grupos.items():
        muestra = grupo["muestra"]
        convenio = str(muestra.convenio or "").strip().upper()
        horas_existentes = sum(
            (
                item.horas_totales or Decimal("0.00")
                for item in horas_pendientes_de_jornada(
                    usuario_id, fecha_jornada, db,
                )
            ),
            start=Decimal("0.00"),
        )
        total = horas_existentes + grupo["horas_borrador"]
        for concepto, divisor in (("MERIENDA", 2), ("COMIDA", 3)):
            cantidad = int(total // Decimal(divisor))
            if (convenio, concepto) not in habilitados or cantidad <= 0:
                continue
            conceptos.append({
                "usuario_nombre": muestra.usuario_nombre,
                "legajo": muestra.legajo,
                "fecha": fecha_jornada,
                "tipo_dia": muestra.tipo_dia,
                "tipo_hora": concepto,
                "cantidad": cantidad,
                "clave": grupo["clave"],
                "divisor": divisor,
                "horas_existentes": horas_existentes,
                "observaciones": f"Cálculo automático sobre {total:.2f} horas extras.",
                "modo": "DIVISOR",
            })
    for (usuario_id, fecha_carga), grupo in grupos_nocturnidad.items():
        concepto = concepto_switcher_vigente(db, usuario_id, fecha_carga)
        if concepto is None:
            continue
        fuentes_existentes = db.scalars(select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.fecha == fecha_carga,
            HoraExtra.estado == "PENDIENTE",
        ))
        cantidad_existente = sum(
            (cantidad_nocturna_registro(item) for item in fuentes_existentes),
            start=Decimal("0.00"),
        )
        total = (cantidad_existente + grupo["cantidad_borrador"]).quantize(
            Decimal("0.01"),
        )
        if total <= 0:
            continue
        muestra = grupo["muestra"]
        conceptos.append({
            "usuario_nombre": muestra.usuario_nombre,
            "legajo": muestra.legajo,
            "fecha": fecha_carga,
            "tipo_dia": muestra.tipo_dia,
            "tipo_hora": concepto.nombre,
            "cantidad": total,
            "clave": grupo["clave"],
            "divisor": None,
            "horas_existentes": cantidad_existente,
            "observaciones": f"Cálculo automático sobre {total:.2f} horas nocturnas.",
            "modo": "NOCTURNIDAD",
        })
    return conceptos, claves, claves_nocturnidad


def expandir_ids_con_jornadas_pendientes(
    ids: list[int],
    db: Session,
) -> list[int]:
    seleccionadas = list(db.scalars(
        select(HoraExtra).where(
            HoraExtra.id.in_(ids),
            HoraExtra.estado == "PENDIENTE",
        )
    ))
    jornadas: set[tuple[int, date]] = set()
    fechas_nocturnas: set[tuple[int, date]] = set()
    for item in seleccionadas:
        if item.tipo_registro == "HORAS":
            jornadas.add((item.usuario_id, fecha_jornada_de_hora_pendiente(item, db)))
        elif es_concepto_automatico(item):
            jornadas.add((item.usuario_id, item.fecha))
        if cantidad_nocturna_registro(item) > 0 or (
            item.tipo_registro == "EXCEPCIONAL"
            and str(item.tipo_hora or "").upper() == CODIGO_SWITCHER_NOCTURNIDAD
        ):
            fechas_nocturnas.add((item.usuario_id, item.fecha))
    expandidos = set(ids)
    for usuario_id, fecha_jornada in jornadas:
        expandidos.update(item.id for item in horas_pendientes_de_jornada(
            usuario_id, fecha_jornada, db,
        ))
        expandidos.update(db.scalars(
            select(HoraExtra.id).where(
                HoraExtra.usuario_id == usuario_id,
                HoraExtra.fecha == fecha_jornada,
                HoraExtra.estado == "PENDIENTE",
                HoraExtra.tipo_registro == "OTRAS",
                func.upper(HoraExtra.tipo_hora).in_(CONCEPTOS_AUTOMATICOS),
            )
        ))
        expandidos.update(db.scalars(
            select(HoraExtra.id).where(
                HoraExtra.usuario_id == usuario_id,
                HoraExtra.fecha == fecha_jornada,
                HoraExtra.estado == "PENDIENTE",
                HoraExtra.tipo_registro == "EXCEPCIONAL",
                func.upper(HoraExtra.tipo_hora) == CODIGO_SWITCHER_NOCTURNIDAD,
            )
        ))
    for usuario_id, fecha in fechas_nocturnas:
        candidatos = db.scalars(select(HoraExtra).where(
            HoraExtra.usuario_id == usuario_id,
            HoraExtra.fecha == fecha,
            HoraExtra.estado == "PENDIENTE",
        ))
        expandidos.update(
            item.id for item in candidatos
            if cantidad_nocturna_registro(item) > 0 or (
                item.tipo_registro == "EXCEPCIONAL"
                and str(item.tipo_hora or "").upper() == CODIGO_SWITCHER_NOCTURNIDAD
            )
        )
    return sorted(expandidos)


def encode_carga(
    usuario_id: int,
    fecha: date,
    hora_inicio: time | None,
    hora_fin: time | None,
    observaciones: str | None,
    tipo_registro: str = "HORAS",
    marcar_como_franco: bool = False,
    tipo_hora: str | None = None,
    cantidad: Decimal | None = None,
    concepto_excepcional_id: int | None = None,
    concepto_excepcional_nombre: str | None = None,
    fecha_descanso_compensatorio: date | None = None,
) -> str:
    return "|".join((
        str(usuario_id),
        fecha.isoformat(),
        hora_inicio.strftime("%H:%M") if hora_inicio else "",
        hora_fin.strftime("%H:%M") if hora_fin else "",
        quote(clean_optional(observaciones) or "", safe=""),
        tipo_registro,
        "SI" if marcar_como_franco else "NO",
        quote(clean_optional(tipo_hora) or "", safe=""),
        str(cantidad) if cantidad is not None else "",
        str(concepto_excepcional_id) if concepto_excepcional_id is not None else "",
        quote(clean_optional(concepto_excepcional_nombre) or "", safe=""),
        fecha_descanso_compensatorio.isoformat() if fecha_descanso_compensatorio else "",
    ))


def decode_carga(value: str) -> tuple[int, date, time | None, time | None, str | None, str, bool, str | None, Decimal | None, int | None, str | None, date | None]:
    try:
        partes = value.split("|", maxsplit=11)
        if len(partes) == 5:
            partes.append("HORAS")
        if len(partes) == 6:
            partes.append("NO")
        while len(partes) < 12:
            partes.append("")
        usuario_id, fecha, hora_inicio, hora_fin, observaciones, tipo_registro, marcar_franco, tipo_hora, cantidad, concepto_id, concepto_nombre, fecha_descanso = partes
        tipo_registro = tipo_registro.strip().upper()
        marcar_franco = marcar_franco.strip().upper()
        if tipo_registro not in {"HORAS", "DIA_TRABAJADO", "REINTEGRO", "OTRAS", "EXCEPCIONAL"} or marcar_franco not in {"SI", "NO"}:
            raise ValueError
        return (
            int(usuario_id),
            date.fromisoformat(fecha),
            time.fromisoformat(hora_inicio) if hora_inicio else None,
            time.fromisoformat(hora_fin) if hora_fin else None,
            clean_optional(unquote(observaciones)),
            tipo_registro,
            marcar_franco == "SI",
            clean_optional(unquote(tipo_hora)),
            Decimal(cantidad) if cantidad else None,
            int(concepto_id) if concepto_id else None,
            clean_optional(unquote(concepto_nombre)),
            date.fromisoformat(fecha_descanso) if fecha_descanso else None,
        )
    except (TypeError, ValueError) as error:
        raise CalculoHorasError("La solicitud contiene una carga con formato inválido.") from error


def calcular_carga_codificada(
    datos,
    db: Session,
    permitir_fc_sin_asignacion: bool = False,
):
    usuario_id, fecha, hora_inicio, hora_fin, _, tipo_registro, marcar_como_franco, tipo_hora, cantidad, concepto_id, concepto_nombre, _ = datos
    usuario = db.get(Usuario, usuario_id)
    convenio = str(usuario.convenio or "").strip().upper() if usuario else ""
    if (
        convenio == "FC"
        and usuario is not None
        and not permitir_fc_sin_asignacion
        and not usuario_tiene_funciones_carga(usuario, db)
    ):
        raise CalculoHorasError("El usuario FC no tiene habilitada la excepción de carga.")
    if convenio == "MONOTRIBUTISTA" and tipo_registro == "REINTEGRO":
        raise CalculoHorasError(
            "MONOTRIBUTISTA no permite solicitar reintegro."
        )
    if tipo_registro == "REINTEGRO":
        return calcular_resultado_reintegro(
            usuario_id, fecha, db, cantidad or Decimal("1.00"),
        )
    if tipo_registro == "DIA_TRABAJADO":
        return calcular_resultado_dia_trabajado(
            usuario_id, fecha, db, marcar_como_franco=marcar_como_franco,
            hora_inicio=hora_inicio, hora_fin=hora_fin,
        )
    if tipo_registro == "OTRAS":
        if tipo_hora is None or cantidad is None:
            raise CalculoHorasError("La otra carga no contiene un tipo y una cantidad válidos.")
        if tipo_hora.strip().upper() in TIPOS_NOCTURNAS_EN_JORNADA:
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("La carga nocturna no contiene un horario válido.")
            return calcular_resultado_nocturnas_en_jornada(
                usuario_id, fecha, hora_inicio, hora_fin, db,
            )
        if tipo_hora.strip().upper() == "DOMINGO":
            return calcular_resultado_domingo(
                usuario_id, fecha, db, hora_inicio=hora_inicio, hora_fin=hora_fin,
            )
        return calcular_resultado_otra_carga(usuario_id, fecha, tipo_hora, cantidad, db)
    if tipo_registro == "EXCEPCIONAL":
        if concepto_id is None or cantidad is None or concepto_nombre is None:
            raise CalculoHorasError("El concepto excepcional está incompleto.")
        return calcular_resultado_concepto_excepcional(
            usuario_id, fecha, concepto_id, cantidad, db,
            nombre_historico=concepto_nombre,
        )
    if hora_inicio is None or hora_fin is None:
        raise CalculoHorasError("La carga de horas no contiene un horario válido.")
    return calcular_resultado_horas_extra(
        usuario_id, fecha, hora_inicio, hora_fin, db,
        marcar_como_franco=marcar_como_franco,
    )


def intervalo_resultado(resultado) -> tuple[datetime, datetime] | None:
    if resultado.hora_inicio is None or resultado.hora_fin is None:
        return None
    inicio = datetime.combine(resultado.fecha, resultado.hora_inicio)
    fin = datetime.combine(resultado.fecha, resultado.hora_fin)
    if fin <= inicio:
        fin += timedelta(days=1)
    return inicio, fin


def evitar_nocturnidad_duplicada_en_jornada(resultados: list) -> None:
    """Deja la nocturnidad en la jornada completa, no en sus horas extra incluidas."""
    jornadas_por_usuario: dict[int, list[tuple[datetime, datetime]]] = {}
    for resultado in resultados:
        es_domingo = (
            resultado.tipo_registro == "OTRAS"
            and str(resultado.tipo_hora or "").strip().upper() == "DOMINGO"
        )
        if resultado.tipo_registro != "DIA_TRABAJADO" and not es_domingo:
            continue
        intervalo = intervalo_resultado(resultado)
        if intervalo is not None:
            jornadas_por_usuario.setdefault(resultado.usuario_id, []).append(intervalo)

    for indice, resultado in enumerate(resultados):
        if resultado.tipo_registro != "HORAS":
            continue
        intervalo_horas = intervalo_resultado(resultado)
        if intervalo_horas is None:
            continue
        inicio_horas, fin_horas = intervalo_horas
        if any(
            inicio_jornada <= inicio_horas and fin_horas <= fin_jornada
            for inicio_jornada, fin_jornada
            in jornadas_por_usuario.get(resultado.usuario_id, [])
        ):
            resultados[indice] = replace(resultado, horas_nocturnas=None)


def periodo_corte(fecha_referencia: date) -> tuple[date, date]:
    corte_actual = date(fecha_referencia.year, fecha_referencia.month, 20)
    if fecha_referencia.day > 20:
        desde = corte_actual
        if fecha_referencia.month == 12:
            hasta = date(fecha_referencia.year + 1, 1, 20)
        else:
            hasta = date(fecha_referencia.year, fecha_referencia.month + 1, 20)
    else:
        hasta = corte_actual
        if fecha_referencia.month == 1:
            desde = date(fecha_referencia.year - 1, 12, 20)
        else:
            desde = date(fecha_referencia.year, fecha_referencia.month - 1, 20)
    return desde, hasta


def obtener_autorizador(request: Request, db: Session) -> tuple[Usuario, UsuarioRol, bool]:
    session_data = request.session.get("user") or {}
    asignacion = db.scalar(
        select(UsuarioRol).where(
            UsuarioRol.id_rol == session_data.get("active_assignment_id"),
            UsuarioRol.usuario_id == session_data.get("id"),
        )
    )
    usuario = db.scalar(
        select(Usuario)
        .options(selectinload(Usuario.roles))
        .where(Usuario.id == session_data.get("id"))
    )
    es_admin = asignacion is not None and asignacion.rol.upper() == "ADMIN"
    if usuario is None or asignacion is None or not usuario.status or (
        not es_admin and asignacion.perfil.upper() != "JEFE"
    ):
        raise HTTPException(
            status_code=403,
            detail="No tenés permisos para acceder a las autorizaciones.",
        )
    return usuario, asignacion, es_admin


def validar_hora_para_autorizador(
    hora: HoraExtra | None,
    asignacion: UsuarioRol,
    es_admin: bool,
    requiere_pendiente: bool = True,
) -> HoraExtra:
    if hora is None:
        raise HTTPException(status_code=404, detail="La carga de horas no existe.")
    roles_usuario = {item.rol.upper() for item in hora.usuario.roles}
    if not es_admin and asignacion.rol.upper() not in roles_usuario:
        raise HTTPException(
            status_code=403,
            detail="No podés administrar horas de usuarios con otro rol.",
        )
    if requiere_pendiente and hora.estado != "PENDIENTE":
        raise HTTPException(
            status_code=409,
            detail="La carga ya no está pendiente de autorización.",
        )
    return hora


@app.get("/login", response_class=HTMLResponse)
def login_page(request: Request, next: str | None = None, error: str | None = None):
    if request.session.get("user"):
        return RedirectResponse(safe_next_url(next), status_code=303)
    return templates.TemplateResponse(request, "login.html", {
        "active_page": "login",
        "next_url": safe_next_url(next),
        "error": error,
    })


@app.post("/login")
def login(
    request: Request,
    username: Annotated[str, Form()],
    password: Annotated[str, Form()],
    next_url: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    redirect_to = safe_next_url(next_url)
    normalized_username = normalize_username(username)
    if is_local_username(normalized_username):
        local_user = authenticate_local_user(db, normalized_username, password)
        if local_user is None:
            return templates.TemplateResponse(
                request,
                "login.html",
                {
                    "active_page": "login",
                    "next_url": redirect_to,
                    "error": "Usuario o contraseña incorrectos.",
                    "username": normalized_username,
                },
                status_code=401,
            )
        request.session["user"] = session_user(local_user, "LOCAL")
        return RedirectResponse(redirect_to, status_code=303)

    try:
        authenticate_ad_user(normalized_username, password)
    except ActiveDirectoryAuthError as error:
        return templates.TemplateResponse(
            request,
            "login.html",
            {
                "active_page": "login",
                "next_url": redirect_to,
                "error": str(error),
                "username": normalized_username,
            },
            status_code=401,
        )

    database_user = find_user(db, normalized_username)
    if database_user is None or not database_user.status:
        return templates.TemplateResponse(
            request,
            "login.html",
            {
                "active_page": "login",
                "next_url": redirect_to,
                "error": "El usuario no está habilitado en la aplicación.",
                "username": normalized_username,
            },
            status_code=403,
        )
    request.session["user"] = session_user(database_user, "AD")
    return RedirectResponse(redirect_to, status_code=303)


@app.post("/logout")
def logout(request: Request):
    request.session.clear()
    return RedirectResponse("/login", status_code=303)


@app.post("/sesion/asignacion")
def cambiar_asignacion(
    request: Request,
    asignacion_id: Annotated[int, Form()],
    db: Session = Depends(get_db),
):
    session_data = request.session.get("user")
    if not session_data:
        return RedirectResponse("/login", status_code=303)
    assignment = db.scalar(
        select(UsuarioRol).where(
            UsuarioRol.id_rol == asignacion_id,
            UsuarioRol.usuario_id == session_data.get("id"),
        )
    )
    if assignment is None:
        raise HTTPException(status_code=403, detail="La asignación no pertenece al usuario activo.")
    session_data["active_assignment_id"] = assignment.id_rol
    session_data["role"] = assignment.rol
    session_data["profile"] = assignment.perfil
    request.session["user"] = session_data
    return RedirectResponse("/", status_code=303)


@app.get("/", response_class=HTMLResponse)
def inicio(request: Request, db: Session = Depends(get_db)):
    assignment = obtener_asignacion_activa(request, db)
    usuario = db.get(Usuario, assignment.usuario_id)
    if usuario is None:
        raise HTTPException(status_code=403, detail="El usuario activo no existe.")
    tutorial_key = tutorial_key_for_assignment(assignment, usuario)
    es_admin = assignment.rol.upper() == "ADMIN"
    es_jefe = assignment.perfil.upper() == "JEFE"
    sin_funciones = (
        assignment.perfil.upper() == "USUARIO"
        and not usuario_tiene_funciones_carga(usuario, db)
    )
    pendientes = 0 if sin_funciones else db.scalar(
        select(func.count(HoraExtra.id)).where(
            HoraExtra.usuario_id == assignment.usuario_id,
            HoraExtra.estado == "PENDIENTE",
        )
    ) or 0
    autorizaciones_pendientes = 0
    if es_admin:
        autorizaciones_pendientes = db.scalar(
            select(func.count(HoraExtra.id)).where(HoraExtra.estado == "PENDIENTE")
        ) or 0
    elif es_jefe:
        autorizaciones_pendientes = db.scalar(
            select(func.count(HoraExtra.id))
            .join(HoraExtra.usuario)
            .where(
                HoraExtra.estado == "PENDIENTE",
                Usuario.roles.any(
                    func.upper(UsuarioRol.rol) == assignment.rol.upper()
                ),
            )
        ) or 0
    return templates.TemplateResponse(request, "inicio.html", {
        "active_page": "inicio",
        "es_jefe": es_jefe,
        "es_admin": es_admin,
        "rol_activo": assignment.rol,
        "sin_funciones": sin_funciones,
        "cargas_pendientes": pendientes,
        "autorizaciones_pendientes": autorizaciones_pendientes,
        "tutorial_label": TUTORIAL_LABELS[tutorial_key],
        "reglas_help_label": (
            "Conocé las reglas de tu convenio"
            if tutorial_key in CONVENIOS_DISPONIBLES
            else "Consultá las reglas por convenio"
        ),
    })


@app.get("/ayuda/tutorial", response_class=HTMLResponse)
def tutorial_interactivo(request: Request, db: Session = Depends(get_db)):
    assignment = obtener_asignacion_activa(request, db)
    usuario = db.get(Usuario, assignment.usuario_id)
    if usuario is None or not usuario.status:
        raise HTTPException(status_code=403, detail="El usuario activo no está habilitado.")
    tutorial_key = tutorial_key_for_assignment(assignment, usuario)
    return HTMLResponse(
        content=tutorial_html_for_key(tutorial_key),
        headers={"Cache-Control": "private, no-store"},
    )


@app.get("/ayuda/reglas", response_class=HTMLResponse)
def reglas_interactivas(request: Request, db: Session = Depends(get_db)):
    assignment = obtener_asignacion_activa(request, db)
    usuario = db.get(Usuario, assignment.usuario_id)
    if usuario is None or not usuario.status:
        raise HTTPException(status_code=403, detail="El usuario activo no está habilitado.")
    convenios = convenios_visibles_para_reglas(assignment, usuario, db)
    reglas = list(db.scalars(
        select(ReglaHora)
        .where(func.upper(ReglaHora.convenio).in_(convenios))
        .order_by(ReglaHora.convenio, ReglaHora.tipo_dia, ReglaHora.tipo_hora)
    )) if convenios else []
    grupos_por_convenio: dict[str, dict[str, dict]] = {
        convenio: {} for convenio in convenios
    }
    for regla in reglas:
        convenio = str(regla.convenio or "").strip().upper()
        if convenio not in grupos_por_convenio:
            continue
        concepto = descripcion_concepto_regla(regla.tipo_hora)
        tipo_dia = str(regla.tipo_dia or "").strip().upper().replace("Á", "A")
        if convenio == "SAT" and str(regla.tipo_hora or "").strip().upper() == "DOMINGO":
            tipo_dia = "DOMINGO"
        nombre_dia, icono_dia, descripcion_dia = descripcion_tipo_dia(
            convenio, tipo_dia,
        )
        grupo = grupos_por_convenio[convenio].setdefault(tipo_dia, {
            "codigo": tipo_dia,
            "nombre": nombre_dia,
            "icono": icono_dia,
            "descripcion": descripcion_dia,
            "conceptos": [],
        })
        grupo["conceptos"].append({
            "tipo_hora": str(regla.tipo_hora or "").strip(),
            "nombre": concepto["nombre"],
            "categoria": concepto["categoria"],
            "descripcion": concepto["descripcion"],
            "icono": concepto["icono"],
            "nocturna_desde": regla.hora_nocturna_desde.strftime("%H:%M") if regla.hora_nocturna_desde else None,
            "nocturna_hasta": regla.hora_nocturna_hasta.strftime("%H:%M") if regla.hora_nocturna_hasta else None,
            "permite_reintegro": bool(regla.permite_reintegro),
            "observaciones": regla.observaciones,
            "modalidad": descripcion_modalidad_regla(
                regla.modalidad_carga, regla.cantidad_predeterminada,
            ),
        })
    for convenio, grupos in grupos_por_convenio.items():
        for tipo_dia in ("FRANCO", "FERIADO"):
            if tipo_dia not in grupos:
                continue
            grupos[tipo_dia]["conceptos"].insert(0, {
                "tipo_hora": "DIA_TRABAJADO",
                "nombre": "Día trabajado completo",
                "categoria": "Día trabajado",
                "descripcion": "Se registra con cantidad 1; las horas extra, si corresponden, quedan en otra fila.",
                "icono": "bi-calendar-check",
                "nocturna_desde": None,
                "nocturna_hasta": None,
                "permite_reintegro": False,
                "observaciones": None,
                "modalidad": None,
            })
        franja = franja_nocturna_convenio(db, convenio)
        if franja:
            nombre_dia, icono_dia, descripcion_dia = descripcion_tipo_dia(
                convenio, "TODOS",
            )
            grupo_todos = grupos.setdefault("TODOS", {
                "codigo": "TODOS", "nombre": nombre_dia,
                "icono": icono_dia, "descripcion": descripcion_dia,
                "conceptos": [],
            })
            grupo_todos["conceptos"].append({
                "tipo_hora": TIPO_NOCTURNAS_EN_JORNADA,
                "nombre": "Horas nocturnas dentro de jornada",
                "categoria": "Adicional nocturno",
                "descripcion": "Permite informar nocturnidad realizada dentro de la jornada habitual, sin generar horas extra.",
                "icono": "bi-moon-stars",
                "nocturna_desde": franja[0].strftime("%H:%M"),
                "nocturna_hasta": franja[1].strftime("%H:%M"),
                "permite_reintegro": False,
                "observaciones": "Puede cargarse por fechas específicas o por un período habitual.",
                "modalidad": None,
            })
    orden_dias = {"HABIL": 0, "FRANCO": 1, "FERIADO": 2, "DOMINGO": 3, "TODOS": 4}
    dias_por_convenio = {
        convenio: sorted(
            grupos.values(), key=lambda grupo: orden_dias.get(grupo["codigo"], 99),
        )
        for convenio, grupos in grupos_por_convenio.items()
    }
    return templates.TemplateResponse(request, "reglas_interactivas.html", {
        "active_page": "reglas_interactivas",
        "convenios": convenios,
        "dias_por_convenio": dias_por_convenio,
        "cantidad_reglas": {
            convenio: len([regla for regla in reglas if str(regla.convenio).strip().upper() == convenio])
            for convenio in convenios
        },
    })


@app.get("/horas", response_class=HTMLResponse)
def home(
    request: Request,
    estado: str | None = None,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    guardado: int | None = None,
    db: Session = Depends(get_db),
):
    horas: list[HoraExtra] = []
    session_user_id = (request.session.get("user") or {}).get("id")
    if not session_user_id:
        raise HTTPException(status_code=403, detail="La sesión no contiene un usuario válido.")
    assignment = obtener_asignacion_activa(request, db)
    es_jefe = assignment.perfil.upper() == "JEFE"
    error = None
    estado_filtro = str(estado or "").strip().upper()
    if estado_filtro not in {"", "PENDIENTE", "APROBADA", "RECHAZADA"}:
        estado_filtro = ""
    usuario_actual = db.scalar(
        select(Usuario)
        .options(selectinload(Usuario.tipo_contratacion))
        .where(Usuario.id == session_user_id)
    )
    if (
        assignment.perfil.upper() == "USUARIO"
        and usuario_actual is not None
        and not usuario_tiene_funciones_carga(usuario_actual, db)
    ):
        return RedirectResponse("/", status_code=303)
    if fecha_desde is None and fecha_hasta is None and usuario_actual is not None:
        fecha_desde, fecha_hasta, _ = limites_fecha_carga_usuario(usuario_actual)
    if fecha_desde and fecha_hasta and fecha_desde > fecha_hasta:
        error = "La fecha desde no puede ser posterior a la fecha hasta."
    try:
        consulta = (
            select(HoraExtra)
            .options(selectinload(HoraExtra.usuario), selectinload(HoraExtra.aprobador))
            .where(HoraExtra.usuario_id == session_user_id)
            .order_by(HoraExtra.fecha.desc(), HoraExtra.hora_inicio.desc())
        )
        if fecha_desde:
            consulta = consulta.where(HoraExtra.fecha >= fecha_desde)
        if fecha_hasta:
            consulta = consulta.where(HoraExtra.fecha <= fecha_hasta)
        if estado_filtro:
            consulta = consulta.where(HoraExtra.estado == estado_filtro)
        horas = list(db.scalars(consulta))
    except SQLAlchemyError:
        error = "No se pudieron cargar las horas extras. Verificá la conexión con la base."
    return templates.TemplateResponse(request, "dashboard.html", {
        "active_page": "horas",
        "horas": horas,
        "estado_filtro": estado_filtro,
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "guardado": guardado,
        "total_horas": sum(
            (hora.horas_totales or Decimal("0.00") for hora in horas),
            start=Decimal("0.00"),
        ),
        "cantidad_pendientes": sum(
            1 for hora in horas if hora.estado == "PENDIENTE"
        ),
        "cantidad_autorizadas": sum(
            1 for hora in horas if hora.estado == "APROBADA"
        ),
        "es_jefe": es_jefe,
        "error": error,
    })


def require_rrhh_role(request: Request) -> None:
    role = str((request.session.get("user") or {}).get("role") or "").upper()
    if role not in {"ADMIN", "RRHH"}:
        raise HTTPException(status_code=403, detail="Se requiere el rol RRHH o ADMIN.")


def acceso_conceptos_excepcionales(
    request: Request, db: Session,
) -> tuple[UsuarioRol, bool]:
    asignacion = obtener_asignacion_activa(request, db)
    administra_catalogo = asignacion.rol.upper() in {"ADMIN", "RRHH"}
    if not administra_catalogo and asignacion.perfil.upper() != "JEFE":
        raise HTTPException(status_code=403, detail="No tenés permiso para administrar excepciones.")
    return asignacion, administra_catalogo


def usuarios_asignables_excepciones(
    asignacion: UsuarioRol, administra_catalogo: bool, db: Session,
) -> list[Usuario]:
    consulta = select(Usuario).where(Usuario.status.is_(True))
    if not administra_catalogo:
        consulta = consulta.where(Usuario.roles.any(and_(
            func.upper(UsuarioRol.rol) == asignacion.rol.upper(),
            func.upper(UsuarioRol.perfil) == "USUARIO",
        )))
    return list(db.scalars(consulta.order_by(Usuario.apellido, Usuario.nombre)))


def periodos_rapidos_rrhh(fecha_consulta: date) -> dict[str, tuple[date, date]]:
    if fecha_consulta.day >= 16:
        nomina_desde = fecha_consulta.replace(day=16)
        primer_dia_siguiente = (
            fecha_consulta.replace(year=fecha_consulta.year + 1, month=1, day=1)
            if fecha_consulta.month == 12
            else fecha_consulta.replace(month=fecha_consulta.month + 1, day=1)
        )
        nomina_hasta = primer_dia_siguiente.replace(day=15)
    else:
        nomina_hasta = fecha_consulta.replace(day=15)
        mes_anterior = (
            fecha_consulta.replace(year=fecha_consulta.year - 1, month=12, day=1)
            if fecha_consulta.month == 1
            else fecha_consulta.replace(month=fecha_consulta.month - 1, day=1)
        )
        nomina_desde = mes_anterior.replace(day=16)
    mes_desde = fecha_consulta.replace(day=1)
    mes_siguiente = (
        mes_desde.replace(year=mes_desde.year + 1, month=1)
        if mes_desde.month == 12
        else mes_desde.replace(month=mes_desde.month + 1)
    )
    return {
        "nomina": (nomina_desde, nomina_hasta),
        "monotributo": (mes_desde, mes_siguiente - timedelta(days=1)),
    }


def aplicar_periodo_rapido_rrhh(
    periodo: str | None,
    fecha_desde: date | None,
    fecha_hasta: date | None,
    contrataciones: list[int],
    tipos: list[TipoContratacion],
) -> tuple[str, date | None, date | None, list[int], dict[str, tuple[date, date]]]:
    periodo_activo = (periodo or "nomina").lower()
    if periodo_activo not in {"nomina", "monotributo", "manual"}:
        periodo_activo = "nomina"
    periodos = periodos_rapidos_rrhh(date.today())
    if periodo_activo in periodos:
        fecha_desde, fecha_hasta = periodos[periodo_activo]
        tipos_buscados = (
            {"nómina", "nomina"}
            if periodo_activo == "nomina"
            else {"monotributo", "consultora"}
        )
        tipos_periodo = [
            item.id_tipo_contratacion for item in tipos
            if item.contratacion.strip().lower() in tipos_buscados
        ]
        # Si el catálogo estuviera incompleto, el acceso rápido no debe mostrar
        # accidentalmente personas de otros tipos de contratación.
        contrataciones = tipos_periodo or [-1]
    return periodo_activo, fecha_desde, fecha_hasta, contrataciones, periodos


def limites_fecha_carga_usuario(
    usuario: Usuario,
    fecha_referencia: date | None = None,
) -> tuple[date | None, date | None, str | None]:
    contratacion = (
        usuario.tipo_contratacion.contratacion.strip().lower()
        if usuario.tipo_contratacion is not None else ""
    )
    periodos = periodos_rapidos_rrhh(fecha_referencia or date.today())
    if contratacion in {"nómina", "nomina"}:
        desde, hasta = periodos["nomina"]
        return desde, hasta, "Nómina"
    if contratacion in {"monotributo", "consultora"}:
        desde, hasta = periodos["monotributo"]
        return desde, hasta, (
            "Consultora" if contratacion == "consultora" else "Monotributo"
        )
    return None, None, None


def validar_fecha_carga_usuario(usuario: Usuario, fecha_carga: date) -> None:
    desde, hasta, contratacion = limites_fecha_carga_usuario(usuario)
    if desde is not None and hasta is not None and not desde <= fecha_carga <= hasta:
        raise CalculoHorasError(
            f"Para {contratacion}, la fecha debe estar entre "
            f"el {desde.strftime('%d/%m/%Y')} y el {hasta.strftime('%d/%m/%Y')}."
        )


def calcular_descanso_articulo(
    fecha_fin: date | None,
    hora_fin: time | None,
    fecha_inicio: date | None,
    hora_inicio: time | None,
) -> tuple[date, Decimal, str] | None:
    if not all((fecha_fin, hora_fin, fecha_inicio, hora_inicio)):
        raise CalculoHorasError(
            "Completá la fecha y hora de fin de jornada y de inicio de la siguiente."
        )
    if hora_fin.minute not in {0, 30} or hora_inicio.minute not in {0, 30}:
        raise CalculoHorasError("Los horarios de HS ARTICULO deben ingresarse cada 30 minutos.")
    fin = datetime.combine(fecha_fin, hora_fin)
    inicio = datetime.combine(fecha_inicio, hora_inicio)
    if inicio <= fin:
        raise CalculoHorasError("El inicio de la jornada siguiente debe ser posterior al fin de la anterior.")
    descanso = Decimal(int((inicio - fin).total_seconds())) / Decimal("3600")
    faltantes = Decimal("12.00") - descanso
    if faltantes <= 0:
        return None
    cantidad = (
        (faltantes * Decimal("2")).quantize(Decimal("1"), rounding=ROUND_HALF_UP)
        / Decimal("2")
    ).quantize(Decimal("0.0"))
    observacion = (
        f"Fin de jornada anterior: {fecha_fin.strftime('%d/%m/%Y')} {hora_fin.strftime('%H:%M')}. "
        f"Inicio de jornada siguiente: {fecha_inicio.strftime('%d/%m/%Y')} {hora_inicio.strftime('%H:%M')}. "
        f"Descanso real: {descanso:.2f} horas. Descanso requerido: 12 horas. "
        f"HS ARTICULO: {cantidad:.1f} horas."
    )
    return fecha_fin, cantidad, observacion


def aplicar_filtros_rrhh(
    consulta,
    fecha_desde: date | None,
    fecha_hasta: date | None,
    contrataciones: list[int],
    convenios: list[str],
    usuario_id: int | None = None,
):
    if fecha_desde:
        consulta = consulta.where(HoraExtra.fecha >= fecha_desde)
    if fecha_hasta:
        consulta = consulta.where(HoraExtra.fecha <= fecha_hasta)
    if contrataciones:
        consulta = consulta.where(Usuario.id_tipo_contratacion.in_(contrataciones))
    if convenios:
        consulta = consulta.where(func.upper(Usuario.convenio).in_(convenios))
    if usuario_id is not None:
        consulta = consulta.where(HoraExtra.usuario_id == usuario_id)
    return consulta


def contexto_filtros_rrhh(
    db: Session,
    fecha_desde: date | None,
    fecha_hasta: date | None,
    contrataciones: list[int],
    convenios: list[str],
    periodo_activo: str,
    periodos_rapidos: dict[str, tuple[date, date]],
    tipos_contratacion: list[TipoContratacion],
) -> dict:
    return {
        "fecha_desde": fecha_desde,
        "fecha_hasta": fecha_hasta,
        "contrataciones_seleccionadas": contrataciones,
        "convenios_seleccionados": convenios,
        "tipos_contratacion": tipos_contratacion,
        "convenios": list(CONVENIOS_DISPONIBLES),
        "periodo_activo": periodo_activo,
        "periodos_rapidos": periodos_rapidos,
    }


@app.get("/rrhh/horas-extras", response_class=HTMLResponse)
def rrhh_horas_extras(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    contratacion: Annotated[list[int] | None, Query()] = None,
    convenio: Annotated[list[str] | None, Query()] = None,
    usuario_id: str | None = None,
    periodo: str | None = None,
    db: Session = Depends(get_db),
):
    usuario_id_normalizado = (
        int(usuario_id) if usuario_id and usuario_id.isdigit() else None
    )
    asignacion = obtener_asignacion_activa(request, db)
    es_rrhh = asignacion.rol.upper() in {"ADMIN", "RRHH"}
    vista_jefe = asignacion.perfil.upper() == "JEFE" and not es_rrhh
    if not es_rrhh and not vista_jefe:
        raise HTTPException(
            status_code=403,
            detail="Se requiere el rol RRHH, ADMIN o un perfil Jefe.",
        )
    contrataciones = list(dict.fromkeys(contratacion or []))
    convenios = [item.upper() for item in dict.fromkeys(convenio or []) if item.upper() in CONVENIOS_DISPONIBLES]
    tipos = list(db.scalars(select(TipoContratacion).order_by(TipoContratacion.contratacion)))
    periodo_activo, fecha_desde, fecha_hasta, contrataciones, periodos = aplicar_periodo_rapido_rrhh(
        periodo, fecha_desde, fecha_hasta, contrataciones, tipos,
    )
    consulta = (
        select(HoraExtra)
        .join(HoraExtra.usuario)
        .options(selectinload(HoraExtra.usuario).selectinload(Usuario.tipo_contratacion))
        .where(HoraExtra.estado.in_(("PENDIENTE", "APROBADA", "RECHAZADA")))
        .order_by(Usuario.apellido, Usuario.nombre, HoraExtra.fecha.desc(), HoraExtra.id.desc())
    )
    if vista_jefe:
        consulta = consulta.where(
            Usuario.roles.any(
                func.upper(UsuarioRol.rol) == asignacion.rol.upper(),
            )
        )
    consulta_usuarios = select(Usuario)
    if vista_jefe:
        consulta_usuarios = consulta_usuarios.where(
            Usuario.roles.any(
                func.upper(UsuarioRol.rol) == asignacion.rol.upper(),
            )
        )
    usuarios_filtro = list(db.scalars(
        consulta_usuarios.order_by(Usuario.apellido, Usuario.nombre)
    ))
    consulta = aplicar_filtros_rrhh(
        consulta, fecha_desde, fecha_hasta, contrataciones, convenios,
        usuario_id_normalizado,
    )
    horas = list(db.scalars(consulta))
    grupos_por_usuario: dict[int, dict] = {}
    for hora in horas:
        grupo = grupos_por_usuario.setdefault(hora.usuario_id, {
            "usuario": hora.usuario,
            "horas": [],
            "pendientes": 0,
            "autorizadas": 0,
            "rechazadas": 0,
            "resumen": {},
        })
        grupo["horas"].append(hora)
        for concepto, cantidad in conceptos_exportables_hora(hora):
            grupo["resumen"][concepto] = (
                grupo["resumen"].get(concepto, Decimal("0.00")) + cantidad
            )
        if hora.estado == "PENDIENTE":
            grupo["pendientes"] += 1
        elif hora.estado == "APROBADA":
            grupo["autorizadas"] += 1
        elif hora.estado == "RECHAZADA":
            grupo["rechazadas"] += 1
    grupos = list(grupos_por_usuario.values())
    for grupo in grupos:
        grupo["resumen_conceptos"] = sorted(grupo.pop("resumen").items())
    contexto = contexto_filtros_rrhh(
        db, fecha_desde, fecha_hasta, contrataciones, convenios,
        periodo_activo, periodos, tipos,
    )
    contexto.update({
        "active_page": "rrhh_horas",
        "horas": horas,
        "grupos": grupos,
        "pendientes": sum(item.estado == "PENDIENTE" for item in horas),
        "autorizadas": sum(item.estado == "APROBADA" for item in horas),
        "rechazadas": sum(item.estado == "RECHAZADA" for item in horas),
        "vista_jefe": vista_jefe,
        "es_rrhh": es_rrhh,
        "rol_activo": asignacion.rol,
        "usuarios_filtro": usuarios_filtro,
        "usuario_id": usuario_id_normalizado,
    })
    return templates.TemplateResponse(request, "rrhh_horas.html", contexto)


class FilaExportacion(NamedTuple):
    legajo: str | None
    nombre: str
    apellido: str
    tipo_hora: str
    cantidad: Decimal


class FilaDetalleExportacion(NamedTuple):
    legajo: str | None
    nombre: str
    apellido: str
    tipo_hora: str
    cantidad: Decimal
    fecha: date
    observaciones: str | None
    fecha_descanso_compensatorio: date | None


NOMBRES_CONCEPTOS_LIQUIDACION = {
    "FRANCO TRABAJADO": "Día Franco o Feriado",
    "FERIADO TRABAJADO": "Día Franco o Feriado",
    "REINTEGRO FRANCO": "Reintegro de día FTN",
    "REINTEGRO FERIADO": "Reintegro de día FTN",
    "50": "Horas extras al 50%",
    "50%": "Horas extras al 50%",
    "100": "Horas extras al 100%",
    "100%": "Horas extras al 100%",
    "HS ARTICULO": "Horas Articulo",
    "HS ARTÍCULO": "Horas Articulo",
    "HORAS NOCTURNAS": "HORAS NOCTURNAS",
    "NOCTURNAS EN JORNADA": "HORAS NOCTURNAS",
    "HORAS NOCTURNAS EN JORNADA": "HORAS NOCTURNAS",
    "DOMINGO": "Domingo",
    "MERIENDA": "Vale de Merienda Imp/ SAT",
    "EXTERIOR COMUN": "Exterior Común",
    "EXTERIOR COMÚN": "Exterior Común",
    "EXTERIOR PRENSA": "Ext/ Prensa Ext/ o Franco",
    "COMIDA": "Vale de Comida Imput SAT",
}


def nombre_concepto_liquidacion(concepto: str) -> str:
    concepto_limpio = concepto.strip()
    return NOMBRES_CONCEPTOS_LIQUIDACION.get(
        concepto_limpio.upper(), concepto_limpio,
    )


def conceptos_exportables_hora(hora: HoraExtra) -> list[tuple[str, Decimal]]:
    conceptos: list[tuple[str, Decimal]] = []
    tipo_registro = str(hora.tipo_registro or "HORAS").upper()
    tipo_dia = str(hora.tipo_dia or "").upper()
    if tipo_registro == "HORAS" and hora.tipo_hora and hora.horas_totales:
        conceptos.append((hora.tipo_hora, hora.horas_totales))
    elif (
        tipo_registro == TIPO_PRECARGA_AUTORIZADA
        and hora.tipo_hora
        and hora.cantidad
    ):
        conceptos.append((hora.tipo_hora, hora.cantidad))
    elif tipo_registro == "OTRAS" and hora.tipo_hora and hora.cantidad:
        conceptos.append((hora.tipo_hora, hora.cantidad))
    elif tipo_registro == "EXCEPCIONAL" and hora.concepto_excepcional_nombre and hora.cantidad:
        conceptos.append((hora.concepto_excepcional_nombre, hora.cantidad))
    elif tipo_registro == "DIA_TRABAJADO":
        concepto = "FERIADO TRABAJADO" if tipo_dia == "FERIADO" else "FRANCO TRABAJADO"
        conceptos.append((concepto, Decimal("1.00")))
    elif tipo_registro == "REINTEGRO":
        concepto = "REINTEGRO FERIADO" if tipo_dia == "FERIADO" else "REINTEGRO FRANCO"
        conceptos.append((concepto, hora.cantidad or Decimal("1.00")))

    # La nocturnidad es un adicional y se exporta además del concepto base.
    if hora.horas_nocturnas and hora.horas_nocturnas > 0:
        conceptos.append(("HORAS NOCTURNAS", hora.horas_nocturnas))
    return [
        (nombre_concepto_liquidacion(concepto), cantidad)
        for concepto, cantidad in conceptos
    ]


def consulta_cargas_exportacion_rrhh(
    fecha_desde: date | None,
    fecha_hasta: date | None,
    contrataciones: list[int],
    convenios: list[str],
):
    consulta = (
        select(HoraExtra)
        .join(HoraExtra.usuario)
        .options(selectinload(HoraExtra.usuario).selectinload(Usuario.tipo_contratacion))
        .where(HoraExtra.estado == "APROBADA")
    )
    return aplicar_filtros_rrhh(
        consulta, fecha_desde, fecha_hasta, contrataciones, convenios,
    )


def filas_exportacion_rrhh(
    db: Session,
    fecha_desde: date | None,
    fecha_hasta: date | None,
    contrataciones: list[int],
    convenios: list[str],
) -> list[FilaExportacion]:
    consulta = consulta_cargas_exportacion_rrhh(
        fecha_desde, fecha_hasta, contrataciones, convenios,
    )
    acumulados: dict[tuple[int, str], Decimal] = {}
    usuarios: dict[int, Usuario] = {}

    def acumular(hora: HoraExtra, concepto: str, cantidad: Decimal) -> None:
        clave = (hora.usuario_id, concepto)
        acumulados[clave] = acumulados.get(clave, Decimal("0.00")) + cantidad
        usuarios[hora.usuario_id] = hora.usuario

    for hora in db.scalars(consulta):
        for concepto, cantidad in conceptos_exportables_hora(hora):
            acumular(hora, concepto, cantidad)

    filas = []
    for (usuario_id, concepto), cantidad in acumulados.items():
        usuario = usuarios[usuario_id]
        filas.append(FilaExportacion(
            usuario.legajo,
            usuario.nombre,
            usuario.apellido,
            concepto,
            cantidad,
        ))
    return sorted(
        filas,
        key=lambda fila: (fila.apellido.lower(), fila.nombre.lower(), fila.tipo_hora.lower()),
    )


def detalle_exportacion_rrhh(
    db: Session,
    fecha_desde: date | None,
    fecha_hasta: date | None,
    contrataciones: list[int],
    convenios: list[str],
) -> list[FilaDetalleExportacion]:
    consulta = consulta_cargas_exportacion_rrhh(
        fecha_desde, fecha_hasta, contrataciones, convenios,
    )
    filas = [
        FilaDetalleExportacion(
            hora.usuario.legajo,
            hora.usuario.nombre,
            hora.usuario.apellido,
            concepto,
            cantidad,
            hora.fecha,
            hora.observaciones,
            hora.fecha_descanso_compensatorio,
        )
        for hora in db.scalars(consulta)
        for concepto, cantidad in conceptos_exportables_hora(hora)
    ]
    return sorted(
        filas,
        key=lambda fila: (
            fila.apellido.lower(), fila.nombre.lower(), fila.fecha, fila.tipo_hora.lower(),
        ),
    )


def fechas_texto_liquidacion(fecha_hasta: date | None) -> tuple[str, str]:
    """Devuelve el rango nominal 01-30 del mes liquidado, aun si el día 30 no existe."""
    mes_liquidacion = fecha_hasta or date.today()
    sufijo = mes_liquidacion.strftime("%m/%Y")
    return f"01/{sufijo}", f"30/{sufijo}"


@app.get("/rrhh/exportacion", response_class=HTMLResponse)
def rrhh_exportacion(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    contratacion: Annotated[list[int] | None, Query()] = None,
    convenio: Annotated[list[str] | None, Query()] = None,
    periodo: str | None = None,
    db: Session = Depends(get_db),
):
    require_rrhh_role(request)
    contrataciones = list(dict.fromkeys(contratacion or []))
    convenios = [item.upper() for item in dict.fromkeys(convenio or []) if item.upper() in CONVENIOS_DISPONIBLES]
    tipos = list(db.scalars(select(TipoContratacion).order_by(TipoContratacion.contratacion)))
    periodo_activo, fecha_desde, fecha_hasta, contrataciones, periodos = aplicar_periodo_rapido_rrhh(
        periodo, fecha_desde, fecha_hasta, contrataciones, tipos,
    )
    filas = filas_exportacion_rrhh(
        db, fecha_desde, fecha_hasta, contrataciones, convenios,
    )
    contexto = contexto_filtros_rrhh(
        db, fecha_desde, fecha_hasta, contrataciones, convenios,
        periodo_activo, periodos, tipos,
    )
    fecha_liquidacion_desde, fecha_liquidacion_hasta = fechas_texto_liquidacion(
        fecha_hasta,
    )
    contexto.update({
        "active_page": "rrhh_exportacion",
        "filas": filas,
        "fecha_liquidacion_desde": fecha_liquidacion_desde,
        "fecha_liquidacion_hasta": fecha_liquidacion_hasta,
    })
    return templates.TemplateResponse(request, "rrhh_exportacion.html", contexto)


@app.get("/rrhh/exportacion.xlsx")
def rrhh_exportacion_xlsx(
    request: Request,
    fecha_desde: date | None = None,
    fecha_hasta: date | None = None,
    contratacion: Annotated[list[int] | None, Query()] = None,
    convenio: Annotated[list[str] | None, Query()] = None,
    periodo: str | None = None,
    db: Session = Depends(get_db),
):
    require_rrhh_role(request)
    contrataciones = list(dict.fromkeys(contratacion or []))
    convenios = [item.upper() for item in dict.fromkeys(convenio or []) if item.upper() in CONVENIOS_DISPONIBLES]
    tipos = list(db.scalars(select(TipoContratacion).order_by(TipoContratacion.contratacion)))
    _, fecha_desde, fecha_hasta, contrataciones, _ = aplicar_periodo_rapido_rrhh(
        periodo, fecha_desde, fecha_hasta, contrataciones, tipos,
    )
    filas = filas_exportacion_rrhh(
        db, fecha_desde, fecha_hasta, contrataciones, convenios,
    )
    detalle = detalle_exportacion_rrhh(
        db, fecha_desde, fecha_hasta, contrataciones, convenios,
    )
    fecha_liquidacion_desde, fecha_liquidacion_hasta = fechas_texto_liquidacion(
        fecha_hasta,
    )
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Horas extras"
    headers = (
        "LEGAJO", "BÚSQUE POR NOMBRE", "CONCEPTO A LIQUIDAR", "FECHA DESDE",
        "FECHA HASTA", "IMPORTE", "CANT.",
    )
    sheet.append(headers)
    for legajo, nombre, apellido, tipo_hora, cantidad in filas:
        sheet.append((
            legajo or "",
            f"{apellido} ,{nombre}".strip(),
            tipo_hora,
            fecha_liquidacion_desde,
            fecha_liquidacion_hasta,
            "",
            float(cantidad),
        ))
    header_fill = PatternFill("solid", fgColor="1D4ED8")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 16
    sheet.column_dimensions["B"].width = 34
    sheet.column_dimensions["C"].width = 28
    sheet.column_dimensions["D"].width = 15
    sheet.column_dimensions["E"].width = 15
    sheet.column_dimensions["F"].width = 14
    sheet.column_dimensions["G"].width = 14
    for cell in sheet["G"][1:]:
        cell.number_format = "0.00"
    detail_sheet = workbook.create_sheet("Detalle")
    detail_headers = (
        "Legajo", "Nombre y apellido", "Concepto", "Cantidad", "Fecha", "Observación",
        "Día previsto para tomar",
    )
    detail_sheet.append(detail_headers)
    for legajo, nombre, apellido, concepto, cantidad, fecha, observaciones, fecha_descanso in detalle:
        detail_sheet.append((
            legajo or "",
            f"{nombre} {apellido}".strip(),
            concepto,
            float(cantidad),
            fecha,
            observaciones or "",
            fecha_descanso,
        ))
    for cell in detail_sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = header_fill
    detail_sheet.freeze_panes = "A2"
    detail_sheet.auto_filter.ref = detail_sheet.dimensions
    detail_sheet.column_dimensions["A"].width = 16
    detail_sheet.column_dimensions["B"].width = 34
    detail_sheet.column_dimensions["C"].width = 24
    detail_sheet.column_dimensions["D"].width = 14
    detail_sheet.column_dimensions["E"].width = 14
    detail_sheet.column_dimensions["F"].width = 55
    detail_sheet.column_dimensions["G"].width = 24
    for cell in detail_sheet["D"][1:]:
        cell.number_format = "0.00"
    for cell in detail_sheet["E"][1:]:
        cell.number_format = "DD/MM/YYYY"
    for cell in detail_sheet["G"][1:]:
        cell.number_format = "DD/MM/YYYY"
    output = io.BytesIO()
    workbook.save(output)
    return Response(
        content=output.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="horas_extras_rrhh.xlsx"'},
    )


@app.get("/solicitudes/nueva", response_class=HTMLResponse)
def nueva_solicitud(
    request: Request,
    destino: str = "propio",
    usuario_id: int | None = None,
    error: str | None = None,
    db: Session = Depends(get_db),
):
    usuarios: list[Usuario] = []
    usuario_seleccionado: Usuario | None = None
    try:
        _, usuarios = usuarios_habilitados_para_carga(request, db, destino)
        if destino == "rrhh":
            usuario_seleccionado = next(
                (usuario for usuario in usuarios if usuario.id == usuario_id),
                None,
            )
            if usuario_id is not None and usuario_seleccionado is None:
                error = "El usuario seleccionado no está habilitado."
        else:
            destino = "propio"
            usuario_seleccionado = usuarios[0] if usuarios else None
    except SQLAlchemyError:
        error = "No se pudieron cargar los usuarios activos. Verificá la conexión con la base."
    fecha_minima, fecha_maxima, _ = limites_fecha_carga_usuario(usuario_seleccionado) if usuario_seleccionado else (None, None, None)
    franja_nocturna = franja_nocturna_convenio(
        db, usuario_seleccionado.convenio if usuario_seleccionado else None,
    )
    convenio_usuario = usuario_seleccionado.convenio if usuario_seleccionado else None
    return templates.TemplateResponse(request, "home.html", {
        "active_page": "rrhh_carga" if destino == "rrhh" else "horas",
        "usuarios": usuarios,
        "usuario_fijo": usuario_seleccionado,
        "seleccion_administrativa": destino == "rrhh" and usuario_seleccionado is None,
        "es_carga_administrativa": destino == "rrhh",
        "destino": destino,
        "fecha_hoy": date.today().isoformat(),
        "fecha_minima": fecha_minima.isoformat() if fecha_minima else "",
        "fecha_maxima": fecha_maxima.isoformat() if fecha_maxima else "",
        "tipos_otras_cargas": tipos_otras_cargas(db, convenio_usuario),
        "conceptos_excepcionales": conceptos_excepcionales_habilitados(db, usuario_seleccionado.id) if usuario_seleccionado else [],
        "feriados": fechas_feriados(db),
        "feriados_sin_devolucion": fechas_feriados_sin_devolucion(db),
        "franja_nocturna": franja_nocturna,
        "permite_hs_articulo": convenio_permite_hs_articulo(db, convenio_usuario),
        "permite_reintegro_convenio": convenio_permite_reintegro(db, convenio_usuario),
        "error": error,
    })


def renderizar_borrador_solicitud(
    request: Request,
    db: Session,
    resultados: list,
    cargas: list[str],
    reintegros: list[str],
    observaciones_resultados: list[str | None],
    error: str | None = None,
    aviso: str | None = None,
    destino: str = "propio",
    status_code: int = 200,
):
    _, usuarios_habilitados = usuarios_habilitados_para_carga(request, db, destino)
    usuario_id = resultados[0].usuario_id if resultados else None
    usuario_seleccionado = next(
        (usuario for usuario in usuarios_habilitados if usuario.id == usuario_id),
        None,
    )
    usuarios = [usuario_seleccionado] if usuario_seleccionado else []
    fecha_minima, fecha_maxima, _ = (
        limites_fecha_carga_usuario(usuarios[0])
        if usuarios else (None, None, None)
    )
    conceptos_automaticos, claves_jornada, claves_nocturnidad = conceptos_automaticos_borrador(
        resultados, db,
    )
    franja_nocturna = franja_nocturna_convenio(
        db, usuarios[0].convenio if usuarios else None,
    )
    convenio_usuario = usuarios[0].convenio if usuarios else None
    fechas_descanso = [decode_carga(carga)[11] for carga in cargas]
    return templates.TemplateResponse(
        request,
        "solicitud.html",
        {
            "active_page": "rrhh_carga" if destino == "rrhh" else "horas",
            "resultados": resultados,
            "claves_jornada": claves_jornada,
            "claves_nocturnidad": claves_nocturnidad,
            "conceptos_automaticos": conceptos_automaticos,
            "observaciones_resultados": observaciones_resultados,
            "fechas_descanso": fechas_descanso,
            "cargas": cargas,
            "reintegros": reintegros,
            "usuarios": usuarios,
            "usuario_fijo": usuario_seleccionado,
            "seleccion_administrativa": False,
            "es_carga_administrativa": destino == "rrhh",
            "destino": destino,
            "fecha_hoy": date.today().isoformat(),
            "fecha_minima": fecha_minima.isoformat() if fecha_minima else "",
            "fecha_maxima": fecha_maxima.isoformat() if fecha_maxima else "",
            "tipos_otras_cargas": tipos_otras_cargas(db, convenio_usuario),
            "conceptos_excepcionales": conceptos_excepcionales_habilitados(
                db, usuarios[0].id,
            ) if usuarios else [],
            "feriados": fechas_feriados(db),
            "feriados_sin_devolucion": fechas_feriados_sin_devolucion(db),
            "franja_nocturna": franja_nocturna,
            "permite_hs_articulo": convenio_permite_hs_articulo(db, convenio_usuario),
            "permite_reintegro_convenio": convenio_permite_reintegro(db, convenio_usuario),
            "total_horas": sum(
                (item.horas_totales or Decimal("0.00") for item in resultados),
                start=Decimal("0.00"),
            ),
            "total_nocturnas": sum(
                (item.horas_nocturnas or Decimal("0.00") for item in resultados),
                start=Decimal("0.00"),
            ),
            "error": error,
            "aviso": aviso,
        },
        status_code=status_code,
    )


@app.post("/solicitudes/procesar", response_class=HTMLResponse)
def procesar_solicitud(
    request: Request,
    usuario_id: Annotated[int, Form()],
    fecha: Annotated[list[date] | None, Form()] = None,
    tipo_registro: Annotated[str, Form()] = "HORAS",
    hora_inicio: Annotated[time | None, Form()] = None,
    hora_fin: Annotated[time | None, Form()] = None,
    observaciones: Annotated[str | None, Form()] = None,
    destino: Annotated[str, Form()] = "propio",
    cargas: Annotated[list[str] | None, Form()] = None,
    reintegros: Annotated[list[str] | None, Form()] = None,
    marcar_como_franco: Annotated[str | None, Form()] = None,
    solicita_reintegro_dia: Annotated[str | None, Form()] = None,
    incluye_horas_extra: Annotated[str | None, Form()] = None,
    tipo_otra_carga: Annotated[str | None, Form()] = None,
    cantidad: Annotated[Decimal | None, Form()] = None,
    jornada_nocturna_inicio: Annotated[time | None, Form()] = None,
    jornada_nocturna_fin: Annotated[time | None, Form()] = None,
    modo_nocturnas: Annotated[str, Form()] = "FECHAS",
    nocturna_fecha_desde: Annotated[date | None, Form()] = None,
    nocturna_fecha_hasta: Annotated[date | None, Form()] = None,
    nocturna_dias_semana: Annotated[list[int] | None, Form()] = None,
    modo_fechas: Annotated[str | None, Form()] = None,
    periodo_fecha_desde: Annotated[date | None, Form()] = None,
    periodo_fecha_hasta: Annotated[date | None, Form()] = None,
    periodo_dias_semana: Annotated[list[int] | None, Form()] = None,
    jornada_hora_inicio: Annotated[time | None, Form()] = None,
    jornada_hora_fin: Annotated[time | None, Form()] = None,
    cantidad_horas_extra: Annotated[Decimal | None, Form()] = None,
    fecha_descanso_compensatorio: Annotated[date | None, Form()] = None,
    concepto_adicional_tipo: Annotated[list[str] | None, Form()] = None,
    concepto_adicional_cantidad: Annotated[list[Decimal] | None, Form()] = None,
    concepto_adicional_observacion: Annotated[list[str] | None, Form()] = None,
    calcular_hs_articulo: Annotated[str | None, Form()] = None,
    articulo_fin_fecha: Annotated[date | None, Form()] = None,
    articulo_fin_hora: Annotated[time | None, Form()] = None,
    articulo_inicio_fecha: Annotated[date | None, Form()] = None,
    articulo_inicio_hora: Annotated[time | None, Form()] = None,
    articulo_jornada_fecha_lista: Annotated[list[date] | None, Form()] = None,
    articulo_fin_fecha_lista: Annotated[list[date] | None, Form()] = None,
    articulo_fin_hora_lista: Annotated[list[time] | None, Form()] = None,
    articulo_inicio_fecha_lista: Annotated[list[date] | None, Form()] = None,
    articulo_inicio_hora_lista: Annotated[list[time] | None, Form()] = None,
    db: Session = Depends(get_db),
):
    fechas_recibidas = list(fecha or [])
    fechas = list(dict.fromkeys(fechas_recibidas))
    resultados = []
    observaciones_resultados: list[str | None] = []
    cargas_codificadas = list(cargas or [])
    selecciones_reintegro = [
        value if value in {"SI", "NO"} else "NO"
        for value in (reintegros or [])
    ]
    if len(selecciones_reintegro) < len(cargas_codificadas):
        selecciones_reintegro.extend(
            ["NO"] * (len(cargas_codificadas) - len(selecciones_reintegro))
        )
    selecciones_reintegro = selecciones_reintegro[:len(cargas_codificadas)]
    error = None
    aviso = None
    feriados_omitidos = 0
    try:
        tipo_registro = tipo_registro.strip().upper()
        if len(fechas_recibidas) != len(fechas):
            raise CalculoHorasError("No repitas una fecha en la misma carga.")
        observacion_limpia = clean_optional(observaciones)
        if observacion_limpia is None:
            raise CalculoHorasError("La observación es obligatoria.")
        _, usuarios_habilitados = usuarios_habilitados_para_carga(request, db, destino)
        usuario_carga = next(
            (item for item in usuarios_habilitados if item.id == usuario_id),
            None,
        )
        if usuario_carga is None:
            raise CalculoHorasError("No tenés permiso para cargar horas al usuario seleccionado.")
        convenio_carga = str(usuario_carga.convenio or "").strip().upper()
        tipo_otra_normalizado = str(tipo_otra_carga or "").strip().upper()
        es_articulo_independiente = (
            tipo_registro == "OTRAS" and tipo_otra_normalizado == "HS ARTICULO"
        )
        regla_otra_seleccionada = None
        concepto_excepcional_seleccionado = None
        if tipo_registro == "OTRAS" and tipo_otra_normalizado.startswith("EXCEPCIONAL:"):
            try:
                concepto_excepcional_id = int(tipo_otra_normalizado.split(":", 1)[1])
            except ValueError as error_conversion:
                raise CalculoHorasError(
                    "El concepto excepcional seleccionado no es válido."
                ) from error_conversion
            concepto_excepcional_seleccionado = db.get(
                ConceptoExcepcional, concepto_excepcional_id,
            )
        if (
            tipo_registro == "OTRAS"
            and tipo_otra_normalizado
            and not tipo_otra_normalizado.startswith("EXCEPCIONAL:")
            and tipo_otra_normalizado not in TIPOS_NOCTURNAS_EN_JORNADA
        ):
            regla_otra_seleccionada = regla_otra_carga(
                db, convenio_carga, tipo_otra_normalizado,
            )
        es_carga_regla_por_dia = (
            regla_otra_seleccionada is not None
            and str(regla_otra_seleccionada.modalidad_carga or "").upper()
            == "CARGA_POR_DIA"
        )
        es_carga_excepcional_por_dia = (
            concepto_excepcional_seleccionado is not None
            and concepto_excepcional_seleccionado.activo
            and concepto_excepcional_seleccionado.tipo == "CARGA_POR_DIA"
        )
        es_carga_por_dia = es_carga_regla_por_dia or es_carga_excepcional_por_dia
        es_nocturnidad_solicitada = (
            tipo_registro == "OTRAS"
            and tipo_otra_normalizado in TIPOS_NOCTURNAS_EN_JORNADA
        )
        modo_seleccion_fechas = str(
            modo_fechas or modo_nocturnas or "FECHAS"
        ).strip().upper()
        usa_selector_periodo = (
            tipo_registro == "HORAS"
            or es_nocturnidad_solicitada
            or es_carga_por_dia
        )
        fecha_periodo_desde = periodo_fecha_desde or nocturna_fecha_desde
        fecha_periodo_hasta = periodo_fecha_hasta or nocturna_fecha_hasta
        dias_periodo = periodo_dias_semana or nocturna_dias_semana
        if usa_selector_periodo and modo_seleccion_fechas == "PERIODO":
            if fecha_periodo_desde is None or fecha_periodo_hasta is None:
                raise CalculoHorasError("Ingresá las fechas desde y hasta del período habitual.")
            if fecha_periodo_hasta < fecha_periodo_desde:
                raise CalculoHorasError("La fecha hasta no puede ser anterior a la fecha desde.")
            if (fecha_periodo_hasta - fecha_periodo_desde).days > 62:
                raise CalculoHorasError("El período habitual no puede superar 63 días.")
            dias_elegidos = set(dias_periodo or [])
            if not dias_elegidos or not dias_elegidos.issubset(set(range(7))):
                raise CalculoHorasError("Seleccioná al menos un día válido de la semana.")
            fechas = []
            fecha_iterada = fecha_periodo_desde
            while fecha_iterada <= fecha_periodo_hasta:
                if (
                    fecha_iterada.weekday() in dias_elegidos
                ):
                    if (
                        es_nocturnidad_solicitada
                        and clasificar_tipo_dia(fecha_iterada, db) == "FERIADO"
                    ):
                        feriados_omitidos += 1
                    else:
                        fechas.append(fecha_iterada)
                fecha_iterada += timedelta(days=1)
        elif usa_selector_periodo and modo_seleccion_fechas != "FECHAS":
            raise CalculoHorasError("La modalidad de selección de días no es válida.")
        if es_articulo_independiente and articulo_fin_fecha is not None:
            fechas = [articulo_fin_fecha]
        if not fechas:
            raise CalculoHorasError(
                "No hay fechas válidas para cargar. Revisá el período, los días elegidos y los feriados."
            )
        fecha_principal = fechas[0]
        if convenio_carga == "FC" and (
            concepto_adicional_tipo
            or str(calcular_hs_articulo or "").upper() == "SI"
        ):
            raise CalculoHorasError(
                "FC no permite agregar conceptos adicionales."
            )
        if (
            str(calcular_hs_articulo or "").upper() == "SI"
            and not convenio_permite_hs_articulo(db, convenio_carga)
        ):
            raise CalculoHorasError(
                f"HS ARTICULO no está configurado para el convenio {convenio_carga}."
            )
        seleccion_reintegro_solicitada = str(
            solicita_reintegro_dia or "NO"
        ).upper()
        if (
            convenio_carga in {"FC", "MONOTRIBUTISTA"}
            and seleccion_reintegro_solicitada in {"ON", "MEDIO", "COMPLETO", "SI"}
        ):
            raise CalculoHorasError(
                f"{convenio_carga} no permite solicitar reintegro."
            )
        for fecha_seleccionada in fechas:
            validar_fecha_carga_usuario(usuario_carga, fecha_seleccionada)
        for carga in cargas_codificadas:
            datos = decode_carga(carga)
            if datos[0] not in ids_habilitados_para_confirmar(request, db, destino):
                raise CalculoHorasError("Una de las cargas contiene un usuario no autorizado.")
            resultados.append(calcular_carga_codificada(
                datos, db, permitir_fc_sin_asignacion=destino == "rrhh",
            ))
            observaciones_resultados.append(datos[4])
        nuevos_items = []
        tipos_adicionales = list(concepto_adicional_tipo or [])
        cantidades_adicionales = list(concepto_adicional_cantidad or [])
        observaciones_adicionales = list(concepto_adicional_observacion or [])
        if not (
            len(tipos_adicionales)
            == len(cantidades_adicionales)
            == len(observaciones_adicionales)
        ):
            raise CalculoHorasError("Los conceptos adicionales están incompletos.")
        adicionales = []
        for tipo, cantidad_adicional, observacion_adicional in zip(
            tipos_adicionales, cantidades_adicionales, observaciones_adicionales,
        ):
            if tipo.strip().upper() == "HS ARTICULO":
                raise CalculoHorasError(
                    "Calculá HS ARTICULO desde el apartado de descanso entre jornadas."
                )
            if cantidad_adicional <= 0 or cantidad_adicional % Decimal("0.5") != 0:
                raise CalculoHorasError(
                    f"La cantidad de {tipo} debe ingresarse de a 0,5."
                )
            observacion_concepto = clean_optional(observacion_adicional)
            if observacion_concepto is None:
                raise CalculoHorasError(
                    f"Justificá la cantidad indicada para {tipo}."
                )
            adicionales.append((tipo, cantidad_adicional, observacion_concepto))

        def agregar_conceptos_adicionales(fecha_concepto: date) -> None:
            for tipo, cantidad_adicional, observacion_concepto in adicionales:
                if tipo.startswith("EXCEPCIONAL:"):
                    try:
                        concepto_id = int(tipo.split(":", 1)[1])
                    except ValueError as error_conversion:
                        raise CalculoHorasError("El concepto excepcional seleccionado no es válido.") from error_conversion
                    concepto = calcular_resultado_concepto_excepcional(
                        usuario_id, fecha_concepto, concepto_id,
                        cantidad_adicional, db,
                    )
                    tipo_item = "EXCEPCIONAL"
                else:
                    concepto = calcular_resultado_otra_carga(
                        usuario_id, fecha_concepto, tipo, cantidad_adicional, db,
                    )
                    tipo_item = "OTRAS"
                nuevos_items.append((
                    concepto, fecha_concepto, None, None, tipo_item, False,
                    observacion_concepto,
                ))

        if tipo_registro == "OTRAS":
            if not tipo_otra_carga:
                raise CalculoHorasError("Seleccioná un tipo de carga.")
            es_nocturnidad_jornada = (
                tipo_otra_carga.strip().upper() in TIPOS_NOCTURNAS_EN_JORNADA
            )
            if (
                not es_nocturnidad_jornada
                and not es_carga_por_dia
                and not es_articulo_independiente
                and cantidad is None
            ):
                raise CalculoHorasError("Ingresá la cantidad de la carga.")
            if es_articulo_independiente:
                calculo_articulo = calcular_descanso_articulo(
                    articulo_fin_fecha, articulo_fin_hora,
                    articulo_inicio_fecha, articulo_inicio_hora,
                )
                if calculo_articulo is None:
                    raise CalculoHorasError(
                        "Se cumplieron las 12 horas de descanso; no corresponde cargar HS ARTICULO."
                    )
                fecha_articulo, cantidad_articulo, detalle_articulo = calculo_articulo
                validar_fecha_carga_usuario(usuario_carga, fecha_articulo)
                observacion_articulo = f"{observacion_limpia}\n{detalle_articulo}"
                if observacion_articulo in observaciones_resultados:
                    raise CalculoHorasError("El borrador ya contiene este cálculo de HS ARTICULO.")
                existente_articulo = db.scalar(
                    select(HoraExtra.id).where(
                        HoraExtra.usuario_id == usuario_id,
                        HoraExtra.tipo_registro == "OTRAS",
                        func.upper(HoraExtra.tipo_hora) == "HS ARTICULO",
                        HoraExtra.observaciones == observacion_articulo,
                        HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
                    ).limit(1)
                )
                if existente_articulo is not None:
                    raise CalculoHorasError("Ya existe esta carga de HS ARTICULO.")
                articulo = calcular_resultado_otra_carga(
                    usuario_id, fecha_articulo, "HS ARTICULO", cantidad_articulo, db,
                )
                nuevos_items.append((
                    articulo, fecha_articulo, None, None, "OTRAS", False,
                    observacion_articulo,
                ))
            elif es_nocturnidad_jornada:
                if jornada_nocturna_inicio is None or jornada_nocturna_fin is None:
                    raise CalculoHorasError("Ingresá el horario completo de tu jornada habitual.")
                superposiciones_omitidas = 0
                for fecha_seleccionada in fechas:
                    try:
                        validar_nocturnidad_sin_superposicion(
                            usuario_id, fecha_seleccionada,
                            jornada_nocturna_inicio, jornada_nocturna_fin, db,
                            resultados + [item[0] for item in nuevos_items],
                        )
                    except CalculoHorasError:
                        if modo_seleccion_fechas != "PERIODO":
                            raise
                        superposiciones_omitidas += 1
                        continue
                    resultado = calcular_resultado_nocturnas_en_jornada(
                        usuario_id, fecha_seleccionada,
                        jornada_nocturna_inicio, jornada_nocturna_fin, db,
                    )
                    nuevos_items.append((
                        resultado, fecha_seleccionada, jornada_nocturna_inicio,
                        jornada_nocturna_fin, "OTRAS", False,
                    ))
                if not nuevos_items:
                    raise CalculoHorasError(
                        "Todas las fechas del período fueron omitidas por feriados o cargas superpuestas."
                    )
                omisiones = []
                if feriados_omitidos:
                    omisiones.append(f"{feriados_omitidos} feriado(s)")
                if superposiciones_omitidas:
                    omisiones.append(
                        f"{superposiciones_omitidas} fecha(s) con cargas superpuestas"
                    )
                if omisiones:
                    aviso = "Se omitieron " + " y ".join(omisiones) + "."
            elif es_carga_por_dia:
                cantidad_por_dia = (
                    concepto_excepcional_seleccionado.cantidad_predeterminada
                    if es_carga_excepcional_por_dia
                    else regla_otra_seleccionada.cantidad_predeterminada
                )
                if cantidad_por_dia is None or cantidad_por_dia <= 0:
                    raise CalculoHorasError("El concepto seleccionado no tiene una cantidad por día válida.")
                for fecha_seleccionada in fechas:
                    if es_carga_excepcional_por_dia:
                        concepto_id = concepto_excepcional_seleccionado.id_concepto_excepcional
                        existe_en_borrador = any(
                            item.usuario_id == usuario_id and item.fecha == fecha_seleccionada
                            and item.tipo_registro == "EXCEPCIONAL"
                            and item.concepto_excepcional_id == concepto_id
                            for item in resultados
                        )
                        existe_guardada = db.scalar(select(HoraExtra.id).where(
                            HoraExtra.usuario_id == usuario_id,
                            HoraExtra.fecha == fecha_seleccionada,
                            HoraExtra.tipo_registro == "EXCEPCIONAL",
                            HoraExtra.concepto_excepcional_id == concepto_id,
                            HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
                        ).limit(1))
                        nombre_concepto = concepto_excepcional_seleccionado.nombre
                    else:
                        existe_en_borrador = any(
                            item.usuario_id == usuario_id and item.fecha == fecha_seleccionada
                            and item.tipo_registro == "OTRAS"
                            and str(item.tipo_hora or "").strip().upper() == tipo_otra_normalizado
                            for item in resultados
                        )
                        existe_guardada = db.scalar(select(HoraExtra.id).where(
                            HoraExtra.usuario_id == usuario_id,
                            HoraExtra.fecha == fecha_seleccionada,
                            HoraExtra.tipo_registro == "OTRAS",
                            func.upper(HoraExtra.tipo_hora) == tipo_otra_normalizado,
                            HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
                        ).limit(1))
                        nombre_concepto = regla_otra_seleccionada.tipo_hora
                    if existe_en_borrador or existe_guardada is not None:
                        raise CalculoHorasError(
                            f"Ya existe {nombre_concepto} para el "
                            f"{fecha_seleccionada.strftime('%d/%m/%Y')}."
                        )
                    if es_carga_excepcional_por_dia:
                        resultado = calcular_resultado_concepto_excepcional(
                            usuario_id, fecha_seleccionada, concepto_id, cantidad_por_dia, db,
                        )
                        tipo_item = "EXCEPCIONAL"
                    else:
                        resultado = calcular_resultado_otra_carga(
                            usuario_id, fecha_seleccionada, regla_otra_seleccionada.tipo_hora,
                            cantidad_por_dia, db,
                        )
                        tipo_item = "OTRAS"
                    nuevos_items.append((
                        resultado, fecha_seleccionada, None, None, tipo_item, False,
                    ))
            elif tipo_otra_carga.startswith("EXCEPCIONAL:"):
                try:
                    concepto_id = int(tipo_otra_carga.split(":", 1)[1])
                except ValueError as error_conversion:
                    raise CalculoHorasError("El concepto excepcional seleccionado no es válido.") from error_conversion
                resultado = calcular_resultado_concepto_excepcional(
                    usuario_id, fecha_principal, concepto_id, cantidad, db,
                )
                nuevos_items.append((resultado, fecha_principal, None, None, "EXCEPCIONAL", False))
            else:
                resultado = calcular_resultado_otra_carga(
                    usuario_id, fecha_principal, tipo_otra_carga, cantidad, db,
                )
                nuevos_items.append((resultado, fecha_principal, None, None, "OTRAS", False))
        elif tipo_registro == "HORAS":
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("Ingresá la hora de inicio y finalización.")
            for fecha_seleccionada in fechas:
                domingo_sat = es_domingo_sat(usuario_carga, fecha_seleccionada)
                if domingo_sat and not existe_domingo_activo_o_en_borrador(
                    usuario_id, fecha_seleccionada, resultados, db,
                ):
                    domingo = calcular_resultado_domingo(
                        usuario_id, fecha_seleccionada, db,
                    )
                    nuevos_items.append((
                        domingo, fecha_seleccionada, None, None, "OTRAS", False,
                    ))
                for fecha_tramo, inicio_tramo, fin_tramo in dividir_carga_en_fechas(
                    fecha_seleccionada, hora_inicio, hora_fin,
                ):
                    resultado = calcular_resultado_horas_extra(
                        usuario_id, fecha_tramo, inicio_tramo, fin_tramo, db,
                        marcar_como_franco=False,
                    )
                    if resultado.tipo_dia != "HABIL" and not domingo_sat:
                        raise CalculoHorasError(
                            f"El {fecha_tramo.strftime('%d/%m/%Y')} es feriado. "
                            "Cargalo desde Franco o feriado trabajado."
                        )
                    nuevos_items.append((
                        resultado, fecha_tramo, inicio_tramo, fin_tramo, "HORAS", False,
                    ))
                agregar_conceptos_adicionales(fecha_seleccionada)
            if str(calcular_hs_articulo or "").upper() == "SI":
                jornadas_articulo = list(articulo_jornada_fecha_lista) if isinstance(articulo_jornada_fecha_lista, (list, tuple)) else []
                fines_fecha = list(articulo_fin_fecha_lista) if isinstance(articulo_fin_fecha_lista, (list, tuple)) else []
                fines_hora = list(articulo_fin_hora_lista) if isinstance(articulo_fin_hora_lista, (list, tuple)) else []
                inicios_fecha = list(articulo_inicio_fecha_lista) if isinstance(articulo_inicio_fecha_lista, (list, tuple)) else []
                inicios_hora = list(articulo_inicio_hora_lista) if isinstance(articulo_inicio_hora_lista, (list, tuple)) else []
                if not jornadas_articulo and all((
                    articulo_fin_fecha, articulo_fin_hora,
                    articulo_inicio_fecha, articulo_inicio_hora,
                )):
                    jornadas_articulo = [fechas[0]]
                    fines_fecha = [articulo_fin_fecha]
                    fines_hora = [articulo_fin_hora]
                    inicios_fecha = [articulo_inicio_fecha]
                    inicios_hora = [articulo_inicio_hora]
                if not jornadas_articulo:
                    raise CalculoHorasError(
                        "Agregá al menos un cálculo de HS ARTICULO."
                    )
                if not (
                    len(jornadas_articulo) == len(fines_fecha) == len(fines_hora)
                    == len(inicios_fecha) == len(inicios_hora)
                ):
                    raise CalculoHorasError("Los cálculos de HS ARTICULO están incompletos.")
                if len(set(jornadas_articulo)) != len(jornadas_articulo):
                    raise CalculoHorasError("No repitas una jornada en los cálculos de HS ARTICULO.")
                fechas_horas = set(fechas)
                intervalos_articulo: set[tuple[date, time, date, time]] = set()
                for jornada, fecha_fin, hora_fin_articulo, fecha_inicio, hora_inicio_articulo in zip(
                    jornadas_articulo, fines_fecha, fines_hora,
                    inicios_fecha, inicios_hora,
                ):
                    if jornada not in fechas_horas:
                        raise CalculoHorasError(
                            "Una jornada de HS ARTICULO no pertenece a las fechas de horas extras."
                        )
                    intervalo = (
                        fecha_fin, hora_fin_articulo,
                        fecha_inicio, hora_inicio_articulo,
                    )
                    if intervalo in intervalos_articulo:
                        raise CalculoHorasError("No repitas el mismo intervalo de descanso.")
                    intervalos_articulo.add(intervalo)
                    calculo_articulo = calcular_descanso_articulo(*intervalo)
                    if calculo_articulo is None:
                        continue
                    fecha_articulo, cantidad_articulo, observacion_articulo = calculo_articulo
                    if observacion_articulo in observaciones_resultados:
                        raise CalculoHorasError(
                            "El borrador ya contiene HS ARTICULO para uno de los intervalos informados."
                        )
                    validar_fecha_carga_usuario(usuario_carga, fecha_articulo)
                    existente_articulo = db.scalar(
                        select(HoraExtra.id).where(
                            HoraExtra.usuario_id == usuario_id,
                            HoraExtra.tipo_registro == "OTRAS",
                            func.upper(HoraExtra.tipo_hora) == "HS ARTICULO",
                            HoraExtra.observaciones == observacion_articulo,
                            HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
                        ).limit(1)
                    )
                    if existente_articulo is not None:
                        raise CalculoHorasError(
                            "Ya existe una carga de HS ARTICULO para uno de los intervalos informados."
                        )
                    articulo = calcular_resultado_otra_carga(
                        usuario_id, fecha_articulo, "HS ARTICULO", cantidad_articulo, db,
                    )
                    nuevos_items.append((
                        articulo, fecha_articulo, None, None, "OTRAS", False,
                        observacion_articulo,
                    ))
        elif tipo_registro == "DIA_TRABAJADO":
            if jornada_hora_inicio is None or jornada_hora_fin is None:
                raise CalculoHorasError("Ingresá la hora de inicio y finalización de la jornada.")
            tramos_jornada = dividir_carga_en_fechas(
                fecha_principal, jornada_hora_inicio, jornada_hora_fin,
            )
            horas_jornada = sum(
                (calcular_horas_totales(inicio, fin) for _, inicio, fin in tramos_jornada),
                start=Decimal("0.00"),
            )
            domingo_inicio = es_domingo_sat(usuario_carga, fecha_principal)
            es_feriado = clasificar_tipo_dia(fecha_principal, db) == "FERIADO"
            franco_declarado = not es_feriado and not domingo_inicio
            jornada_corta = (
                convenio_carga not in {"FC", "MONOTRIBUTISTA"}
                and not domingo_inicio
                and horas_jornada < Decimal("4.00")
            )
            es_franco = False
            if jornada_corta:
                for fecha_tramo, inicio_tramo, fin_tramo in tramos_jornada:
                    tramo_franco = franco_declarado
                    resultado = calcular_resultado_horas_extra(
                        usuario_id, fecha_tramo, inicio_tramo, fin_tramo, db,
                        marcar_como_franco=tramo_franco,
                    )
                    if "100" not in str(resultado.tipo_hora or ""):
                        raise CalculoHorasError(
                            "La regla de franco o feriado debe corresponder a horas al 100%."
                        )
                    nuevos_items.append((
                        resultado, fecha_tramo, inicio_tramo, fin_tramo,
                        "HORAS", tramo_franco,
                    ))
            elif domingo_inicio:
                if existe_domingo_activo_o_en_borrador(
                    usuario_id, fecha_principal, resultados, db,
                ):
                    raise CalculoHorasError("Ya existe una carga de domingo para esa fecha.")
                dia_trabajado = calcular_resultado_domingo(
                    usuario_id, fecha_principal, db,
                    hora_inicio=jornada_hora_inicio,
                    hora_fin=jornada_hora_fin,
                )
                es_franco = False
                nuevos_items.append((
                    dia_trabajado, fecha_principal, jornada_hora_inicio,
                    jornada_hora_fin, "OTRAS", False,
                ))
            else:
                dia_trabajado = calcular_resultado_dia_trabajado(
                    usuario_id, fecha_principal, db,
                    marcar_como_franco=franco_declarado,
                    hora_inicio=jornada_hora_inicio,
                    hora_fin=jornada_hora_fin,
                )
                es_franco = dia_trabajado.tipo_dia == "FRANCO"
                reintegro_permitido = False
                try:
                    calcular_resultado_reintegro(
                        usuario_id, fecha_principal, db, Decimal("1.00"),
                    )
                    reintegro_permitido = True
                except CalculoHorasError:
                    pass
                fecha_descanso = None
                if seleccion_reintegro_solicitada == "NO" and reintegro_permitido:
                    if fecha_descanso_compensatorio is None:
                        raise CalculoHorasError(
                            "Indicá qué día vas a tomar el franco o feriado trabajado."
                        )
                    fecha_descanso = fecha_descanso_compensatorio
                nuevos_items.append((
                    dia_trabajado, fecha_principal, jornada_hora_inicio,
                    jornada_hora_fin, "DIA_TRABAJADO", es_franco,
                    observacion_limpia, fecha_descanso,
                ))
            fechas_jornada = {
                tramo[0] for tramo in tramos_jornada
            }
            domingos_sat = sorted(
                fecha_jornada for fecha_jornada in fechas_jornada
                if es_domingo_sat(usuario_carga, fecha_jornada)
                and not (domingo_inicio and fecha_jornada == fecha_principal)
            )
            for fecha_domingo in domingos_sat:
                if existe_domingo_activo_o_en_borrador(
                    usuario_id, fecha_domingo, resultados, db,
                ):
                    raise CalculoHorasError(
                        f"Ya existe una carga de domingo para el {fecha_domingo.strftime('%d/%m/%Y')}."
                    )
                domingo = calcular_resultado_domingo(usuario_id, fecha_domingo, db)
                nuevos_items.append((
                    domingo, fecha_domingo, None, None, "OTRAS", False,
                ))
            if incluye_horas_extra == "on" and jornada_corta:
                raise CalculoHorasError(
                    "La jornada informada ya se registra íntegramente como horas extras al 100%."
                )
            if incluye_horas_extra == "on" and not jornada_corta:
                if cantidad_horas_extra is None:
                    raise CalculoHorasError(
                        "Ingresá la cantidad de horas extras."
                    )
                if (
                    cantidad_horas_extra <= 0
                    or cantidad_horas_extra % Decimal("0.5") != 0
                    or cantidad_horas_extra > horas_jornada
                ):
                    raise CalculoHorasError(
                        "Las horas extras deben ingresarse de a 0,5 y no pueden superar la jornada."
                    )
                inicio_dt = datetime.combine(fecha_principal, jornada_hora_inicio)
                fin_dt = datetime.combine(fecha_principal, jornada_hora_fin)
                if fin_dt <= inicio_dt:
                    fin_dt += timedelta(days=1)
                inicio_extra_dt = fin_dt - timedelta(
                    minutes=int(cantidad_horas_extra * Decimal("60"))
                )
                for fecha_tramo, inicio_tramo, fin_tramo in dividir_carga_en_fechas(
                    inicio_extra_dt.date(), inicio_extra_dt.time(), fin_dt.time(),
                ):
                    tramo_franco = es_franco and fecha_tramo == fecha_principal
                    resultado = calcular_resultado_horas_extra(
                        usuario_id, fecha_tramo, inicio_tramo, fin_tramo, db,
                        marcar_como_franco=tramo_franco,
                    )
                    nuevos_items.append((
                        resultado, fecha_tramo, inicio_tramo, fin_tramo,
                        "HORAS", tramo_franco,
                    ))
            seleccion_reintegro = seleccion_reintegro_solicitada
            if seleccion_reintegro in {"ON", "MEDIO", "COMPLETO"}:
                if jornada_corta:
                    raise CalculoHorasError(
                        "La jornada informada no genera reintegro del día."
                    )
                if domingo_inicio:
                    raise CalculoHorasError("El trabajo en domingo no permite solicitar reintegro.")
                cantidad_reintegro = (
                    Decimal("0.50") if seleccion_reintegro == "MEDIO"
                    else Decimal("1.00")
                )
                reintegro = calcular_resultado_reintegro(
                    usuario_id, fecha_principal, db, cantidad_reintegro,
                )
                nuevos_items.append((
                    reintegro, fecha_principal, None, None, "REINTEGRO", es_franco,
                ))
            agregar_conceptos_adicionales(fecha_principal)
        else:
            raise CalculoHorasError("El tipo de registro seleccionado no es válido.")

        nuevos_resultados = [item[0] for item in nuevos_items]
        resultados.extend(nuevos_resultados)
        evitar_nocturnidad_duplicada_en_jornada(resultados)
        observaciones_nuevas = [
            item[6] if len(item) > 6 else observacion_limpia
            for item in nuevos_items
        ]
        observaciones_resultados.extend(observaciones_nuevas)
        cargas_codificadas.extend(
            encode_carga(
                usuario_id, fecha_tramo, inicio_tramo, fin_tramo,
                item[6] if len(item) > 6 else observacion_limpia,
                tipo_item,
                marcar_como_franco=es_franco_item,
                tipo_hora=resultado.tipo_hora if tipo_item == "OTRAS" else None,
                cantidad=resultado.cantidad if tipo_item in {"OTRAS", "EXCEPCIONAL", "REINTEGRO"} else None,
                concepto_excepcional_id=resultado.concepto_excepcional_id,
                concepto_excepcional_nombre=resultado.concepto_excepcional_nombre,
                fecha_descanso_compensatorio=(
                    item[7] if len(item) > 7 and tipo_item == "DIA_TRABAJADO" else None
                ),
            )
            for item in nuevos_items
            for resultado, fecha_tramo, inicio_tramo, fin_tramo, tipo_item, es_franco_item in [item[:6]]
        )
        selecciones_reintegro.extend(
            "SI" if item[4] == "REINTEGRO" else "NO"
            for item in nuevos_items
        )
    except CalculoHorasError as exc:
        error = str(exc)

    return renderizar_borrador_solicitud(
        request, db, resultados, cargas_codificadas, selecciones_reintegro,
        observaciones_resultados, error=error, destino=destino,
        aviso=aviso,
        status_code=422 if error and not resultados else 200,
    )


@app.post("/solicitudes/confirmar")
def confirmar_solicitud(
    request: Request,
    cargas: Annotated[list[str] | None, Form()] = None,
    reintegros: Annotated[list[str] | None, Form()] = None,
    destino: Annotated[str, Form()] = "propio",
    db: Session = Depends(get_db),
):
    cargas = list(cargas or [])
    if not cargas:
        return RedirectResponse("/solicitudes/nueva", status_code=303)
    selecciones = list(reintegros or [])
    if len(selecciones) != len(cargas):
        return RedirectResponse(
            "/solicitudes/nueva?error="
            + quote("La selección de reintegros es inconsistente."),
            status_code=303,
        )

    try:
        carga_administrativa = destino == "rrhh"
        asignacion_activa = obtener_asignacion_activa(request, db)
        if carga_administrativa and asignacion_activa.rol.upper() not in {"ADMIN", "RRHH"}:
            raise CalculoHorasError("Sólo RRHH o ADMIN pueden confirmar cargas administrativas.")
        autorizador_id = asignacion_activa.usuario_id if carga_administrativa else None
        fecha_resolucion = datetime.now() if carga_administrativa else None
        estado_inicial = "APROBADA" if carga_administrativa else "PENDIENTE"
        ids_permitidos = ids_habilitados_para_confirmar(request, db, destino)
        datos_cargas = [decode_carga(carga) for carga in cargas]
        cierres_medianoche = {
            (datos[0], datos[1])
            for datos in datos_cargas
            if datos[5] == "HORAS" and datos[3] == time(0, 0)
        }
        resultados_confirmacion = []
        conceptos_unicos = set()
        for indice, datos in enumerate(datos_cargas):
            if datos[0] not in ids_permitidos:
                raise CalculoHorasError("Una de las cargas contiene un usuario no autorizado.")
            if not datos[4]:
                raise CalculoHorasError("Todas las cargas deben tener una observación.")
            usuario_carga = db.get(Usuario, datos[0])
            if usuario_carga is None:
                raise CalculoHorasError("El usuario de una carga no existe.")
            convenio_carga = str(usuario_carga.convenio or "").strip().upper()
            if (
                convenio_carga in {"FC", "MONOTRIBUTISTA"}
                and selecciones[indice] == "SI"
            ):
                raise CalculoHorasError(
                    f"{convenio_carga} no permite solicitar reintegro."
                )
            fecha_a_validar = datos[1]
            if (
                datos[5] == "HORAS"
                and datos[2] == time(0, 0)
                and (datos[0], datos[1] - timedelta(days=1)) in cierres_medianoche
            ):
                fecha_a_validar -= timedelta(days=1)
            validar_fecha_carga_usuario(usuario_carga, fecha_a_validar)
            resultado = calcular_carga_codificada(
                datos, db, permitir_fc_sin_asignacion=carga_administrativa,
            )
            if (
                resultado.tipo_registro == "OTRAS"
                and str(resultado.tipo_hora or "").upper()
                in TIPOS_NOCTURNAS_EN_JORNADA
            ):
                validar_nocturnidad_sin_superposicion(
                    resultado.usuario_id, resultado.fecha,
                    resultado.hora_inicio, resultado.hora_fin, db,
                    resultados_confirmacion,
                )
            resultados_confirmacion.append(resultado)
            es_domingo = (
                resultado.tipo_registro == "OTRAS"
                and str(resultado.tipo_hora or "").upper() == "DOMINGO"
            )
            if resultado.tipo_registro in {"DIA_TRABAJADO", "REINTEGRO"} or es_domingo:
                concepto_registro = "DOMINGO" if es_domingo else resultado.tipo_registro
                clave = (resultado.usuario_id, resultado.fecha, concepto_registro)
                if clave in conceptos_unicos:
                    raise CalculoHorasError(
                        "El borrador contiene más de un aviso o reintegro para la misma fecha."
                    )
                conceptos_unicos.add(clave)
                existente = db.scalar(
                    select(HoraExtra.id).where(
                        HoraExtra.usuario_id == resultado.usuario_id,
                        HoraExtra.fecha == resultado.fecha,
                        (
                            func.upper(HoraExtra.tipo_hora) == "DOMINGO"
                            if es_domingo
                            else HoraExtra.tipo_registro == resultado.tipo_registro
                        ),
                        HoraExtra.estado.in_(("PENDIENTE", "APROBADA")),
                    ).limit(1)
                )
                if existente is not None:
                    concepto = (
                        "carga de domingo" if es_domingo
                        else "aviso de día trabajado" if resultado.tipo_registro == "DIA_TRABAJADO"
                        else "pedido de reintegro"
                    )
                    raise CalculoHorasError(
                        f"Ya existe un {concepto} para esa fecha."
                    )

        reintegros_en_borrador = {
            (resultado.usuario_id, resultado.fecha)
            for resultado in resultados_confirmacion
            if resultado.tipo_registro == "REINTEGRO"
        }
        for datos, resultado in zip(datos_cargas, resultados_confirmacion):
            if resultado.tipo_registro != "DIA_TRABAJADO":
                continue
            tiene_reintegro = (
                resultado.usuario_id, resultado.fecha
            ) in reintegros_en_borrador
            reintegro_permitido = False
            try:
                calcular_resultado_reintegro(
                    resultado.usuario_id, resultado.fecha, db, Decimal("1.00"),
                )
                reintegro_permitido = True
            except CalculoHorasError:
                pass
            if reintegro_permitido and not tiene_reintegro and datos[11] is None:
                raise CalculoHorasError(
                    "Indicá qué día vas a tomar el franco o feriado trabajado."
                )

        evitar_nocturnidad_duplicada_en_jornada(resultados_confirmacion)

        fecha_carga = datetime.now()
        registros_horas_creados: list[HoraExtra] = []
        for indice, (datos, resultado) in enumerate(zip(datos_cargas, resultados_confirmacion)):
            solicita_reintegro = (
                resultado.tipo_registro == "REINTEGRO"
                or (resultado.permite_reintegro and selecciones[indice] == "SI")
            )
            registro = HoraExtra(
                usuario_id=resultado.usuario_id,
                fecha=resultado.fecha,
                hora_inicio=resultado.hora_inicio,
                hora_fin=resultado.hora_fin,
                horas_totales=resultado.horas_totales,
                cantidad=resultado.cantidad,
                concepto_excepcional_id=resultado.concepto_excepcional_id,
                concepto_excepcional_nombre=resultado.concepto_excepcional_nombre,
                tipo_dia=resultado.tipo_dia,
                tipo_hora=resultado.tipo_hora,
                horas_nocturnas=resultado.horas_nocturnas,
                solicita_reintegro=solicita_reintegro,
                fecha_descanso_compensatorio=(
                    datos[11] if resultado.tipo_registro == "DIA_TRABAJADO" else None
                ),
                observaciones=datos[4],
                estado=estado_inicial,
                aprobado_por=autorizador_id,
                fecha_carga=fecha_carga,
                fecha_resolucion=fecha_resolucion,
                tipo_registro=resultado.tipo_registro,
            )
            db.add(registro)
            if resultado.tipo_registro == "HORAS":
                registros_horas_creados.append(registro)
        db.flush()
        jornadas_a_recalcular: set[tuple[int, date]] = set()
        for tramo in registros_horas_creados:
            jornadas_a_recalcular.add((
                tramo.usuario_id,
                fecha_jornada_de_hora_estado(tramo, db, estado_inicial),
            ))
            if tramo.hora_fin == time(0, 0):
                jornadas_a_recalcular.add((
                    tramo.usuario_id,
                    tramo.fecha + timedelta(days=1),
                ))
        for usuario_id, fecha_jornada in jornadas_a_recalcular:
            recalcular_conceptos_jornada_estado(
                usuario_id,
                fecha_jornada,
                db,
                estado_inicial,
                aprobado_por=autorizador_id,
                fecha_resolucion=fecha_resolucion,
            )
        for usuario_id, fecha_switcher in {
            (resultado.usuario_id, resultado.fecha)
            for resultado in resultados_confirmacion
        }:
            recalcular_switcher_nocturnidad_estado(
                usuario_id, fecha_switcher, db, estado_inicial,
                aprobado_por=autorizador_id,
                fecha_resolucion=fecha_resolucion,
            )
        db.commit()
    except (CalculoHorasError, SQLAlchemyError) as error:
        db.rollback()
        mensaje = str(error) if isinstance(error, CalculoHorasError) else "No se pudo guardar la solicitud en la base de datos."
        try:
            datos_borrador = [decode_carga(carga) for carga in cargas]
            resultados_borrador = [
                calcular_carga_codificada(
                    datos, db, permitir_fc_sin_asignacion=destino == "rrhh",
                ) for datos in datos_borrador
            ]
            return renderizar_borrador_solicitud(
                request,
                db,
                resultados_borrador,
                cargas,
                selecciones,
                [datos[4] for datos in datos_borrador],
                error=mensaje,
                destino=destino,
                status_code=422,
            )
        except (CalculoHorasError, SQLAlchemyError):
            db.rollback()
            return RedirectResponse(
                "/solicitudes/nueva?destino=" + quote(destino) + "&error=" + quote(mensaje),
                status_code=303,
            )
    if destino == "rrhh":
        return RedirectResponse(
            f"/rrhh/horas-extras?guardado={len(cargas)}",
            status_code=303,
        )
    return RedirectResponse(f"/horas?guardado={len(cargas)}", status_code=303)


@app.get("/configuracion/precargas-autorizadas", response_class=HTMLResponse)
def administrar_precargas_autorizadas(
    request: Request,
    guardado: str | None = None,
    error: str | None = None,
    editar: int | None = None,
    db: Session = Depends(get_db),
):
    asignacion = obtener_asignacion_activa(request, db)
    if asignacion.rol.upper() not in {"ADMIN", "RRHH"}:
        raise HTTPException(
            status_code=403,
            detail="Sólo RRHH o ADMIN pueden administrar precargas autorizadas.",
        )
    usuarios = list(db.scalars(
        select(Usuario)
        .options(selectinload(Usuario.tipo_contratacion))
        .where(Usuario.status.is_(True), Usuario.convenio.is_not(None))
        .order_by(Usuario.apellido, Usuario.nombre)
    ))
    reglas = list(db.execute(
        select(ReglaHora.convenio, ReglaHora.tipo_hora)
        .where(ReglaHora.tipo_hora.is_not(None))
        .distinct()
        .order_by(ReglaHora.convenio, ReglaHora.tipo_hora)
    ))
    conceptos_por_convenio: dict[str, list[str]] = {}
    for convenio, tipo_hora in reglas:
        convenio_normalizado = str(convenio or "").strip().upper()
        concepto_normalizado = str(tipo_hora or "").strip().upper()
        if convenio_normalizado and concepto_normalizado:
            conceptos_por_convenio.setdefault(convenio_normalizado, []).append(
                concepto_normalizado,
            )
    precargas = list(db.scalars(
        select(HoraExtra)
        .join(HoraExtra.usuario)
        .options(
            selectinload(HoraExtra.usuario).selectinload(Usuario.tipo_contratacion)
        )
        .where(HoraExtra.tipo_registro == TIPO_PRECARGA_AUTORIZADA)
        .order_by(HoraExtra.fecha.desc(), Usuario.apellido, Usuario.nombre)
    ))
    grupos_pack: dict[tuple, dict] = {}
    for precarga in precargas:
        clave = (
            precarga.usuario_id,
            str(precarga.tipo_hora or "").upper(),
            precarga.cantidad,
            precarga.observaciones,
            precarga.fecha_resolucion or precarga.id,
        )
        grupo = grupos_pack.setdefault(clave, {
            "id": precarga.id,
            "usuario": precarga.usuario,
            "tipo_hora": precarga.tipo_hora,
            "cantidad": precarga.cantidad,
            "observaciones": precarga.observaciones,
            "fecha_desde": precarga.fecha,
            "fecha_hasta": precarga.fecha,
            "filas": [],
        })
        grupo["filas"].append(precarga)
        grupo["fecha_desde"] = min(grupo["fecha_desde"], precarga.fecha)
        grupo["fecha_hasta"] = max(grupo["fecha_hasta"], precarga.fecha)
    packs = list(grupos_pack.values())
    pack_edicion = next((item for item in packs if item["id"] == editar), None)
    if editar is not None and pack_edicion is None:
        error = "El pack seleccionado no existe."
    periodos_edicion = []
    conceptos_edicion = []
    if pack_edicion is not None:
        usuario_edicion = pack_edicion["usuario"]
        contratacion = (
            usuario_edicion.tipo_contratacion.contratacion.strip().lower()
            if usuario_edicion.tipo_contratacion is not None else ""
        )
        tipo_periodo = "nomina" if contratacion in {"nómina", "nomina"} else "calendario"
        periodos_edicion = opciones_periodos_precarga()[tipo_periodo]
        conceptos_edicion = conceptos_por_convenio.get(
            str(usuario_edicion.convenio or "").upper(), [],
        )
    return templates.TemplateResponse(request, "precargas_autorizadas.html", {
        "active_page": "precargas_autorizadas",
        "usuarios": usuarios,
        "conceptos_por_convenio": conceptos_por_convenio,
        "packs": packs,
        "pack_edicion": pack_edicion,
        "periodos_edicion": periodos_edicion,
        "conceptos_edicion": conceptos_edicion,
        "fecha_hoy": date.today().isoformat(),
        "periodos_precarga": opciones_periodos_precarga(),
        "guardado": guardado,
        "error": error,
    })


@app.post("/configuracion/precargas-autorizadas")
def crear_pack_precargas_autorizadas(
    request: Request,
    usuario_id: Annotated[int, Form()],
    tipo_hora: Annotated[str, Form()],
    cantidad: Annotated[Decimal, Form()],
    periodo_desde: Annotated[date, Form()],
    periodo_hasta: Annotated[date, Form()],
    observacion: Annotated[str, Form()],
    db: Session = Depends(get_db),
):
    asignacion = obtener_asignacion_activa(request, db)
    if asignacion.rol.upper() not in {"ADMIN", "RRHH"}:
        raise HTTPException(
            status_code=403,
            detail="Sólo RRHH o ADMIN pueden administrar precargas autorizadas.",
        )
    usuario = db.scalar(
        select(Usuario)
        .options(selectinload(Usuario.tipo_contratacion))
        .where(Usuario.id == usuario_id, Usuario.status.is_(True))
    )
    if usuario is None or not usuario.convenio:
        raise HTTPException(status_code=422, detail="El usuario no está habilitado.")
    inicio_desde, _ = periodo_precarga_usuario(usuario, periodo_desde)
    inicio_hasta, fecha_hasta_real = periodo_precarga_usuario(usuario, periodo_hasta)
    if periodo_desde != inicio_desde or periodo_hasta != inicio_hasta:
        raise HTTPException(
            status_code=422,
            detail="Seleccioná períodos completos válidos para la contratación del usuario.",
        )
    if periodo_hasta < periodo_desde:
        raise HTTPException(
            status_code=422,
            detail="El período final no puede ser anterior al período inicial.",
        )
    if len(periodos_vigencia_precarga(usuario, periodo_desde, fecha_hasta_real)) > 12:
        raise HTTPException(
            status_code=422,
            detail="Un pack puede abarcar como máximo 12 períodos.",
        )
    if cantidad <= 0:
        raise HTTPException(status_code=422, detail="La cantidad debe ser mayor que cero.")
    observacion_limpia = observacion.strip()
    if not observacion_limpia:
        raise HTTPException(status_code=422, detail="El motivo del pack es obligatorio.")
    tipo_hora_normalizado = tipo_hora.strip().upper()
    concepto_habilitado = db.scalar(
        select(ReglaHora.id).where(
            func.upper(ReglaHora.convenio) == usuario.convenio.strip().upper(),
            func.upper(ReglaHora.tipo_hora) == tipo_hora_normalizado,
        ).limit(1)
    )
    if concepto_habilitado is None:
        raise HTTPException(
            status_code=422,
            detail="El concepto no está disponible para el convenio del usuario.",
        )
    creadas = generar_pack_precargas_autorizadas(
        usuario, tipo_hora_normalizado, cantidad, periodo_desde, fecha_hasta_real,
        observacion_limpia, asignacion.usuario_id, db,
    )
    db.commit()
    mensaje = f"pack creado: {creadas} precargas autorizadas generadas"
    return RedirectResponse(
        "/configuracion/precargas-autorizadas?guardado=" + quote(mensaje),
        status_code=303,
    )


def filas_de_pack_precarga(pack_id: int, db: Session) -> list[HoraExtra]:
    representante = db.get(HoraExtra, pack_id)
    if (
        representante is None
        or representante.tipo_registro != TIPO_PRECARGA_AUTORIZADA
    ):
        return []
    condiciones = [
        HoraExtra.usuario_id == representante.usuario_id,
        HoraExtra.tipo_registro == TIPO_PRECARGA_AUTORIZADA,
        HoraExtra.tipo_hora == representante.tipo_hora,
        HoraExtra.cantidad == representante.cantidad,
        HoraExtra.observaciones == representante.observaciones,
    ]
    if representante.fecha_resolucion is None:
        condiciones.append(HoraExtra.id == representante.id)
    else:
        condiciones.append(HoraExtra.fecha_resolucion == representante.fecha_resolucion)
    return list(db.scalars(
        select(HoraExtra).where(*condiciones).order_by(HoraExtra.fecha)
    ))


@app.post("/configuracion/precargas-autorizadas/{pack_id}/editar")
def editar_pack_precargas_autorizadas(
    pack_id: int,
    request: Request,
    tipo_hora: Annotated[str, Form()],
    cantidad: Annotated[Decimal, Form()],
    periodo_desde: Annotated[date, Form()],
    periodo_hasta: Annotated[date, Form()],
    observacion: Annotated[str, Form()],
    db: Session = Depends(get_db),
):
    asignacion = obtener_asignacion_activa(request, db)
    if asignacion.rol.upper() not in {"ADMIN", "RRHH"}:
        raise HTTPException(status_code=403, detail="No tenés permiso para editar packs.")
    filas = filas_de_pack_precarga(pack_id, db)
    if not filas:
        raise HTTPException(status_code=404, detail="El pack no existe.")
    usuario = db.scalar(
        select(Usuario)
        .options(selectinload(Usuario.tipo_contratacion))
        .where(Usuario.id == filas[0].usuario_id)
    )
    if usuario is None or not usuario.convenio:
        raise HTTPException(status_code=422, detail="El usuario no está habilitado.")
    inicio_desde, _ = periodo_precarga_usuario(usuario, periodo_desde)
    inicio_hasta, fecha_hasta_real = periodo_precarga_usuario(usuario, periodo_hasta)
    if periodo_desde != inicio_desde or periodo_hasta != inicio_hasta:
        raise HTTPException(status_code=422, detail="Los períodos seleccionados no son válidos.")
    if periodo_hasta < periodo_desde or cantidad <= 0 or not observacion.strip():
        raise HTTPException(status_code=422, detail="Revisá el rango, la cantidad y el motivo.")
    if len(periodos_vigencia_precarga(usuario, periodo_desde, fecha_hasta_real)) > 12:
        raise HTTPException(
            status_code=422,
            detail="Un pack puede abarcar como máximo 12 períodos.",
        )
    tipo_hora_normalizado = tipo_hora.strip().upper()
    if db.scalar(select(ReglaHora.id).where(
        func.upper(ReglaHora.convenio) == usuario.convenio.strip().upper(),
        func.upper(ReglaHora.tipo_hora) == tipo_hora_normalizado,
    ).limit(1)) is None:
        raise HTTPException(status_code=422, detail="El concepto no corresponde al convenio.")
    for fila in filas:
        db.delete(fila)
    db.flush()
    creadas = generar_pack_precargas_autorizadas(
        usuario, tipo_hora_normalizado, cantidad, periodo_desde,
        fecha_hasta_real, observacion.strip(), asignacion.usuario_id, db,
    )
    db.commit()
    return RedirectResponse(
        "/configuracion/precargas-autorizadas?guardado="
        + quote(f"pack actualizado: {creadas} precargas generadas"),
        status_code=303,
    )


@app.post("/configuracion/precargas-autorizadas/{pack_id}/eliminar")
def eliminar_pack_precargas_autorizadas(
    pack_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    asignacion = obtener_asignacion_activa(request, db)
    if asignacion.rol.upper() not in {"ADMIN", "RRHH"}:
        raise HTTPException(status_code=403, detail="No tenés permiso para eliminar packs.")
    filas = filas_de_pack_precarga(pack_id, db)
    if not filas:
        raise HTTPException(status_code=404, detail="El pack no existe.")
    for fila in filas:
        db.delete(fila)
    cantidad_eliminada = len(filas)
    db.commit()
    return RedirectResponse(
        "/configuracion/precargas-autorizadas?guardado="
        + quote(f"pack eliminado: {cantidad_eliminada} precargas removidas"),
        status_code=303,
    )


@app.get("/configuracion/conceptos-excepcionales", response_class=HTMLResponse)
def administrar_conceptos_excepcionales(
    request: Request,
    guardado: str | None = None,
    error: str | None = None,
    db: Session = Depends(get_db),
):
    asignacion, administra_catalogo = acceso_conceptos_excepcionales(request, db)
    usuarios = usuarios_asignables_excepciones(asignacion, administra_catalogo, db)
    ids_usuarios = {usuario.id for usuario in usuarios}
    conceptos = list(db.scalars(
        select(ConceptoExcepcional).order_by(
            ConceptoExcepcional.activo.desc(), ConceptoExcepcional.nombre,
        )
    ))
    tipos_hora_destino = list(db.scalars(
        select(ReglaHora.tipo_hora)
        .where(ReglaHora.tipo_hora.is_not(None))
        .distinct()
        .order_by(ReglaHora.tipo_hora)
    ))
    consulta_asignaciones = (
        select(UsuarioConceptoExcepcional)
        .options(
            selectinload(UsuarioConceptoExcepcional.usuario),
            selectinload(UsuarioConceptoExcepcional.concepto),
            selectinload(UsuarioConceptoExcepcional.asignador),
        )
        .order_by(
            UsuarioConceptoExcepcional.activo.desc(),
            UsuarioConceptoExcepcional.fecha_asignacion.desc(),
        )
    )
    if not administra_catalogo:
        consulta_asignaciones = consulta_asignaciones.where(
            UsuarioConceptoExcepcional.usuario_id.in_(ids_usuarios or {-1})
        )
    asignaciones = list(db.scalars(consulta_asignaciones))
    por_usuario: dict[int, dict] = {}
    por_concepto: dict[int, dict] = {}
    for item in asignaciones:
        grupo_usuario = por_usuario.setdefault(
            item.usuario_id, {"entidad": item.usuario, "asignaciones": []},
        )
        grupo_usuario["asignaciones"].append(item)
        grupo_concepto = por_concepto.setdefault(
            item.concepto_excepcional_id,
            {"entidad": item.concepto, "asignaciones": []},
        )
        grupo_concepto["asignaciones"].append(item)
    asignaciones_por_usuario = sorted(
        por_usuario.values(),
        key=lambda grupo: (
            grupo["entidad"].apellido.lower(), grupo["entidad"].nombre.lower(),
        ),
    )
    asignaciones_por_concepto = sorted(
        por_concepto.values(), key=lambda grupo: grupo["entidad"].nombre.lower(),
    )
    return templates.TemplateResponse(request, "conceptos_excepcionales.html", {
        "active_page": "conceptos_excepcionales",
        "conceptos": conceptos,
        "conceptos_activos": [
            item for item in conceptos
            if item.activo and item.tipo != "PRECARGA_PERIODO"
        ],
        "tipos_concepto": TIPOS_CONCEPTO_EXCEPCIONAL_CONFIGURABLES,
        "tipos_hora_destino": tipos_hora_destino,
        "usuarios": usuarios,
        "asignaciones": asignaciones,
        "asignaciones_por_usuario": asignaciones_por_usuario,
        "asignaciones_por_concepto": asignaciones_por_concepto,
        "administra_catalogo": administra_catalogo,
        "guardado": guardado,
        "error": error,
        "fecha_hoy": date.today().isoformat(),
        "solo_precargas": False,
    })


@app.post("/configuracion/conceptos-excepcionales")
def crear_concepto_excepcional(
    request: Request,
    nombre: Annotated[str, Form()],
    codigo: Annotated[str | None, Form()] = None,
    tipo: Annotated[str, Form()] = "CARGA_MANUAL",
    descripcion: Annotated[str | None, Form()] = None,
    cantidad_predeterminada: Annotated[Decimal | None, Form()] = None,
    tipo_hora_destino: Annotated[str | None, Form()] = None,
    requiere_observacion: Annotated[str | None, Form()] = None,
    requiere_autorizacion: Annotated[str | None, Form()] = None,
    vista: str | None = None,
    db: Session = Depends(get_db),
):
    _, administra_catalogo = acceso_conceptos_excepcionales(request, db)
    if not administra_catalogo:
        raise HTTPException(status_code=403, detail="Sólo RRHH o ADMIN pueden crear conceptos.")
    if tipo.strip().upper() == TIPO_SWITCHER_NOCTURNIDAD:
        raise HTTPException(
            status_code=409,
            detail="El concepto switcher_nocturnidad es administrado por el sistema.",
        )
    if tipo.strip().upper() == "PRECARGA_PERIODO":
        raise HTTPException(
            status_code=409,
            detail="Las precargas se arman con conceptos generales del convenio.",
        )
    concepto = ConceptoExcepcional(
        activo=True,
        creado_por=(request.session.get("user") or {}).get("id"),
    )
    try:
        configurar_concepto_excepcional(
            concepto, nombre, codigo, tipo, descripcion,
            cantidad_predeterminada, tipo_hora_destino,
            requiere_observacion == "on", requiere_autorizacion == "on",
        )
        validar_destino_precarga(concepto, db)
    except CalculoHorasError as error_configuracion:
        return RedirectResponse(
            "/configuracion/conceptos-excepcionales?"
            + ("vista=precargas&" if vista == "precargas" else "") + "error="
            + quote(str(error_configuracion)), status_code=303,
        )
    db.add(concepto)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse("/configuracion/conceptos-excepcionales?" + ("vista=precargas&" if vista == "precargas" else "") + "error=Ya+existe+un+concepto+con+ese+nombre+o+código", status_code=303)
    return RedirectResponse("/configuracion/conceptos-excepcionales?" + ("vista=precargas&" if vista == "precargas" else "") + "guardado=concepto+creado", status_code=303)


@app.post("/configuracion/conceptos-excepcionales/{concepto_id:int}")
def actualizar_concepto_excepcional(
    concepto_id: int,
    request: Request,
    nombre: Annotated[str, Form()],
    codigo: Annotated[str | None, Form()] = None,
    tipo: Annotated[str, Form()] = "CARGA_MANUAL",
    descripcion: Annotated[str | None, Form()] = None,
    cantidad_predeterminada: Annotated[Decimal | None, Form()] = None,
    tipo_hora_destino: Annotated[str | None, Form()] = None,
    requiere_observacion: Annotated[str | None, Form()] = None,
    requiere_autorizacion: Annotated[str | None, Form()] = None,
    activo: Annotated[str | None, Form()] = None,
    vista: str | None = None,
    db: Session = Depends(get_db),
):
    _, administra_catalogo = acceso_conceptos_excepcionales(request, db)
    if not administra_catalogo:
        raise HTTPException(status_code=403, detail="Sólo RRHH o ADMIN pueden modificar conceptos.")
    concepto = db.get(ConceptoExcepcional, concepto_id)
    if concepto is None:
        raise HTTPException(status_code=404, detail="El concepto no existe.")
    if concepto.tipo == TIPO_SWITCHER_NOCTURNIDAD:
        raise HTTPException(
            status_code=409,
            detail="El concepto switcher_nocturnidad es administrado por el sistema.",
        )
    try:
        configurar_concepto_excepcional(
            concepto, nombre, codigo, tipo, descripcion,
            cantidad_predeterminada, tipo_hora_destino,
            requiere_observacion == "on", requiere_autorizacion == "on",
        )
        validar_destino_precarga(concepto, db)
        if concepto.tipo == "PRECARGA_PERIODO" and db.scalar(
            select(UsuarioConceptoExcepcional.id_asignacion).where(
                UsuarioConceptoExcepcional.concepto_excepcional_id == concepto.id_concepto_excepcional,
                UsuarioConceptoExcepcional.activo.is_(True),
                UsuarioConceptoExcepcional.cantidad_asignada.is_(None),
            ).limit(1)
        ) is not None:
            raise CalculoHorasError(
                "Desactivá y volvé a crear las asignaciones vigentes con una cantidad por período."
            )
    except CalculoHorasError as error_configuracion:
        db.rollback()
        return RedirectResponse(
            "/configuracion/conceptos-excepcionales?"
            + ("vista=precargas&" if vista == "precargas" else "") + "error="
            + quote(str(error_configuracion)), status_code=303,
        )
    concepto.activo = activo == "on"
    try:
        db.commit()
    except IntegrityError as error_db:
        db.rollback()
        raise HTTPException(status_code=422, detail="Ya existe un concepto con ese nombre o código.") from error_db
    return RedirectResponse("/configuracion/conceptos-excepcionales?" + ("vista=precargas&" if vista == "precargas" else "") + "guardado=concepto+actualizado", status_code=303)


@app.post("/configuracion/conceptos-excepcionales/asignaciones")
def asignar_concepto_excepcional(
    request: Request,
    usuario_id: Annotated[int, Form()],
    concepto_excepcional_id: Annotated[int, Form()],
    fecha_desde: Annotated[date, Form()],
    observacion: Annotated[str, Form()],
    fecha_hasta: Annotated[date | None, Form()] = None,
    cantidad_asignada: Annotated[Decimal | None, Form()] = None,
    vista: str | None = None,
    db: Session = Depends(get_db),
):
    asignacion_activa, administra_catalogo = acceso_conceptos_excepcionales(request, db)
    usuarios_permitidos = {
        item.id for item in usuarios_asignables_excepciones(
            asignacion_activa, administra_catalogo, db,
        )
    }
    if usuario_id not in usuarios_permitidos:
        raise HTTPException(status_code=403, detail="No podés asignar excepciones a ese usuario.")
    concepto = db.get(ConceptoExcepcional, concepto_excepcional_id)
    if concepto is None or not concepto.activo:
        raise HTTPException(status_code=422, detail="El concepto no existe o está inactivo.")
    usuario = db.scalar(
        select(Usuario)
        .options(selectinload(Usuario.tipo_contratacion))
        .where(Usuario.id == usuario_id)
    )
    if usuario is None:
        raise HTTPException(status_code=404, detail="El usuario no existe.")
    observacion = observacion.strip()
    if not observacion:
        raise HTTPException(status_code=422, detail="La observación administrativa es obligatoria.")
    if fecha_hasta is not None and fecha_hasta < fecha_desde:
        raise HTTPException(status_code=422, detail="La fecha hasta no puede ser anterior a la fecha desde.")
    cantidad_configurada = None
    if concepto.tipo == "PRECARGA_PERIODO":
        if not administra_catalogo:
            raise HTTPException(
                status_code=403,
                detail="Sólo RRHH o ADMIN pueden asignar precargas autorizadas.",
            )
        if fecha_hasta is None:
            raise HTTPException(
                status_code=422,
                detail="La precarga autorizada necesita una fecha hasta.",
            )
        if cantidad_asignada is None or cantidad_asignada <= 0:
            raise HTTPException(
                status_code=422,
                detail="Ingresá una cantidad por período mayor que cero.",
            )
        cantidad_configurada = cantidad_asignada.quantize(Decimal("0.01"))
    consulta_solapada = select(UsuarioConceptoExcepcional.id_asignacion).where(
        UsuarioConceptoExcepcional.usuario_id == usuario_id,
        UsuarioConceptoExcepcional.concepto_excepcional_id == concepto_excepcional_id,
        UsuarioConceptoExcepcional.activo.is_(True),
        (UsuarioConceptoExcepcional.fecha_hasta.is_(None) | (UsuarioConceptoExcepcional.fecha_hasta >= fecha_desde)),
    )
    if fecha_hasta is not None:
        consulta_solapada = consulta_solapada.where(
            UsuarioConceptoExcepcional.fecha_desde <= fecha_hasta,
        )
    solapada = db.scalar(consulta_solapada.limit(1))
    if solapada is not None:
        raise HTTPException(status_code=422, detail="El usuario ya tiene una asignación vigente que se superpone.")
    nueva_asignacion = UsuarioConceptoExcepcional(
        usuario_id=usuario_id,
        concepto_excepcional_id=concepto_excepcional_id,
        fecha_desde=fecha_desde,
        fecha_hasta=fecha_hasta,
        activo=True,
        cantidad_asignada=cantidad_configurada,
        observacion=observacion,
        asignado_por=(request.session.get("user") or {}).get("id"),
    )
    db.add(nueva_asignacion)
    db.flush()
    precargas_creadas = generar_precargas_autorizadas(
        nueva_asignacion,
        concepto,
        usuario,
        nueva_asignacion.asignado_por,
        db,
    )
    db.commit()
    mensaje = "asignación creada"
    if concepto.tipo == "PRECARGA_PERIODO":
        mensaje = f"asignación creada y {precargas_creadas} precargas autorizadas generadas"
    return RedirectResponse(
        "/configuracion/conceptos-excepcionales?"
        + ("vista=precargas&" if vista == "precargas" else "")
        + "guardado=" + quote(mensaje),
        status_code=303,
    )


@app.post("/configuracion/conceptos-excepcionales/asignaciones/{asignacion_id}/desactivar")
def desactivar_asignacion_excepcional(
    asignacion_id: int,
    request: Request,
    vista: str | None = None,
    db: Session = Depends(get_db),
):
    asignacion_activa, administra_catalogo = acceso_conceptos_excepcionales(request, db)
    registro = db.get(UsuarioConceptoExcepcional, asignacion_id)
    if registro is None:
        raise HTTPException(status_code=404, detail="La asignación no existe.")
    permitidos = {
        item.id for item in usuarios_asignables_excepciones(
            asignacion_activa, administra_catalogo, db,
        )
    }
    if registro.usuario_id not in permitidos:
        raise HTTPException(status_code=403, detail="No podés modificar esa asignación.")
    concepto = db.get(ConceptoExcepcional, registro.concepto_excepcional_id)
    if concepto is not None and concepto.tipo == "PRECARGA_PERIODO":
        precargas_futuras = list(db.scalars(
            select(HoraExtra).where(
                HoraExtra.usuario_id == registro.usuario_id,
                HoraExtra.concepto_excepcional_id == registro.concepto_excepcional_id,
                HoraExtra.tipo_registro == TIPO_PRECARGA_AUTORIZADA,
                HoraExtra.fecha > date.today(),
            )
        ))
        for precarga in precargas_futuras:
            db.delete(precarga)
    registro.activo = False
    if registro.fecha_hasta is None or registro.fecha_hasta > date.today():
        registro.fecha_hasta = date.today()
    db.commit()
    return RedirectResponse("/configuracion/conceptos-excepcionales?" + ("vista=precargas&" if vista == "precargas" else "") + "guardado=asignación+desactivada", status_code=303)


@app.get("/configuracion/parametros", response_class=HTMLResponse)
def parametros(
    request: Request,
    editar: int | None = None,
    guardado: str | None = None,
    db: Session = Depends(get_db),
):
    error = None
    reglas: list[ReglaHora] = []
    regla_edicion = None
    try:
        reglas = list(db.scalars(select(ReglaHora).order_by(ReglaHora.convenio, ReglaHora.tipo_dia, ReglaHora.id)))
        if editar is not None:
            regla_edicion = db.get(ReglaHora, editar)
            if regla_edicion is None:
                error = "La regla seleccionada no existe."
    except SQLAlchemyError:
        error = "No se pudieron cargar las reglas. Verificá la conexión con la base de datos."

    return templates.TemplateResponse(
        request,
        "parametros.html",
        {
            "active_page": "parametros",
            "reglas": reglas,
            "regla_edicion": regla_edicion,
            "guardado": guardado,
            "error": error,
            "roles_disponibles": ROLES_DISPONIBLES,
            "perfiles_disponibles": PERFILES_DISPONIBLES,
            "convenios_disponibles": CONVENIOS_DISPONIBLES,
        },
    )


@app.post("/configuracion/parametros/reglas")
def crear_regla(
    convenio: Annotated[str, Form()],
    tipo_dia: Annotated[str, Form()],
    tipo_hora: Annotated[str, Form()],
    hora_nocturna_desde: Annotated[str | None, Form()] = None,
    hora_nocturna_hasta: Annotated[str | None, Form()] = None,
    permite_reintegro: Annotated[str | None, Form()] = None,
    observaciones: Annotated[str | None, Form()] = None,
    modalidad_carga: Annotated[str | None, Form()] = None,
    cantidad_predeterminada: Annotated[Decimal | None, Form()] = None,
    db: Session = Depends(get_db),
):
    nueva_regla = ReglaHora(**rule_form_data(
        convenio, tipo_dia, tipo_hora, hora_nocturna_desde,
        hora_nocturna_hasta, permite_reintegro, observaciones,
        modalidad_carga, cantidad_predeterminada,
    ))
    try:
        db.add(nueva_regla)
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudo guardar la regla.") from error
    return RedirectResponse("/configuracion/parametros?guardado=creada", status_code=303)


@app.post("/configuracion/parametros/reglas/{regla_id}")
def actualizar_regla(
    regla_id: int,
    convenio: Annotated[str, Form()],
    tipo_dia: Annotated[str, Form()],
    tipo_hora: Annotated[str, Form()],
    hora_nocturna_desde: Annotated[str | None, Form()] = None,
    hora_nocturna_hasta: Annotated[str | None, Form()] = None,
    permite_reintegro: Annotated[str | None, Form()] = None,
    observaciones: Annotated[str | None, Form()] = None,
    modalidad_carga: Annotated[str | None, Form()] = None,
    cantidad_predeterminada: Annotated[Decimal | None, Form()] = None,
    db: Session = Depends(get_db),
):
    regla = db.get(ReglaHora, regla_id)
    if regla is None:
        raise HTTPException(status_code=404, detail="La regla no existe.")
    for field, value in rule_form_data(
        convenio, tipo_dia, tipo_hora, hora_nocturna_desde,
        hora_nocturna_hasta, permite_reintegro, observaciones,
        modalidad_carga, cantidad_predeterminada,
    ).items():
        setattr(regla, field, value)
    try:
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudo actualizar la regla.") from error
    return RedirectResponse("/configuracion/parametros?guardado=actualizada", status_code=303)


@app.post("/configuracion/parametros/reglas/{regla_id}/eliminar")
def eliminar_regla(
    regla_id: int,
    db: Session = Depends(get_db),
):
    regla = db.get(ReglaHora, regla_id)
    if regla is None:
        raise HTTPException(status_code=404, detail="La regla no existe.")
    try:
        db.delete(regla)
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(
            status_code=500,
            detail="No se pudo eliminar la regla.",
        ) from error
    return RedirectResponse(
        "/configuracion/parametros?guardado=eliminada",
        status_code=303,
    )


@app.get("/configuracion/usuarios", response_class=HTMLResponse)
def usuarios_y_roles(
    request: Request,
    editar_usuario: int | None = None,
    consultar_ad: bool = False,
    guardado: str | None = None,
    error: str | None = None,
    filtro_rol: str | None = None,
    filtro_perfil: str | None = None,
    filtro_tipo_contratacion: int | None = None,
    filtro_convenio: str | None = None,
    filtro_estado: str = "ACTIVO",
    db: Session = Depends(get_db),
):
    usuarios: list[Usuario] = []
    tipos_contratacion: list[TipoContratacion] = []
    convenios: list[str] = list(CONVENIOS_DISPONIBLES)
    usuario_edicion = None
    usuarios_ad = None
    error_ad = None
    filtro_rol = (filtro_rol or "").strip().upper()
    filtro_perfil = (filtro_perfil or "").strip().upper()
    filtro_convenio = (filtro_convenio or "").strip().upper()
    filtro_estado = filtro_estado.strip().upper()
    if filtro_rol not in ROLES_DISPONIBLES:
        filtro_rol = ""
    if filtro_perfil not in PERFILES_DISPONIBLES:
        filtro_perfil = ""
    if filtro_convenio not in CONVENIOS_DISPONIBLES:
        filtro_convenio = ""
    if filtro_estado not in {"ACTIVO", "INACTIVO"}:
        filtro_estado = "ACTIVO"
    try:
        tipos_contratacion = list(db.scalars(
            select(TipoContratacion).order_by(TipoContratacion.id_tipo_contratacion)
        ))
        ids_tipos_contratacion = {
            item.id_tipo_contratacion for item in tipos_contratacion
        }
        if filtro_tipo_contratacion not in ids_tipos_contratacion:
            filtro_tipo_contratacion = None
        consulta_usuarios = (
            select(Usuario)
            .options(selectinload(Usuario.roles), selectinload(Usuario.tipo_contratacion))
            .where(Usuario.status.is_(filtro_estado == "ACTIVO"))
            .order_by(Usuario.apellido, Usuario.nombre)
        )
        condiciones_asignacion = []
        if filtro_rol:
            condiciones_asignacion.append(func.upper(UsuarioRol.rol) == filtro_rol)
        if filtro_perfil:
            condiciones_asignacion.append(func.upper(UsuarioRol.perfil) == filtro_perfil)
        if condiciones_asignacion:
            consulta_usuarios = consulta_usuarios.where(
                Usuario.roles.any(and_(*condiciones_asignacion))
            )
        if filtro_tipo_contratacion is not None:
            consulta_usuarios = consulta_usuarios.where(
                Usuario.id_tipo_contratacion == filtro_tipo_contratacion
            )
        if filtro_convenio:
            consulta_usuarios = consulta_usuarios.where(
                func.upper(Usuario.convenio) == filtro_convenio
            )
        usuarios = list(db.scalars(consulta_usuarios))
        if editar_usuario is not None:
            usuario_edicion = db.get(Usuario, editar_usuario)
            if usuario_edicion is None:
                error = "El usuario seleccionado no existe."
    except SQLAlchemyError:
        error = "No se pudieron cargar usuarios y roles. Verificá la conexión con la base."

    if consultar_ad:
        try:
            usuarios_ad = list_ad_group_users()
        except ActiveDirectoryAuthError as ad_error:
            usuarios_ad = []
            error_ad = str(ad_error)

    return templates.TemplateResponse(request, "usuarios.html", {
        "active_page": "usuarios",
        "usuarios": usuarios,
        "tipos_contratacion": tipos_contratacion,
        "convenios": convenios,
        "usuario_edicion": usuario_edicion,
        "usuarios_ad": usuarios_ad,
        "error_ad": error_ad,
        "guardado": guardado,
        "error": error,
        "roles_disponibles": ROLES_DISPONIBLES,
        "perfiles_disponibles": PERFILES_DISPONIBLES,
        "filtro_rol": filtro_rol,
        "filtro_perfil": filtro_perfil,
        "filtro_tipo_contratacion": filtro_tipo_contratacion,
        "filtro_convenio": filtro_convenio,
        "filtro_estado": filtro_estado,
    })


def require_user_management_role(request: Request) -> None:
    role = str(request.session.get("user", {}).get("role") or "").upper()
    if role not in {"ADMIN", "RRHH"}:
        raise HTTPException(status_code=403, detail="Se requiere el rol ADMIN o RRHH.")


@app.get("/configuracion/usuarios/sincronizar", response_class=HTMLResponse)
def previsualizar_usuarios_ad(
    request: Request,
    guardado: str | None = None,
    error: str | None = None,
    db: Session = Depends(get_db),
):
    require_user_management_role(request)
    try:
        ad_users = list_ad_group_users()
    except ActiveDirectoryAuthError as ad_error:
        return templates.TemplateResponse(request, "usuarios_ad_sync.html", {
            "active_page": "usuarios",
            "error": str(ad_error),
            "guardado": guardado,
            "new_users": [],
            "existing_users": [],
            "inactive_users": [],
            "reactivable_users": [],
            "convenios": [],
            "tipos_contratacion": [],
        }, status_code=502)

    local_users = list(db.scalars(select(Usuario).options(selectinload(Usuario.roles))))
    local_by_username = {user.username.strip().lower(): user for user in local_users}
    ad_by_username = {user.username.strip().lower(): user for user in ad_users}
    new_users = [user for key, user in ad_by_username.items() if key not in local_by_username]
    existing_users = [
        local_by_username[key] for key in ad_by_username
        if key in local_by_username and local_by_username[key].status
    ]
    reactivable_users = [
        local_by_username[key] for key in ad_by_username
        if key in local_by_username and not local_by_username[key].status
    ]
    protected = local_usernames()
    inactive_users = [
        user for key, user in local_by_username.items()
        if user.origen.upper() in {"LDAP", "AD"}
        and user.status and key not in ad_by_username and key not in protected
    ]
    convenios = list(CONVENIOS_DISPONIBLES)
    tipos = list(db.scalars(
        select(TipoContratacion).order_by(TipoContratacion.id_tipo_contratacion)
    ))
    return templates.TemplateResponse(request, "usuarios_ad_sync.html", {
        "active_page": "usuarios",
        "error": error,
        "guardado": guardado,
        "new_users": new_users,
        "existing_users": existing_users,
        "inactive_users": inactive_users,
        "reactivable_users": reactivable_users,
        "convenios": convenios,
        "tipos_contratacion": tipos,
        "roles_disponibles": ROLES_DISPONIBLES,
        "perfiles_disponibles": PERFILES_DISPONIBLES,
    })


@app.post("/configuracion/usuarios/sincronizar/incorporar")
def incorporar_usuario_ad(
    request: Request,
    username: Annotated[str, Form()],
    nombre: Annotated[str, Form()],
    apellido: Annotated[str, Form()],
    correo: Annotated[str | None, Form()] = None,
    legajo: Annotated[str | None, Form()] = None,
    convenio: Annotated[str, Form()] = "",
    id_tipo_contratacion: Annotated[int, Form()] = 0,
    rol: Annotated[str, Form()] = "",
    perfil: Annotated[str, Form()] = "",
    db: Session = Depends(get_db),
):
    require_user_management_role(request)
    normalized = normalize_username(username)
    if find_user(db, normalized) is not None:
        return RedirectResponse("/configuracion/usuarios/sincronizar?error=El+usuario+ya+existe+y+no+fue+modificado", status_code=303)
    try:
        ad_user = next(
            (item for item in list_ad_group_users() if normalize_username(item.username) == normalized),
            None,
        )
    except ActiveDirectoryAuthError:
        return RedirectResponse("/configuracion/usuarios/sincronizar?error=No+se+pudo+verificar+el+usuario+en+AD", status_code=303)
    if ad_user is None:
        return RedirectResponse("/configuracion/usuarios/sincronizar?error=El+usuario+ya+no+pertenece+al+grupo+de+AD", status_code=303)
    role_name, profile_name = rol.strip().upper(), perfil.strip().upper()
    if role_name not in ROLES_DISPONIBLES or profile_name not in PERFILES_DISPONIBLES:
        raise HTTPException(status_code=422, detail="El rol o perfil no es válido.")
    if db.get(TipoContratacion, id_tipo_contratacion) is None:
        raise HTTPException(status_code=422, detail="La contratación no existe.")
    if convenio.strip().upper() not in CONVENIOS_DISPONIBLES:
        raise HTTPException(status_code=422, detail="El convenio no existe.")
    user = Usuario(
        legajo=clean_optional(legajo),
        nombre=nombre.strip() or ad_user.first_name,
        apellido=apellido.strip() or ad_user.last_name,
        correo=clean_optional(correo) or ad_user.email,
        convenio=convenio.strip().upper(),
        id_tipo_contratacion=id_tipo_contratacion,
        username=normalized,
        hashed_password=hash_password(secrets.token_urlsafe(32)),
        origen="AD",
        status=True,
    )
    user.roles.append(UsuarioRol(rol=role_name, perfil=profile_name))
    db.add(user)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse("/configuracion/usuarios/sincronizar?error=El+legajo,+correo+o+usuario+ya+está+registrado", status_code=303)
    return RedirectResponse("/configuracion/usuarios/sincronizar?guardado=usuario+incorporado", status_code=303)


@app.post("/configuracion/usuarios/sincronizar/{usuario_id}/reactivar")
def reactivar_usuario_ad(usuario_id: int, request: Request, db: Session = Depends(get_db)):
    require_user_management_role(request)
    user = db.get(Usuario, usuario_id)
    if user is None or user.origen.upper() not in {"LDAP", "AD"}:
        raise HTTPException(status_code=404, detail="El usuario LDAP no existe.")
    try:
        present = any(normalize_username(item.username) == normalize_username(user.username) for item in list_ad_group_users())
    except ActiveDirectoryAuthError as error:
        raise HTTPException(status_code=502, detail=str(error)) from error
    if not present:
        raise HTTPException(status_code=409, detail="El usuario no pertenece actualmente al grupo de AD.")
    user.status = True
    db.commit()
    return RedirectResponse("/configuracion/usuarios/sincronizar?guardado=usuario+reactivado", status_code=303)


@app.post("/configuracion/usuarios/sincronizar/desactivar-ausentes")
def desactivar_usuarios_ad_ausentes(request: Request, db: Session = Depends(get_db)):
    require_user_management_role(request)
    try:
        present = {normalize_username(item.username) for item in list_ad_group_users()}
    except ActiveDirectoryAuthError as error:
        raise HTTPException(status_code=502, detail=str(error)) from error
    protected = local_usernames()
    users = list(db.scalars(select(Usuario).where(Usuario.status.is_(True))))
    changed = 0
    for user in users:
        username = normalize_username(user.username)
        if user.origen.upper() in {"LDAP", "AD"} and username not in present and username not in protected:
            user.status = False
            changed += 1
    db.commit()
    return RedirectResponse(f"/configuracion/usuarios/sincronizar?guardado={changed}+usuarios+desactivados", status_code=303)


@app.post("/configuracion/usuarios/{usuario_id}/roles")
def crear_asignacion_usuario(
    usuario_id: int,
    rol: Annotated[str, Form()],
    perfil: Annotated[str, Form()],
    db: Session = Depends(get_db),
):
    usuario = db.get(Usuario, usuario_id)
    if usuario is None:
        raise HTTPException(status_code=404, detail="El usuario no existe.")
    nombre_rol = rol.strip().upper()
    nombre_perfil = perfil.strip().upper()
    if nombre_rol not in ROLES_DISPONIBLES or nombre_perfil not in PERFILES_DISPONIBLES:
        raise HTTPException(status_code=422, detail="La asignación de rol y perfil no es válida.")
    db.add(UsuarioRol(usuario_id=usuario_id, rol=nombre_rol, perfil=nombre_perfil))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(
            f"/configuracion/usuarios?editar_usuario={usuario_id}&error=La+asignación+ya+existe#asignaciones",
            status_code=303,
        )
    return RedirectResponse(
        f"/configuracion/usuarios?editar_usuario={usuario_id}&guardado=asignación+creada#asignaciones",
        status_code=303,
    )


@app.post("/configuracion/usuarios/{usuario_id}/roles/{rol_id}/eliminar")
def eliminar_asignacion_usuario(
    usuario_id: int,
    rol_id: int,
    db: Session = Depends(get_db),
):
    registro = db.get(UsuarioRol, rol_id)
    if registro is None or registro.usuario_id != usuario_id:
        raise HTTPException(status_code=404, detail="La asignación no existe.")
    db.delete(registro)
    try:
        db.commit()
    except SQLAlchemyError:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudo eliminar la asignación.")
    return RedirectResponse(
        f"/configuracion/usuarios?editar_usuario={usuario_id}&guardado=asignación+eliminada#asignaciones",
        status_code=303,
    )


def update_user_fields(
    usuario: Usuario,
    legajo: str | None,
    nombre: str,
    apellido: str,
    correo: str | None,
    convenio: str | None,
    id_tipo_contratacion: int,
    observaciones: str | None,
    username: str,
    origen: str,
    status: str | None,
) -> None:
    required = [nombre.strip(), apellido.strip(), username.strip(), origen.strip()]
    if not all(required):
        raise HTTPException(status_code=422, detail="Completá los campos obligatorios.")
    usuario.legajo = clean_optional(legajo)
    usuario.nombre = nombre.strip()
    usuario.apellido = apellido.strip()
    usuario.correo = clean_optional(correo)
    usuario.convenio = clean_optional(convenio)
    usuario.id_tipo_contratacion = id_tipo_contratacion
    usuario.observaciones = clean_optional(observaciones)
    usuario.username = username.strip()
    usuario.origen = origen.strip()
    usuario.status = status == "on"


@app.post("/configuracion/usuarios")
def crear_usuario(
    nombre: Annotated[str, Form()],
    apellido: Annotated[str, Form()],
    username: Annotated[str, Form()],
    password: Annotated[str, Form()],
    origen: Annotated[str, Form()],
    rol: Annotated[str, Form()],
    perfil: Annotated[str, Form()],
    legajo: Annotated[str | None, Form()] = None,
    correo: Annotated[str | None, Form()] = None,
    convenio: Annotated[str | None, Form()] = None,
    id_tipo_contratacion: Annotated[int, Form()] = 0,
    observaciones: Annotated[str | None, Form()] = None,
    status: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    nombre_rol = rol.strip().upper()
    nombre_perfil = perfil.strip().upper()
    if nombre_rol not in ROLES_DISPONIBLES or nombre_perfil not in PERFILES_DISPONIBLES:
        raise HTTPException(status_code=422, detail="La asignación de rol y perfil no es válida.")
    if len(password) < 8:
        return RedirectResponse("/configuracion/usuarios?error=La+contraseña+debe+tener+al+menos+8+caracteres", status_code=303)
    if db.get(TipoContratacion, id_tipo_contratacion) is None:
        raise HTTPException(status_code=422, detail="La contratación seleccionada no existe.")
    if not convenio or convenio.strip().upper() not in CONVENIOS_DISPONIBLES:
        raise HTTPException(status_code=422, detail="El convenio seleccionado no es válido.")
    convenio = convenio.strip().upper()
    usuario = Usuario(hashed_password=hash_password(password))
    update_user_fields(
        usuario, legajo, nombre, apellido, correo, convenio,
        id_tipo_contratacion, observaciones, username, origen, status,
    )
    db.add(usuario)
    try:
        db.flush()
        db.add(UsuarioRol(
            usuario_id=usuario.id,
            rol=nombre_rol,
            perfil=nombre_perfil,
        ))
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse("/configuracion/usuarios?error=El+legajo,+correo+o+usuario+ya+está+registrado", status_code=303)
    return RedirectResponse("/configuracion/usuarios?guardado=usuario+creado", status_code=303)


@app.post("/configuracion/usuarios/{usuario_id}")
def actualizar_usuario(
    usuario_id: int,
    nombre: Annotated[str, Form()],
    apellido: Annotated[str, Form()],
    username: Annotated[str, Form()],
    origen: Annotated[str, Form()],
    legajo: Annotated[str | None, Form()] = None,
    correo: Annotated[str | None, Form()] = None,
    convenio: Annotated[str | None, Form()] = None,
    id_tipo_contratacion: Annotated[int, Form()] = 0,
    observaciones: Annotated[str | None, Form()] = None,
    password: Annotated[str | None, Form()] = None,
    status: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    usuario = db.get(Usuario, usuario_id)
    if usuario is None:
        raise HTTPException(status_code=404, detail="El usuario no existe.")
    if db.get(TipoContratacion, id_tipo_contratacion) is None:
        raise HTTPException(status_code=422, detail="La contratación seleccionada no existe.")
    if not convenio or convenio.strip().upper() not in CONVENIOS_DISPONIBLES:
        raise HTTPException(status_code=422, detail="El convenio seleccionado no es válido.")
    convenio = convenio.strip().upper()
    if password:
        if len(password) < 8:
            return RedirectResponse(f"/configuracion/usuarios?editar_usuario={usuario_id}&error=La+contraseña+debe+tener+al+menos+8+caracteres", status_code=303)
        usuario.hashed_password = hash_password(password)
    update_user_fields(
        usuario, legajo, nombre, apellido, correo, convenio,
        id_tipo_contratacion, observaciones, username, origen, status,
    )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(f"/configuracion/usuarios?editar_usuario={usuario_id}&error=El+legajo,+correo+o+usuario+ya+está+registrado", status_code=303)
    return RedirectResponse("/configuracion/usuarios?guardado=usuario+actualizado", status_code=303)


@app.get("/configuracion/feriados", response_class=HTMLResponse)
def administrar_feriados(
    request: Request,
    editar: int | None = None,
    guardado: str | None = None,
    error: str | None = None,
    anio_actualizado: int | None = None,
    cantidad: int | None = None,
    duplicados: int | None = None,
    db: Session = Depends(get_db),
):
    feriados: list[Feriado] = []
    feriado_edicion = None
    try:
        feriados = list(db.scalars(select(Feriado).order_by(Feriado.fecha)))
        if editar is not None:
            feriado_edicion = db.get(Feriado, editar)
            if feriado_edicion is None:
                error = "El feriado seleccionado no existe"
    except SQLAlchemyError:
        error = "No se pudieron cargar los feriados. Verificá la conexión con la base"
    return templates.TemplateResponse(request, "feriados.html", {
        "active_page": "feriados",
        "feriados": feriados,
        "feriado_edicion": feriado_edicion,
        "guardado": guardado,
        "error": error,
        "anio_actualizado": anio_actualizado,
        "cantidad": cantidad or 0,
        "duplicados": duplicados or 0,
        "anio_actualizacion": min(
            (feriado.fecha.year for feriado in feriados),
            default=None,
        ),
        "es_jefe": str((request.session.get("user") or {}).get("profile") or "").upper() == "JEFE",
    })


def pending_owned_hour(request: Request, hora_id: int, db: Session) -> HoraExtra:
    assignment = obtener_asignacion_activa(request, db)
    if assignment.perfil.upper() != "USUARIO":
        raise HTTPException(
            status_code=403,
            detail="El perfil jefe no puede modificar ni eliminar horas extras.",
        )
    user_id = (request.session.get("user") or {}).get("id")
    current_user = db.get(Usuario, user_id)
    if current_user is not None and not usuario_tiene_funciones_carga(current_user, db):
        raise HTTPException(status_code=403, detail="No tenés funciones de carga habilitadas.")
    hour = db.scalar(
        select(HoraExtra)
        .options(selectinload(HoraExtra.usuario))
        .where(HoraExtra.id == hora_id, HoraExtra.usuario_id == user_id)
    )
    if hour is None:
        raise HTTPException(status_code=404, detail="La carga no existe.")
    if hour.estado != "PENDIENTE":
        raise HTTPException(status_code=409, detail="Sólo se pueden modificar cargas pendientes.")
    return hour


@app.get("/horas/{hora_id}/editar", response_class=HTMLResponse)
def editar_hora_propia(hora_id: int, request: Request, db: Session = Depends(get_db)):
    hour = pending_owned_hour(request, hora_id, db)
    fecha_minima, fecha_maxima, _ = limites_fecha_carga_usuario(hour.usuario)
    return templates.TemplateResponse(request, "hora_editar.html", {
        "active_page": "horas",
        "hora": hour,
        "fecha_minima": fecha_minima.isoformat() if fecha_minima else "",
        "fecha_maxima": fecha_maxima.isoformat() if fecha_maxima else "",
        "edicion_rrhh": False,
    })


def hora_editable_por_rrhh(
    request: Request,
    hora_id: int,
    db: Session,
) -> tuple[HoraExtra, UsuarioRol]:
    asignacion = obtener_asignacion_activa(request, db)
    if asignacion.rol.upper() not in {"ADMIN", "RRHH"}:
        raise HTTPException(status_code=403, detail="Sólo RRHH o ADMIN pueden corregir estas cargas.")
    hora = db.scalar(
        select(HoraExtra)
        .options(selectinload(HoraExtra.usuario))
        .where(HoraExtra.id == hora_id)
    )
    if hora is None:
        raise HTTPException(status_code=404, detail="La carga no existe.")
    if es_concepto_automatico(hora):
        raise HTTPException(
            status_code=409,
            detail="El concepto automático se corrige editando las horas de la jornada.",
        )
    return hora, asignacion


@app.get("/rrhh/horas/{hora_id}/editar", response_class=HTMLResponse)
def editar_hora_rrhh(hora_id: int, request: Request, db: Session = Depends(get_db)):
    hora, _ = hora_editable_por_rrhh(request, hora_id, db)
    fecha_minima, fecha_maxima, _ = limites_fecha_carga_usuario(hora.usuario)
    return templates.TemplateResponse(request, "hora_editar.html", {
        "active_page": "rrhh_horas",
        "hora": hora,
        "fecha_minima": fecha_minima.isoformat() if fecha_minima else "",
        "fecha_maxima": fecha_maxima.isoformat() if fecha_maxima else "",
        "edicion_rrhh": True,
    })


@app.post("/rrhh/horas/{hora_id}/editar")
def guardar_hora_rrhh(
    hora_id: int,
    request: Request,
    fecha: Annotated[date, Form()],
    hora_inicio: Annotated[time | None, Form()] = None,
    hora_fin: Annotated[time | None, Form()] = None,
    cantidad: Annotated[Decimal | None, Form()] = None,
    fecha_descanso_compensatorio: Annotated[date | None, Form()] = None,
    observaciones: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    hora, asignacion = hora_editable_por_rrhh(request, hora_id, db)
    estado_anterior = hora.estado
    fecha_anterior = hora.fecha
    cerraba_medianoche = hora.hora_fin == time(0, 0)
    jornada_anterior = (
        fecha_jornada_de_hora_estado(hora, db, estado_anterior)
        if hora.tipo_registro == "HORAS" else None
    )
    try:
        observacion_limpia = clean_optional(observaciones)
        if observacion_limpia is None:
            raise CalculoHorasError("La observación es obligatoria.")
        validar_fecha_carga_usuario(hora.usuario, fecha)
        if hora.tipo_registro == "REINTEGRO":
            resultado = calcular_resultado_reintegro(
                hora.usuario_id, fecha, db, cantidad or hora.cantidad or Decimal("1.00"),
            )
        elif hora.tipo_registro == "DIA_TRABAJADO":
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("Ingresá la hora de inicio y finalización de la jornada.")
            resultado = calcular_resultado_dia_trabajado(
                hora.usuario_id, fecha, db,
                marcar_como_franco=clasificar_tipo_dia(fecha, db) != "FERIADO",
                hora_inicio=hora_inicio, hora_fin=hora_fin,
            )
        elif hora.tipo_registro == "EXCEPCIONAL":
            if hora.concepto_excepcional_id is None:
                raise CalculoHorasError("La carga no tiene un concepto excepcional asociado.")
            resultado = calcular_resultado_concepto_excepcional(
                hora.usuario_id, fecha, hora.concepto_excepcional_id,
                cantidad or hora.cantidad or Decimal("0.00"), db,
            )
        elif hora.tipo_registro == "OTRAS":
            if str(hora.tipo_hora or "").upper() in TIPOS_NOCTURNAS_EN_JORNADA:
                if hora_inicio is None or hora_fin is None:
                    raise CalculoHorasError("Ingresá el horario completo de la jornada.")
                validar_nocturnidad_sin_superposicion(
                    hora.usuario_id, fecha, hora_inicio, hora_fin, db,
                    excluir_hora_id=hora.id,
                )
                resultado = calcular_resultado_nocturnas_en_jornada(
                    hora.usuario_id, fecha, hora_inicio, hora_fin, db,
                )
            else:
                resultado = calcular_resultado_otra_carga(
                    hora.usuario_id, fecha, hora.tipo_hora,
                    cantidad or hora.cantidad or Decimal("0.00"), db,
                )
        else:
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("Ingresá la hora de inicio y finalización.")
            resultado = calcular_resultado_horas_extra(
                hora.usuario_id, fecha, hora_inicio, hora_fin, db,
                marcar_como_franco=hora.tipo_dia == "FRANCO",
            )

        hora.fecha = resultado.fecha
        hora.hora_inicio = resultado.hora_inicio
        hora.hora_fin = resultado.hora_fin
        hora.horas_totales = resultado.horas_totales
        hora.cantidad = resultado.cantidad
        hora.tipo_dia = resultado.tipo_dia
        hora.tipo_hora = resultado.tipo_hora
        hora.horas_nocturnas = resultado.horas_nocturnas
        hora.observaciones = observacion_limpia
        if hora.tipo_registro == "DIA_TRABAJADO" and hora.fecha_descanso_compensatorio is not None:
            if fecha_descanso_compensatorio is None:
                raise CalculoHorasError("Indicá qué día se tomará el franco o feriado trabajado.")
            hora.fecha_descanso_compensatorio = fecha_descanso_compensatorio
        momento = datetime.now()
        hora.estado = "APROBADA"
        hora.aprobado_por = asignacion.usuario_id
        hora.fecha_resolucion = momento
        hora.observacion_rechazo = None
        db.flush()

        if hora.tipo_registro == "HORAS":
            if estado_anterior == "PENDIENTE" and jornada_anterior is not None:
                recalcular_conceptos_jornada_estado(
                    hora.usuario_id, jornada_anterior, db, estado_anterior,
                )
            jornada_nueva = fecha_jornada_de_hora_estado(hora, db, "APROBADA")
            jornadas_aprobadas = {jornada_nueva}
            if estado_anterior == "APROBADA":
                jornadas_aprobadas.add(jornada_anterior)
            if cerraba_medianoche:
                jornadas_aprobadas.add(fecha_anterior + timedelta(days=1))
            if hora.hora_fin == time(0, 0):
                jornadas_aprobadas.add(hora.fecha + timedelta(days=1))
            for fecha_jornada in jornadas_aprobadas:
                if fecha_jornada is not None:
                    recalcular_conceptos_jornada_estado(
                        hora.usuario_id, fecha_jornada, db, "APROBADA",
                        aprobado_por=asignacion.usuario_id,
                        fecha_resolucion=momento,
                    )
        for fecha_switcher in {fecha_anterior, hora.fecha}:
            recalcular_switcher_nocturnidad_estado(
                hora.usuario_id, fecha_switcher, db, "APROBADA",
                aprobado_por=asignacion.usuario_id,
                fecha_resolucion=momento,
            )
        db.commit()
    except (CalculoHorasError, SQLAlchemyError) as error:
        db.rollback()
        mensaje = str(error) if isinstance(error, CalculoHorasError) else "No se pudo corregir la carga."
        raise HTTPException(status_code=422, detail=mensaje) from error
    return RedirectResponse("/rrhh/horas-extras", status_code=303)


@app.post("/horas/{hora_id}/editar")
def guardar_hora_propia(
    hora_id: int,
    request: Request,
    fecha: Annotated[date, Form()],
    hora_inicio: Annotated[time | None, Form()] = None,
    hora_fin: Annotated[time | None, Form()] = None,
    marcar_como_franco: Annotated[str | None, Form()] = None,
    solicita_reintegro: Annotated[str | None, Form()] = None,
    fecha_descanso_compensatorio: Annotated[date | None, Form()] = None,
    observaciones: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    hour = pending_owned_hour(request, hora_id, db)
    if es_concepto_automatico(hour):
        raise HTTPException(status_code=409, detail="Los conceptos automáticos se administran desde la jornada.")
    jornada_anterior = (
        fecha_jornada_de_hora_pendiente(hour, db)
        if hour.tipo_registro == "HORAS" else None
    )
    fecha_anterior = hour.fecha
    cerraba_medianoche = hour.hora_fin == time(0, 0)
    try:
        observacion_limpia = clean_optional(observaciones)
        if observacion_limpia is None:
            raise CalculoHorasError("La observación es obligatoria.")
        validar_fecha_carga_usuario(hour.usuario, fecha)
        if hour.tipo_registro == "REINTEGRO":
            result = calcular_resultado_reintegro(
                hour.usuario_id, fecha, db, hour.cantidad or Decimal("1.00"),
            )
            hour.hora_inicio = None
            hour.hora_fin = None
            hour.horas_totales = None
            hour.horas_nocturnas = None
            hour.cantidad = result.cantidad
            hour.solicita_reintegro = True
        elif (
            hour.tipo_registro == "OTRAS"
            and str(hour.tipo_hora or "").upper() in TIPOS_NOCTURNAS_EN_JORNADA
        ):
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("Ingresá el horario completo de la jornada.")
            validar_nocturnidad_sin_superposicion(
                hour.usuario_id, fecha, hora_inicio, hora_fin, db,
                excluir_hora_id=hour.id,
            )
            result = calcular_resultado_nocturnas_en_jornada(
                hour.usuario_id, fecha, hora_inicio, hora_fin, db,
            )
            hour.hora_inicio = result.hora_inicio
            hour.hora_fin = result.hora_fin
            hour.horas_totales = None
            hour.horas_nocturnas = None
            hour.cantidad = result.cantidad
            hour.solicita_reintegro = False
        elif hour.tipo_registro == "DIA_TRABAJADO":
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("Ingresá la hora de inicio y finalización de la jornada.")
            result = calcular_resultado_dia_trabajado(
                hour.usuario_id, fecha, db,
                marcar_como_franco=clasificar_tipo_dia(fecha, db) != "FERIADO",
                hora_inicio=hora_inicio,
                hora_fin=hora_fin,
            )
            hour.hora_inicio = result.hora_inicio
            hour.hora_fin = result.hora_fin
            hour.horas_totales = result.horas_totales
            hour.horas_nocturnas = result.horas_nocturnas
            hour.solicita_reintegro = False
        else:
            if hora_inicio is None or hora_fin is None:
                raise CalculoHorasError("Ingresá la hora de inicio y finalización.")
            result = calcular_resultado_horas_extra(
                hour.usuario_id, fecha, hora_inicio, hora_fin, db,
                # Las horas de un franco sólo nacen desde el flujo de día trabajado.
                # Al editarlas conservamos esa clasificación sin exponer el selector.
                marcar_como_franco=hour.tipo_dia == "FRANCO",
            )
            hour.hora_inicio = result.hora_inicio
            hour.hora_fin = result.hora_fin
            hour.horas_totales = result.horas_totales
            hour.horas_nocturnas = result.horas_nocturnas
        hour.fecha = result.fecha
        hour.tipo_dia = result.tipo_dia
        hour.tipo_hora = result.tipo_hora
        hour.observaciones = observacion_limpia
        if hour.tipo_registro == "DIA_TRABAJADO" and hour.fecha_descanso_compensatorio is not None:
            if fecha_descanso_compensatorio is None:
                raise CalculoHorasError("Indicá qué día vas a tomar el franco o feriado trabajado.")
            hour.fecha_descanso_compensatorio = fecha_descanso_compensatorio
        db.flush()
        if hour.tipo_registro == "HORAS":
            jornada_nueva = fecha_jornada_de_hora_pendiente(hour, db)
            jornadas_afectadas = {jornada_anterior, jornada_nueva}
            if cerraba_medianoche:
                jornadas_afectadas.add(fecha_anterior + timedelta(days=1))
            if hour.hora_fin == time(0, 0):
                jornadas_afectadas.add(hour.fecha + timedelta(days=1))
            for fecha_jornada in jornadas_afectadas:
                if fecha_jornada is not None:
                    recalcular_conceptos_jornada_pendiente(
                        hour.usuario_id, fecha_jornada, db,
                    )
        for fecha_switcher in {fecha_anterior, hour.fecha}:
            recalcular_switcher_nocturnidad_estado(
                hour.usuario_id, fecha_switcher, db, "PENDIENTE",
            )
        db.commit()
    except (CalculoHorasError, SQLAlchemyError) as error:
        db.rollback()
        message = str(error) if isinstance(error, CalculoHorasError) else "No se pudo actualizar la carga."
        raise HTTPException(status_code=422, detail=message) from error
    return RedirectResponse("/horas", status_code=303)


@app.post("/horas/{hora_id}/eliminar")
def eliminar_hora_propia(hora_id: int, request: Request, db: Session = Depends(get_db)):
    hour = pending_owned_hour(request, hora_id, db)
    if es_concepto_automatico(hour):
        raise HTTPException(status_code=409, detail="Los conceptos automáticos se administran desde la jornada.")
    usuario_id = hour.usuario_id
    fecha_jornada = (
        fecha_jornada_de_hora_pendiente(hour, db)
        if hour.tipo_registro == "HORAS" else None
    )
    fecha_siguiente_afectada = (
        hour.fecha + timedelta(days=1)
        if hour.tipo_registro == "HORAS" and hour.hora_fin == time(0, 0)
        else None
    )
    domingos_asociados = []
    if (
        hour.tipo_registro == "DIA_TRABAJADO"
        and hour.hora_inicio is not None
        and hour.hora_fin is not None
        and str(hour.usuario.convenio or "").strip().upper() == "SAT"
    ):
        fechas_jornada = {
            fecha_tramo
            for fecha_tramo, _, _ in dividir_carga_en_fechas(
                hour.fecha, hour.hora_inicio, hour.hora_fin,
            )
        }
        domingos_asociados = list(db.scalars(
            select(HoraExtra).where(
                HoraExtra.usuario_id == usuario_id,
                HoraExtra.fecha.in_(fechas_jornada),
                HoraExtra.estado == "PENDIENTE",
                HoraExtra.tipo_registro == "OTRAS",
                func.upper(HoraExtra.tipo_hora) == "DOMINGO",
                HoraExtra.hora_inicio.is_(None),
            )
        ))
    try:
        for domingo in domingos_asociados:
            db.delete(domingo)
        db.delete(hour)
        db.flush()
        if fecha_jornada is not None:
            recalcular_conceptos_jornada_pendiente(usuario_id, fecha_jornada, db)
        if fecha_siguiente_afectada is not None:
            recalcular_conceptos_jornada_pendiente(
                usuario_id, fecha_siguiente_afectada, db,
            )
        recalcular_switcher_nocturnidad_estado(
            usuario_id, hour.fecha, db, "PENDIENTE",
        )
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudo eliminar la carga.") from error
    return RedirectResponse("/horas", status_code=303)


@app.post("/configuracion/feriados")
def crear_feriado(
    fecha: Annotated[date, Form()],
    observacion: Annotated[str, Form()],
    devuelve: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    observacion = observacion.strip()
    if not observacion:
        raise HTTPException(status_code=422, detail="La observación es obligatoria.")
    db.add(Feriado(
        fecha=fecha,
        observacion=observacion,
        devuelve=devuelve == "on",
    ))
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse("/configuracion/feriados?error=Ya+existe+un+feriado+en+esa+fecha", status_code=303)
    return RedirectResponse("/configuracion/feriados?guardado=creado", status_code=303)


@app.post("/configuracion/feriados/actualizar-anio")
def actualizar_anio_feriados(
    anio_origen: Annotated[int, Form()],
    db: Session = Depends(get_db),
):
    if anio_origen < 1900 or anio_origen > 9998:
        raise HTTPException(status_code=422, detail="El año indicado no es válido.")
    feriados_origen = list(db.scalars(
        select(Feriado)
        .where(
            Feriado.fecha >= date(anio_origen, 1, 1),
            Feriado.fecha <= date(anio_origen, 12, 31),
        )
        .order_by(Feriado.fecha)
    ))
    if not feriados_origen:
        return RedirectResponse(
            "/configuracion/feriados?error=No+hay+feriados+para+el+año+seleccionado",
            status_code=303,
        )

    anio_destino = anio_origen + 1
    fechas_destino = {
        feriado.fecha: feriado
        for feriado in db.scalars(
            select(Feriado).where(
                Feriado.fecha >= date(anio_destino, 1, 1),
                Feriado.fecha <= date(anio_destino, 12, 31),
            )
        )
    }
    actualizados = 0
    duplicados = 0
    try:
        for feriado in feriados_origen:
            try:
                nueva_fecha = feriado.fecha.replace(year=anio_destino)
            except ValueError as error:
                raise HTTPException(
                    status_code=422,
                    detail=f"La fecha {feriado.fecha.strftime('%d/%m/%Y')} no existe en {anio_destino}.",
                ) from error
            if nueva_fecha in fechas_destino:
                db.delete(feriado)
                duplicados += 1
            else:
                feriado.fecha = nueva_fecha
                fechas_destino[nueva_fecha] = feriado
                actualizados += 1
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudo actualizar el año de los feriados.") from error
    return RedirectResponse(
        f"/configuracion/feriados?anio_actualizado={anio_destino}&cantidad={actualizados}&duplicados={duplicados}",
        status_code=303,
    )


@app.post("/configuracion/feriados/{feriado_id}")
def actualizar_feriado(
    feriado_id: int,
    fecha: Annotated[date, Form()],
    observacion: Annotated[str, Form()],
    devuelve: Annotated[str | None, Form()] = None,
    db: Session = Depends(get_db),
):
    feriado = db.get(Feriado, feriado_id)
    if feriado is None:
        raise HTTPException(status_code=404, detail="El feriado no existe.")
    feriado.fecha = fecha
    feriado.observacion = observacion.strip()
    feriado.devuelve = devuelve == "on"
    if not feriado.observacion:
        raise HTTPException(status_code=422, detail="La observación es obligatoria.")
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        return RedirectResponse(
            f"/configuracion/feriados?editar={feriado_id}&error=Ya+existe+un+feriado+en+esa+fecha",
            status_code=303,
        )
    return RedirectResponse("/configuracion/feriados?guardado=actualizado", status_code=303)


@app.get("/autorizaciones", response_class=HTMLResponse)
def autorizaciones(
    request: Request,
    guardado: int | None = None,
    rechazado: int | None = None,
    usuario_id: int | None = None,
    db: Session = Depends(get_db),
):
    autorizador, asignacion_activa, es_admin = obtener_autorizador(request, db)
    grupos = []
    condiciones_alcance = []
    if not es_admin:
        condiciones_alcance.extend([
            Usuario.roles.any(func.upper(UsuarioRol.rol) == asignacion_activa.rol.upper()),
        ])
    pendientes_accesibles = list(db.scalars(
        select(HoraExtra)
        .join(HoraExtra.usuario)
        .options(selectinload(HoraExtra.usuario).selectinload(Usuario.roles))
        .where(HoraExtra.estado == "PENDIENTE", *condiciones_alcance)
        .order_by(Usuario.apellido, Usuario.nombre, HoraExtra.fecha, HoraExtra.hora_inicio)
    ))
    usuarios_filtro = []
    usuarios_vistos = set()
    for hora in pendientes_accesibles:
        if hora.usuario_id not in usuarios_vistos:
            usuarios_vistos.add(hora.usuario_id)
            usuarios_filtro.append(hora.usuario)
    if usuario_id not in usuarios_vistos:
        usuario_id = None
    pendientes = [
        hora for hora in pendientes_accesibles
        if usuario_id is None or hora.usuario_id == usuario_id
    ]
    agrupados: dict[int, dict] = {}
    for hora in pendientes:
        grupo = agrupados.setdefault(hora.usuario_id, {
            "usuario": hora.usuario,
            "horas": [],
            "total": Decimal("0.00"),
            "resumen": {},
        })
        grupo["horas"].append(hora)
        grupo["total"] += hora.horas_totales or Decimal("0.00")
        if hora.tipo_registro == "EXCEPCIONAL":
            concepto = hora.concepto_excepcional_nombre or "Concepto excepcional"
            cantidad = hora.cantidad or Decimal("0.00")
        elif hora.tipo_registro == "DIA_TRABAJADO":
            concepto, cantidad = "Día trabajado", Decimal("1.00")
        elif hora.tipo_registro == "REINTEGRO":
            concepto, cantidad = "Reintegro", hora.cantidad or Decimal("1.00")
        elif hora.tipo_registro == "OTRAS":
            concepto = "Domingo trabajado" if str(hora.tipo_hora or "").upper() == "DOMINGO" else (hora.tipo_hora or "Otras cargas")
            cantidad = hora.cantidad or Decimal("0.00")
        else:
            tipo_hora = str(hora.tipo_hora or "").strip()
            concepto = f"Horas al {tipo_hora}{'' if '%' in tipo_hora else '%'}" if tipo_hora else "Horas extras"
            cantidad = hora.horas_totales or Decimal("0.00")
        grupo["resumen"][concepto] = grupo["resumen"].get(concepto, Decimal("0.00")) + cantidad
        if hora.horas_nocturnas and hora.horas_nocturnas > 0:
            grupo["resumen"]["Horas nocturnas"] = grupo["resumen"].get(
                "Horas nocturnas", Decimal("0.00")
            ) + hora.horas_nocturnas
    grupos = list(agrupados.values())
    for grupo in grupos:
        grupo["resumen_conceptos"] = sorted(grupo.pop("resumen").items())

    # El control autorizado se muestra por el período vigente aplicable a cada
    # persona: 16–15 para Nómina y mes calendario para Monotributo/Consultora.
    hoy = date.today()
    mes_actual = periodos_rapidos_rrhh(hoy)["monotributo"]
    periodos_por_usuario: dict[int, tuple[date, date]] = {}
    for grupo in grupos:
        desde, hasta, _ = limites_fecha_carga_usuario(grupo["usuario"], hoy)
        periodos_por_usuario[grupo["usuario"].id] = (
            desde or mes_actual[0],
            hasta or mes_actual[1],
        )
        grupo["autorizadas_periodo"] = 0
        grupo["detalle_autorizado"] = []
        grupo["periodo_autorizado"] = periodos_por_usuario[grupo["usuario"].id]
    if periodos_por_usuario:
        desde_general = min(periodo[0] for periodo in periodos_por_usuario.values())
        hasta_general = max(periodo[1] for periodo in periodos_por_usuario.values())
        autorizadas_periodo = list(db.scalars(
            select(HoraExtra)
            .where(
                HoraExtra.estado == "APROBADA",
                HoraExtra.usuario_id.in_(list(periodos_por_usuario)),
                HoraExtra.fecha >= desde_general,
                HoraExtra.fecha <= hasta_general,
            )
        ))
        grupos_por_id = {grupo["usuario"].id: grupo for grupo in grupos}
        for hora in autorizadas_periodo:
            desde, hasta = periodos_por_usuario[hora.usuario_id]
            if not desde <= hora.fecha <= hasta:
                continue
            grupo = grupos_por_id[hora.usuario_id]
            grupo["autorizadas_periodo"] += 1
            grupo["detalle_autorizado"].append(hora)
    for grupo in grupos:
        grupo["detalle_autorizado"].sort(
            key=lambda hora: (hora.fecha, hora.id), reverse=True,
        )

    condiciones_autorizadas = [HoraExtra.estado == "APROBADA", *condiciones_alcance]
    if usuario_id is not None:
        condiciones_autorizadas.append(HoraExtra.usuario_id == usuario_id)
    autorizadas_total = db.scalar(
        select(func.count(HoraExtra.id))
        .join(HoraExtra.usuario)
        .where(*condiciones_autorizadas)
    ) or 0

    return templates.TemplateResponse(request, "autorizaciones.html", {
        "active_page": "autorizaciones",
        "jefe": autorizador,
        "asignacion_activa": asignacion_activa,
        "es_admin": es_admin,
        "grupos": grupos,
        "guardado": guardado,
        "rechazado": rechazado,
        "usuarios_filtro": usuarios_filtro,
        "usuario_id": usuario_id,
        "autorizadas_total": autorizadas_total,
        "cantidad_pendientes": sum(len(grupo["horas"]) for grupo in grupos),
        "total_pendiente": sum(
            (grupo["total"] for grupo in grupos),
            start=Decimal("0.00"),
        ),
    })


@app.post("/autorizaciones/aprobar")
def aprobar_horas(
    request: Request,
    hora_ids: Annotated[list[int] | None, Form()] = None,
    filtro_usuario_id: Annotated[int | None, Form()] = None,
    db: Session = Depends(get_db),
):
    ids = list(dict.fromkeys(hora_ids or []))
    if not ids:
        raise HTTPException(status_code=422, detail="Seleccioná al menos una carga para autorizar.")
    autorizador, asignacion_activa, es_admin = obtener_autorizador(request, db)
    try:
        ids = expandir_ids_con_jornadas_pendientes(ids, db)
        horas = list(db.scalars(
            select(HoraExtra)
            .options(selectinload(HoraExtra.usuario).selectinload(Usuario.roles))
            .where(HoraExtra.id.in_(ids))
            .with_for_update()
        ))
        if len(horas) != len(ids):
            raise HTTPException(status_code=404, detail="Una de las cargas seleccionadas no existe.")
        fecha_resolucion = datetime.now()
        for hora in horas:
            validar_hora_para_autorizador(hora, asignacion_activa, es_admin)
            hora.estado = "APROBADA"
            hora.aprobado_por = autorizador.id
            hora.observacion_rechazo = None
            hora.fecha_resolucion = fecha_resolucion
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudieron autorizar las horas seleccionadas.") from error
    return RedirectResponse(
        f"/autorizaciones?guardado={len(horas)}"
        + (f"&usuario_id={filtro_usuario_id}" if filtro_usuario_id else ""),
        status_code=303,
    )


@app.post("/autorizaciones/{hora_id}/eliminar")
def eliminar_hora_pendiente(
    hora_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    obtener_autorizador(request, db)
    raise HTTPException(
        status_code=403,
        detail="Los autorizadores sólo pueden aprobar o rechazar cargas.",
    )


@app.post("/autorizaciones/rechazar")
def rechazar_horas(
    request: Request,
    hora_ids: Annotated[list[int] | None, Form()] = None,
    fila_ids: Annotated[list[int] | None, Form()] = None,
    observaciones_rechazo: Annotated[list[str] | None, Form()] = None,
    filtro_usuario_id: Annotated[int | None, Form()] = None,
    db: Session = Depends(get_db),
):
    ids = list(dict.fromkeys(hora_ids or []))
    if not ids:
        raise HTTPException(status_code=422, detail="Seleccioná al menos una carga para rechazar.")
    filas = list(fila_ids or [])
    observaciones = list(observaciones_rechazo or [])
    if len(filas) != len(observaciones):
        raise HTTPException(
            status_code=422,
            detail="Las observaciones de rechazo no coinciden con las cargas mostradas.",
        )
    observacion_por_hora = {
        hora_id: clean_optional(observacion)
        for hora_id, observacion in zip(filas, observaciones)
    }
    autorizador, asignacion_activa, es_admin = obtener_autorizador(request, db)
    try:
        ids = expandir_ids_con_jornadas_pendientes(ids, db)
        horas = list(db.scalars(
            select(HoraExtra)
            .options(selectinload(HoraExtra.usuario).selectinload(Usuario.roles))
            .where(HoraExtra.id.in_(ids))
            .with_for_update()
        ))
        if len(horas) != len(ids):
            raise HTTPException(status_code=404, detail="Una de las cargas seleccionadas no existe.")
        fecha_resolucion = datetime.now()
        for hora in horas:
            validar_hora_para_autorizador(hora, asignacion_activa, es_admin)
            hora.estado = "RECHAZADA"
            hora.aprobado_por = autorizador.id
            hora.observacion_rechazo = observacion_por_hora.get(hora.id)
            hora.fecha_resolucion = fecha_resolucion
        db.commit()
    except SQLAlchemyError as error:
        db.rollback()
        raise HTTPException(status_code=500, detail="No se pudieron rechazar las horas seleccionadas.") from error
    return RedirectResponse(
        f"/autorizaciones?rechazado={len(horas)}"
        + (f"&usuario_id={filtro_usuario_id}" if filtro_usuario_id else ""),
        status_code=303,
    )


@app.post("/autorizaciones/{hora_id}/editar")
def editar_hora_pendiente(
    hora_id: int,
    request: Request,
    db: Session = Depends(get_db),
):
    obtener_autorizador(request, db)
    raise HTTPException(
        status_code=403,
        detail="Los autorizadores sólo pueden aprobar o rechazar cargas.",
    )


@app.get("/health")
def health_check():
    return {"estado": "ok", "aplicacion": "activa"}


@app.get("/health/db")
@app.get("/health/database")
def database_health_check():
    try:
        with engine.connect() as connection:
            connection.execute(text("SELECT 1"))
        return {"estado": "ok", "base_de_datos": "conectada"}
    except SQLAlchemyError as error:
        raise HTTPException(
            status_code=503,
            detail={"estado": "error", "base_de_datos": "desconectada"},
        ) from error
